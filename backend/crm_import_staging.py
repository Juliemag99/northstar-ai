"""Administrator company/contact import staging (Checkpoint B).

Upload and read-only preview only. Never writes companies, contacts,
relationships, statuses, workflows, activities, or notes. Never stores the
original file. Never evaluates spreadsheet formulas.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import re
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from models import (
    CrmImportBatchView,
    CrmImportRowsPage,
    CrmImportRowView,
    CrmImportUploadResult,
    NorthStarUser,
)

log = logging.getLogger("northstar.crm_import")

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ZIP_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ZIP_MEMBERS = 50
MAX_COLUMNS = 40
MAX_SOURCE_ROWS = 10_000
MAX_CELL_CHARS = 2_000
CELL_PREVIEW_CHARS = 200
SAMPLE_ROWS = 25
MAX_PAGE_SIZE = 100
RETENTION_DAYS = 7
CSV_WORKSHEET_NAME = "CSV"

STATUS_PREVIEWED = "previewed"
STATUS_FAILED = "failed"
STATUS_CANCELLED = "cancelled"
STATUS_EXPIRED = "expired"
USABLE_STATUS = STATUS_PREVIEWED

GENERIC_PARSE_ERROR = "The spreadsheet could not be read. Check the file and try again."
TOO_MANY_COLUMNS = (
    f"This spreadsheet has more than {MAX_COLUMNS} columns. "
    "Reduce columns and upload again."
)
EXTRA_SOURCE_COLUMNS = (
    "This spreadsheet has values in columns beyond the header row. "
    "Give every source column a header, or remove the extra columns."
)
TOO_MANY_ROWS = (
    f"This spreadsheet has more than {MAX_SOURCE_ROWS} data rows. "
    "Split the file and upload again."
)
FILE_TOO_LARGE = "This file is too large to upload."
UNSUPPORTED_TYPE = "Upload a CSV or XLSX file."
NEED_WORKSHEET = "This workbook has multiple sheets. Choose one visible sheet to continue."
HIDDEN_SHEET = "Hidden worksheets cannot be imported. Choose a visible sheet."
BATCH_NOT_REUSABLE = "This import batch can no longer be previewed."
CLIENT_REQUIRED = "Choose a specific client before uploading."
MAPPING_COMPANY_REQUIRED = "Map a Company name column before continuing."
MAPPING_UNKNOWN_FIELD = "That destination field is not supported."
MAPPING_HEADER_MISSING = "That source column is not in this spreadsheet."
MAPPING_EMPTY_HEADER = "Each mapped column must have a name."
MAPPING_DUPLICATE_HEADER = "Each source column can map to only one field."
MAPPING_NAME_CONFLICT = (
    "Use either a full name column or first and last name columns, not both."
)
MAPPING_NO_ROWS = "This import has no staged rows to map."

CANONICAL_MAPPING_FIELDS = frozenset(
    {
        "company_name",
        "website",
        "phone",
        "address",
        "city",
        "state",
        "zip",
        "contact_first_name",
        "contact_last_name",
        "contact_full_name",
        "contact_title",
        "contact_email",
        "contact_phone",
    }
)


class BatchNotReusable(ValueError):
    pass


class WorksheetSelectionNeeded(Exception):
    def __init__(self, *, filename: str, file_type: str, visible_sheets: list[str]):
        super().__init__(NEED_WORKSHEET)
        self.filename = filename
        self.file_type = file_type
        self.visible_sheets = visible_sheets


def _now() -> datetime:
    return datetime.now(timezone.utc).replace(microsecond=0)


def _iso(value: datetime) -> str:
    return value.strftime("%Y-%m-%dT%H:%M:%SZ")


def _blank(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _safe_filename(name: str) -> str:
    base = Path(_blank(name) or "upload").name
    base = re.sub(r"[^\w.\- ()]+", "_", base).strip("._ ") or "upload"
    return base[:180]


def _file_type(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return "csv"
    if suffix == ".xlsx":
        return "xlsx"
    return ""


def _json_list(value: object) -> list[str]:
    if isinstance(value, list):
        return [str(v) for v in value if str(v).strip()]
    text = _blank(value)
    if not text:
        return []
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        return []
    if isinstance(parsed, list):
        return [str(v) for v in parsed if str(v).strip()]
    return []


def _json_dict(value: object) -> dict[str, str]:
    text = _blank(value)
    if not text:
        return {}
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        return {}
    if not isinstance(parsed, dict):
        return {}
    return {str(k): "" if v is None else str(v) for k, v in parsed.items()}


def _json_mapping(value: object) -> dict[str, str]:
    parsed = _json_dict(value)
    return {str(k): str(v) for k, v in parsed.items() if str(k).strip() and str(v).strip()}


def _is_formula_text(text: str) -> bool:
    stripped = text.lstrip()
    return stripped.startswith("=")


def ensure_crm_import_schema(conn=None) -> None:
    owns = conn is None
    if owns:
        from db import get_connection

        conn = get_connection()
    try:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS crm_import_batches (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                client_id INTEGER NOT NULL,
                uploaded_by_user_id INTEGER,
                uploaded_by_name TEXT NOT NULL DEFAULT '',
                original_filename TEXT NOT NULL DEFAULT '',
                file_type TEXT NOT NULL DEFAULT '',
                worksheet_name TEXT NOT NULL DEFAULT '',
                file_size_bytes INTEGER NOT NULL DEFAULT 0,
                sha256 TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL DEFAULT 'previewed',
                error_message TEXT NOT NULL DEFAULT '',
                headers_json TEXT NOT NULL DEFAULT '[]',
                warnings_json TEXT NOT NULL DEFAULT '[]',
                total_rows INTEGER NOT NULL DEFAULT 0,
                source_row_count INTEGER NOT NULL DEFAULT 0,
                blank_row_count INTEGER NOT NULL DEFAULT 0,
                error_row_count INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL DEFAULT '',
                updated_at TEXT NOT NULL DEFAULT '',
                cancelled_at TEXT NOT NULL DEFAULT '',
                expires_at TEXT NOT NULL DEFAULT '',
                mapping_json TEXT NOT NULL DEFAULT '{}',
                mapping_updated_at TEXT NOT NULL DEFAULT '',
                mapping_updated_by_user_id INTEGER,
                FOREIGN KEY (client_id) REFERENCES clients(id) ON DELETE CASCADE,
                FOREIGN KEY (uploaded_by_user_id) REFERENCES users(id) ON DELETE SET NULL,
                FOREIGN KEY (mapping_updated_by_user_id) REFERENCES users(id) ON DELETE SET NULL
            );

            CREATE TABLE IF NOT EXISTS crm_import_rows (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                batch_id INTEGER NOT NULL,
                client_id INTEGER NOT NULL,
                source_row_number INTEGER NOT NULL DEFAULT 0,
                raw_json TEXT NOT NULL DEFAULT '{}',
                warnings_json TEXT NOT NULL DEFAULT '[]',
                errors_json TEXT NOT NULL DEFAULT '[]',
                is_blank INTEGER NOT NULL DEFAULT 0,
                has_blocking_error INTEGER NOT NULL DEFAULT 0,
                FOREIGN KEY (batch_id) REFERENCES crm_import_batches(id) ON DELETE CASCADE,
                FOREIGN KEY (client_id) REFERENCES clients(id) ON DELETE CASCADE
            );

            CREATE INDEX IF NOT EXISTS idx_crm_import_rows_batch
                ON crm_import_rows(batch_id, source_row_number);
            CREATE INDEX IF NOT EXISTS idx_crm_import_batches_client
                ON crm_import_batches(client_id, status, expires_at);
            """
        )
        existing = {
            str(r["name"])
            for r in conn.execute("PRAGMA table_info(crm_import_batches)").fetchall()
        }
        for name, declaration in (
            ("mapping_json", "TEXT NOT NULL DEFAULT '{}'"),
            ("mapping_updated_at", "TEXT NOT NULL DEFAULT ''"),
            ("mapping_updated_by_user_id", "INTEGER"),
        ):
            if name not in existing:
                conn.execute(f"ALTER TABLE crm_import_batches ADD COLUMN {name} {declaration}")
        if owns:
            conn.commit()
    finally:
        if owns:
            conn.close()


def purge_expired_staging_rows(conn=None) -> int:
    """Delete expired staged rows; keep sanitized batch audit rows. Never touches CRM."""
    owns = conn is None
    if owns:
        from db import get_connection

        conn = get_connection()
    try:
        ensure_crm_import_schema(conn)
        now = _iso(_now())
        rows = conn.execute(
            """
            SELECT id FROM crm_import_batches
            WHERE expires_at != '' AND expires_at < ?
              AND status NOT IN (?, ?)
            """,
            (now, STATUS_CANCELLED, STATUS_EXPIRED),
        ).fetchall()
        ids = [int(r["id"]) for r in rows]
        if not ids:
            return 0
        placeholders = ",".join("?" * len(ids))
        conn.execute(
            f"DELETE FROM crm_import_rows WHERE batch_id IN ({placeholders})",
            ids,
        )
        conn.execute(
            f"""
            UPDATE crm_import_batches SET
                status = ?,
                headers_json = '[]',
                warnings_json = '[]',
                updated_at = ?
            WHERE id IN ({placeholders})
            """,
            [STATUS_EXPIRED, now, *ids],
        )
        if owns:
            conn.commit()
        log.info("purged expired crm import staging rows batches=%s", len(ids))
        return len(ids)
    finally:
        if owns:
            conn.close()


def _sanitize_headers(headers: list[str]) -> tuple[list[str], list[str]]:
    warnings: list[str] = []
    seen: dict[str, int] = {}
    out: list[str] = []
    for index, raw in enumerate(headers, start=1):
        name = _blank(raw)
        if not name:
            name = f"Column_{index}"
            warnings.append(f"Blank header in column {index} was labeled {name}.")
        key = name
        count = seen.get(key.lower(), 0)
        if count:
            name = f"{key}_{count + 1}"
            warnings.append(f"Duplicate header {key!r} was renamed to {name}.")
        seen[key.lower()] = count + 1
        out.append(name[:120])
    return out, warnings


def _cell_record(header: str, raw: object, *, formula: bool) -> tuple[str, list[str], list[str]]:
    warnings: list[str] = []
    errors: list[str] = []
    text = "" if raw is None else str(raw)
    if formula or _is_formula_text(text):
        warnings.append(f"Formula in {header} was stored as text and was not calculated.")
        text = text if text else str(raw or "")
    if len(text) > MAX_CELL_CHARS:
        errors.append(
            f"{header} exceeds {MAX_CELL_CHARS} characters. Re-upload a shorter value before import."
        )
        text = text[:CELL_PREVIEW_CHARS] + "…"
    return text, warnings, errors


def _row_values(
    headers: list[str], cells: list[tuple[object, bool]]
) -> tuple[dict[str, str], list[str], list[str], bool]:
    values: dict[str, str] = {}
    warnings: list[str] = []
    errors: list[str] = []
    for header, (raw, formula) in zip(headers, cells):
        stored, cell_warnings, cell_errors = _cell_record(header, raw, formula=formula)
        values[header] = stored
        warnings.extend(cell_warnings)
        errors.extend(cell_errors)
    return values, warnings, errors, bool(errors)


def _decode_csv_bytes(content: bytes) -> tuple[str, list[str]]:
    warnings: list[str] = []
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            return content.decode(encoding), warnings
        except UnicodeDecodeError:
            continue
    try:
        text = content.decode("cp1252")
        warnings.append("File encoding was treated as Windows-1252.")
        return text, warnings
    except UnicodeDecodeError as exc:
        raise ValueError(GENERIC_PARSE_ERROR) from exc


def _parse_csv(content: bytes) -> tuple[list[str], list[dict[str, Any]], dict[str, int], list[str]]:
    text, warnings = _decode_csv_bytes(content)
    sample = text[:8192]
    dialect = csv.excel
    if sample.strip():
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
            if getattr(dialect, "delimiter", ",") in {"\t", ";"}:
                warnings.append("A non-comma delimiter was detected.")
        except csv.Error:
            dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = list(reader)
    return _materialize_grid(rows, warnings)


def _xlsx_zip_guard(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            names = zf.namelist()
            if len(names) > MAX_ZIP_MEMBERS:
                raise ValueError(GENERIC_PARSE_ERROR)
            uncompressed = sum(max(0, int(info.file_size)) for info in zf.infolist())
            if uncompressed > MAX_ZIP_UNCOMPRESSED_BYTES:
                raise ValueError(FILE_TOO_LARGE)
            if "xl/vbaProject.bin" in names:
                raise ValueError(UNSUPPORTED_TYPE)
    except zipfile.BadZipFile as exc:
        raise ValueError(GENERIC_PARSE_ERROR) from exc


def _parse_xlsx(
    content: bytes, *, worksheet: str
) -> tuple[list[str], list[dict[str, Any]], dict[str, int], list[str], str]:
    _xlsx_zip_guard(content)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError(GENERIC_PARSE_ERROR) from exc

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_vba=False)
    try:
        visible: list[str] = []
        hidden: set[str] = set()
        for ws in wb.worksheets:
            state = _blank(getattr(ws, "sheet_state", "visible")).lower()
            if state in {"hidden", "veryhidden"}:
                hidden.add(ws.title)
            else:
                visible.append(ws.title)
        if not visible:
            raise ValueError(GENERIC_PARSE_ERROR)
        requested = _blank(worksheet)
        if requested:
            if requested in hidden:
                raise ValueError(HIDDEN_SHEET)
            if requested not in visible:
                raise ValueError(NEED_WORKSHEET)
            sheet_name = requested
        elif len(visible) == 1:
            sheet_name = visible[0]
        else:
            raise WorksheetSelectionNeeded(
                filename="",
                file_type="xlsx",
                visible_sheets=visible,
            )
        ws = wb[sheet_name]
        grid: list[list[tuple[object, bool]]] = []
        for row in ws.iter_rows():
            cells: list[tuple[object, bool]] = []
            for cell in row:
                formula = str(getattr(cell, "data_type", "")).lower() == "f"
                cells.append((cell.value, formula))
            grid.append(cells)
        warnings: list[str] = []
        if hidden:
            warnings.append("Hidden worksheets were ignored and cannot be selected.")
        headers, staged, counts, extra_warnings = _materialize_typed_grid(grid, warnings)
        return headers, staged, counts, extra_warnings, sheet_name
    finally:
        wb.close()


def _materialize_grid(
    rows: list[list[Any]], warnings: list[str]
) -> tuple[list[str], list[dict[str, Any]], dict[str, int], list[str]]:
    typed: list[list[tuple[object, bool]]] = []
    for row in rows:
        typed.append([(cell, _is_formula_text(_blank(cell))) for cell in row])
    return _materialize_typed_grid(typed, warnings)


def _materialize_typed_grid(
    rows: list[list[tuple[object, bool]]], warnings: list[str]
) -> tuple[list[str], list[dict[str, Any]], dict[str, int], list[str]]:
    trailing_blank = 0
    while len(rows) > 1 and not any(_blank(c[0]) or c[1] for c in rows[-1]):
        rows.pop()
        trailing_blank += 1
    if not rows:
        raise ValueError(GENERIC_PARSE_ERROR)
    header_cells = [c[0] for c in rows[0]]
    if len(header_cells) > MAX_COLUMNS:
        raise ValueError(TOO_MANY_COLUMNS)
    headers, header_warnings = _sanitize_headers([_blank(h) for h in header_cells])
    warnings.extend(header_warnings)
    if not any(headers):
        raise ValueError(GENERIC_PARSE_ERROR)

    total_rows = 0
    blank_row_count = 0
    staged: list[dict[str, Any]] = []
    for offset, raw_row in enumerate(rows[1:], start=2):
        total_rows += 1
        if len(raw_row) > MAX_COLUMNS:
            raise ValueError(TOO_MANY_COLUMNS)
        extra = raw_row[len(headers) :]
        if any(_blank(c[0]) or c[1] for c in extra):
            raise ValueError(EXTRA_SOURCE_COLUMNS)
        padded = list(raw_row[: len(headers)])
        while len(padded) < len(headers):
            padded.append(("", False))
        if not any(_blank(c[0]) or c[1] for c in padded):
            blank_row_count += 1
            continue
        if len(staged) >= MAX_SOURCE_ROWS:
            raise ValueError(TOO_MANY_ROWS)
        values, row_warnings, row_errors, blocking = _row_values(headers, padded)
        staged.append(
            {
                "source_row_number": offset,
                "values": values,
                "warnings": row_warnings,
                "errors": row_errors,
                "has_blocking_error": blocking,
            }
        )
    counts = {
        "total_rows": total_rows + trailing_blank,
        "source_row_count": len(staged),
        "blank_row_count": blank_row_count + trailing_blank,
        "error_row_count": sum(1 for r in staged if r["has_blocking_error"]),
    }
    return headers, staged, counts, warnings


def _insert_batch(
    conn,
    *,
    client_id: int,
    actor: NorthStarUser,
    filename: str,
    file_type: str,
    worksheet_name: str,
    size: int,
    digest: str,
    status: str,
    error_message: str,
    headers: list[str],
    warnings: list[str],
    counts: dict[str, int],
    staged: list[dict[str, Any]],
) -> int:
    now = _now()
    expires = now + timedelta(days=RETENTION_DAYS)
    cur = conn.execute(
        """
        INSERT INTO crm_import_batches (
            client_id, uploaded_by_user_id, uploaded_by_name,
            original_filename, file_type, worksheet_name, file_size_bytes, sha256,
            status, error_message, headers_json, warnings_json,
            total_rows, source_row_count, blank_row_count, error_row_count,
            created_at, updated_at, cancelled_at, expires_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, '', ?)
        """,
        (
            client_id,
            int(actor.id),
            _blank(actor.full_name) or _blank(actor.email),
            filename,
            file_type,
            worksheet_name,
            size,
            digest,
            status,
            error_message,
            json.dumps(headers),
            json.dumps(warnings),
            int(counts.get("total_rows") or 0),
            int(counts.get("source_row_count") or 0),
            int(counts.get("blank_row_count") or 0),
            int(counts.get("error_row_count") or 0),
            _iso(now),
            _iso(now),
            _iso(expires),
        ),
    )
    batch_id = int(cur.lastrowid)
    if staged:
        conn.executemany(
            """
            INSERT INTO crm_import_rows (
                batch_id, client_id, source_row_number, raw_json,
                warnings_json, errors_json, is_blank, has_blocking_error
            ) VALUES (?, ?, ?, ?, ?, ?, 0, ?)
            """,
            [
                (
                    batch_id,
                    client_id,
                    int(row["source_row_number"]),
                    json.dumps(row["values"], separators=(",", ":")),
                    json.dumps(row["warnings"]),
                    json.dumps(row["errors"]),
                    1 if row["has_blocking_error"] else 0,
                )
                for row in staged
            ],
        )
    return batch_id


def _batch_view(conn, row, *, include_sample: bool) -> CrmImportBatchView:
    batch_id = int(row["id"])
    status = _blank(row["status"]) or STATUS_FAILED
    reusable = status == USABLE_STATUS
    headers = _json_list(row["headers_json"]) if reusable else []
    warnings = _json_list(row["warnings_json"]) if reusable else []
    sample: list[CrmImportRowView] = []
    if include_sample and reusable:
        sample_rows = conn.execute(
            """
            SELECT * FROM crm_import_rows
            WHERE batch_id = ?
            ORDER BY source_row_number, id
            LIMIT ?
            """,
            (batch_id, SAMPLE_ROWS),
        ).fetchall()
        sample = [_row_view(r) for r in sample_rows]
    return CrmImportBatchView(
        batch_id=batch_id,
        client_id=int(row["client_id"]),
        status=status,
        original_filename=_blank(row["original_filename"]),
        file_type=_blank(row["file_type"]),
        worksheet_name=_blank(row["worksheet_name"]),
        file_size_bytes=int(row["file_size_bytes"] or 0),
        sha256=_blank(row["sha256"]),
        headers=headers,
        warnings=warnings,
        error_message=_blank(row["error_message"]),
        total_rows=int(row["total_rows"] or 0),
        source_row_count=int(row["source_row_count"] or 0),
        blank_row_count=int(row["blank_row_count"] or 0),
        error_row_count=int(row["error_row_count"] or 0),
        reusable=reusable,
        expires_at=_blank(row["expires_at"]),
        created_at=_blank(row["created_at"]),
        updated_at=_blank(row["updated_at"]),
        cancelled_at=_blank(row["cancelled_at"]),
        uploaded_by_user_id=(
            int(row["uploaded_by_user_id"])
            if row["uploaded_by_user_id"] is not None
            else None
        ),
        uploaded_by_name=_blank(row["uploaded_by_name"]),
        mapping=_json_mapping(
            row["mapping_json"]
            if reusable and "mapping_json" in row.keys()
            else "{}"
        )
        if reusable
        else {},
        mapping_updated_at=_blank(row["mapping_updated_at"]) if "mapping_updated_at" in row.keys() else "",
        mapping_updated_by_user_id=(
            int(row["mapping_updated_by_user_id"])
            if reusable
            and "mapping_updated_by_user_id" in row.keys()
            and row["mapping_updated_by_user_id"] is not None
            else None
        ),
        sample_rows=sample,
    )


def _row_view(row) -> CrmImportRowView:
    return CrmImportRowView(
        row_id=int(row["id"]),
        source_row_number=int(row["source_row_number"] or 0),
        values=_json_dict(row["raw_json"]),
        warnings=_json_list(row["warnings_json"]),
        errors=_json_list(row["errors_json"]),
        has_blocking_error=bool(row["has_blocking_error"]),
    )


def _load_batch_row(conn, client_id: int, batch_id: int):
    return conn.execute(
        """
        SELECT * FROM crm_import_batches
        WHERE id = ? AND client_id = ?
        """,
        (batch_id, client_id),
    ).fetchone()


def upload_crm_import(
    *,
    client_id: int,
    actor: NorthStarUser,
    filename: str,
    content: bytes,
    worksheet: str = "",
) -> CrmImportUploadResult:
    if client_id <= 0:
        raise ValueError(CLIENT_REQUIRED)
    if actor is None or not bool(actor.active) or not bool(actor.is_administrator):
        raise PermissionError("Not authorized.")
    if content is None or len(content) == 0:
        raise ValueError("Choose a CSV or XLSX file to upload.")
    if len(content) > MAX_FILE_BYTES:
        raise ValueError(FILE_TOO_LARGE)

    safe_name = _safe_filename(filename)
    file_type = _file_type(safe_name)
    if file_type not in {"csv", "xlsx"}:
        raise ValueError(UNSUPPORTED_TYPE)
    digest = hashlib.sha256(content).hexdigest()
    log.info(
        "crm import upload client_id=%s bytes=%s type=%s",
        client_id,
        len(content),
        file_type,
    )

    from db import get_connection

    try:
        if file_type == "csv":
            headers, staged, counts, warnings = _parse_csv(content)
            sheet_name = CSV_WORKSHEET_NAME
        else:
            try:
                headers, staged, counts, warnings, sheet_name = _parse_xlsx(
                    content, worksheet=worksheet
                )
            except WorksheetSelectionNeeded as exc:
                return CrmImportUploadResult(
                    kind="needs_worksheet",
                    needs_worksheet=True,
                    visible_sheets=list(exc.visible_sheets),
                    filename=safe_name,
                    file_type="xlsx",
                    message=NEED_WORKSHEET,
                )
    except ValueError as exc:
        message = str(exc) or GENERIC_PARSE_ERROR
        with get_connection() as conn:
            ensure_crm_import_schema(conn)
            purge_expired_staging_rows(conn)
            batch_id = _insert_batch(
                conn,
                client_id=client_id,
                actor=actor,
                filename=safe_name,
                file_type=file_type,
                worksheet_name=_blank(worksheet) or (CSV_WORKSHEET_NAME if file_type == "csv" else ""),
                size=len(content),
                digest=digest,
                status=STATUS_FAILED,
                error_message=message,
                headers=[],
                warnings=[],
                counts={
                    "total_rows": 0,
                    "source_row_count": 0,
                    "blank_row_count": 0,
                    "error_row_count": 0,
                },
                staged=[],
            )
        batch = get_crm_import_batch(client_id, batch_id, include_sample=False)
        return CrmImportUploadResult(
            kind="failed",
            filename=safe_name,
            file_type=file_type,
            message=message,
            batch=batch,
        )

    with get_connection() as conn:
        ensure_crm_import_schema(conn)
        purge_expired_staging_rows(conn)
        batch_id = _insert_batch(
            conn,
            client_id=client_id,
            actor=actor,
            filename=safe_name,
            file_type=file_type,
            worksheet_name=sheet_name,
            size=len(content),
            digest=digest,
            status=STATUS_PREVIEWED,
            error_message="",
            headers=headers,
            warnings=warnings,
            counts=counts,
            staged=staged,
        )
    batch = get_crm_import_batch(client_id, batch_id, include_sample=True)
    return CrmImportUploadResult(
        kind="previewed",
        filename=safe_name,
        file_type=file_type,
        message="Preview ready. Nothing was written to companies or contacts.",
        batch=batch,
    )


def get_crm_import_batch(
    client_id: int, batch_id: int, *, include_sample: bool = True
) -> CrmImportBatchView:
    from db import get_connection

    with get_connection() as conn:
        ensure_crm_import_schema(conn)
        purge_expired_staging_rows(conn)
        row = _load_batch_row(conn, client_id, batch_id)
        if row is None:
            raise LookupError("Import batch not found.")
        return _batch_view(conn, row, include_sample=include_sample)


def list_crm_import_rows(
    client_id: int, batch_id: int, *, offset: int = 0, limit: int = SAMPLE_ROWS
) -> CrmImportRowsPage:
    from db import get_connection

    page = max(1, min(int(limit or SAMPLE_ROWS), MAX_PAGE_SIZE))
    start = max(0, int(offset or 0))
    with get_connection() as conn:
        ensure_crm_import_schema(conn)
        purge_expired_staging_rows(conn)
        row = _load_batch_row(conn, client_id, batch_id)
        if row is None:
            raise LookupError("Import batch not found.")
        if _blank(row["status"]) != USABLE_STATUS:
            raise BatchNotReusable(BATCH_NOT_REUSABLE)
        total = int(
            conn.execute(
                "SELECT COUNT(*) AS n FROM crm_import_rows WHERE batch_id = ?",
                (batch_id,),
            ).fetchone()["n"]
        )
        rows = conn.execute(
            """
            SELECT * FROM crm_import_rows
            WHERE batch_id = ?
            ORDER BY source_row_number, id
            LIMIT ? OFFSET ?
            """,
            (batch_id, page, start),
        ).fetchall()
        return CrmImportRowsPage(
            batch_id=batch_id,
            client_id=client_id,
            offset=start,
            limit=page,
            total=total,
            rows=[_row_view(r) for r in rows],
        )


def cancel_crm_import(client_id: int, batch_id: int, *, actor: NorthStarUser) -> CrmImportBatchView:
    if actor is None or not bool(actor.active) or not bool(actor.is_administrator):
        raise PermissionError("Not authorized.")
    from db import get_connection

    now = _iso(_now())
    with get_connection() as conn:
        ensure_crm_import_schema(conn)
        purge_expired_staging_rows(conn)
        row = _load_batch_row(conn, client_id, batch_id)
        if row is None:
            raise LookupError("Import batch not found.")
        status = _blank(row["status"])
        if status == STATUS_CANCELLED:
            return _batch_view(conn, row, include_sample=False)
        if status not in {STATUS_PREVIEWED, STATUS_FAILED}:
            raise BatchNotReusable(BATCH_NOT_REUSABLE)
        conn.execute("DELETE FROM crm_import_rows WHERE batch_id = ? AND client_id = ?", (batch_id, client_id))
        conn.execute(
            """
            UPDATE crm_import_batches SET
                status = ?,
                cancelled_at = ?,
                updated_at = ?,
                headers_json = '[]',
                warnings_json = '[]',
                error_message = CASE
                    WHEN error_message = '' THEN 'Cancelled. Staged source rows were removed.'
                    ELSE error_message
                END
            WHERE id = ? AND client_id = ?
            """,
            (STATUS_CANCELLED, now, now, batch_id, client_id),
        )
    return get_crm_import_batch(client_id, batch_id, include_sample=False)


def _exact_header(headers: list[str], submitted: str) -> str | None:
    for header in headers:
        if header == submitted:
            return header
    for header in headers:
        if header.strip() == submitted:
            return header
    return None


def _normalize_mapping(raw: dict[str, str], headers: list[str]) -> dict[str, str]:
    if not isinstance(raw, dict):
        raise ValueError(MAPPING_UNKNOWN_FIELD)
    normalized: dict[str, str] = {}
    used_headers: dict[str, str] = {}
    for dest_raw, src_raw in raw.items():
        dest = _blank(dest_raw)
        if dest not in CANONICAL_MAPPING_FIELDS:
            raise ValueError(MAPPING_UNKNOWN_FIELD)
        submitted = _blank(src_raw)
        if not submitted:
            raise ValueError(MAPPING_EMPTY_HEADER)
        exact = _exact_header(headers, submitted)
        if exact is None:
            raise ValueError(MAPPING_HEADER_MISSING)
        prior = used_headers.get(exact)
        if prior and prior != dest:
            raise ValueError(MAPPING_DUPLICATE_HEADER)
        used_headers[exact] = dest
        normalized[dest] = exact
    if "company_name" not in normalized:
        raise ValueError(MAPPING_COMPANY_REQUIRED)
    if "contact_full_name" in normalized and (
        "contact_first_name" in normalized or "contact_last_name" in normalized
    ):
        raise ValueError(MAPPING_NAME_CONFLICT)
    return normalized


def save_crm_import_mapping(
    client_id: int,
    batch_id: int,
    *,
    actor: NorthStarUser,
    mapping: dict[str, str],
) -> CrmImportBatchView:
    if actor is None or not bool(actor.active) or not bool(actor.is_administrator):
        raise PermissionError("Not authorized.")
    from db import get_connection

    now = _iso(_now())
    with get_connection() as conn:
        ensure_crm_import_schema(conn)
        purge_expired_staging_rows(conn)
        row = _load_batch_row(conn, client_id, batch_id)
        if row is None:
            raise LookupError("Import batch not found.")
        if _blank(row["status"]) != USABLE_STATUS:
            raise BatchNotReusable(BATCH_NOT_REUSABLE)
        remaining = int(
            conn.execute(
                "SELECT COUNT(*) AS n FROM crm_import_rows WHERE batch_id = ? AND client_id = ?",
                (batch_id, client_id),
            ).fetchone()["n"]
        )
        if remaining <= 0:
            raise BatchNotReusable(MAPPING_NO_ROWS)
        headers = _json_list(row["headers_json"])
        normalized = _normalize_mapping(mapping or {}, headers)
        log.info(
            "crm import mapping client_id=%s batch_id=%s fields=%s",
            client_id,
            batch_id,
            ",".join(sorted(normalized)),
        )
        conn.execute(
            """
            UPDATE crm_import_batches SET
                mapping_json = ?,
                mapping_updated_at = ?,
                mapping_updated_by_user_id = ?,
                updated_at = ?
            WHERE id = ? AND client_id = ?
            """,
            (
                json.dumps(normalized, separators=(",", ":")),
                now,
                int(actor.id),
                now,
                batch_id,
                client_id,
            ),
        )
    return get_crm_import_batch(client_id, batch_id, include_sample=True)

