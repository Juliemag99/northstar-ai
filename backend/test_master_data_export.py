"""Administrator Master Data Export.

Run: python test_master_data_export.py

Uses isolated testdb copies only — never writes production northstar.db,
never talks to live :8007, and never stores a real staff password in source.
"""

from __future__ import annotations

import testdb
import csv
import io
import json
import os
import secrets
import sys
from datetime import date
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from auth_http import ADMIN_REQUIRED_DETAIL, AUTH_REQUIRED_DETAIL
from auth_passwords import hash_password
from db import get_connection
from main import app
from master_data_export import (
    CAMPAIGN_HEADERS,
    COMPANY_HEADERS,
    CONTACT_HEADERS,
    RELATIONSHIP_HEADERS,
    SHEET_CAMPAIGNS,
    SHEET_COMPANIES,
    SHEET_CONTACTS,
    SHEET_RELATIONSHIPS,
    snapshot_export_counts,
)

EXPORT = "/api/admin/master-data-export"
MARKER = "NSMDE"
CARMECO_ID = 1
BROWN_ID = 2


def _fail(msg: str) -> None:
    raise AssertionError(msg)


def _dump(payload: object) -> str:
    return json.dumps(payload, default=str)


def _assert_no_secrets(payload: object, *secrets_out: str) -> None:
    dumped = _dump(payload).lower()
    if "password_hash" in dumped:
        _fail("JSON leaked password_hash.")
    if "$argon2id$" in dumped:
        _fail("JSON leaked an Argon2 hash.")
    if "csrf_secret" in dumped:
        _fail("JSON leaked csrf_secret.")
    for item in secrets_out:
        if item and str(item).lower() in dumped:
            _fail("JSON leaked a raw secret.")


def _secret_password() -> str:
    return f"NsTest9{secrets.token_hex(10)}"


def _client() -> TestClient:
    return TestClient(app)


def _create_user(*, email: str, password: str, administrator: int = 0, active: int = 1) -> int:
    digest = hash_password(password, email=email)
    with get_connection() as conn:
        conn.execute(
            """
            INSERT INTO users (
                email, full_name, is_administrator, is_internal_northstar, active,
                password_hash, failed_login_count, locked_until
            ) VALUES (?, ?, ?, 1, ?, ?, 0, '')
            """,
            (email, "Master Export Test User", int(administrator), int(active), digest),
        )
        user_id = int(
            conn.execute("SELECT id FROM users WHERE email = ?", (email,)).fetchone()["id"]
        )
        conn.commit()
    return user_id


def _delete_user(user_id: int) -> None:
    with get_connection() as conn:
        conn.execute("DELETE FROM staff_sessions WHERE user_id = ?", (user_id,))
        conn.execute("DELETE FROM user_client_assignments WHERE user_id = ?", (user_id,))
        conn.execute("DELETE FROM users WHERE id = ?", (user_id,))
        conn.commit()


def _login(http: TestClient, email: str, password: str):
    return http.post("/api/auth/login", json={"email": email, "password": password})


def _julie_id() -> int:
    with get_connection() as conn:
        row = conn.execute(
            "SELECT id FROM users WHERE lower(trim(email)) = 'juliem@n-star.us' LIMIT 1"
        ).fetchone()
    if row is None:
        _fail("Isolated testdb needs Julie user.")
    return int(row["id"])


def _insert_fixtures() -> dict[str, Any]:
    stamp = secrets.token_hex(4)
    master_rn = f"{MARKER}-{stamp}-MASTER"
    ccr_rn = f"{MARKER}-{stamp}-CCR"
    empty_rn = f"{MARKER}-{stamp}-EMPTY"
    contact_rn = f"{MARKER}-{stamp}-CONTACT"
    campaign_name = f"{MARKER} Campaign {stamp}"
    julie_id = _julie_id()
    ids: dict[str, int] = {}
    with get_connection() as conn:
        cur = conn.execute(
            """
            INSERT INTO companies (
                external_record_no, company_name, website, legacy_phone, legacy_alt_phone,
                address, city, state, zip, type_of_industry, source
            ) VALUES (?, ?, 'multi.example.test', '555-0100', '555-0101',
                      '1 Export Way', 'Cincinnati', 'OH', '45202', 'Stamping', 'test')
            """,
            (master_rn, f"{MARKER} Multi {stamp}"),
        )
        ids["multi_company_id"] = int(cur.lastrowid)
        conn.execute(
            """
            INSERT INTO client_company_relationships (
                client_id, company_id, external_record_no, status, assigned_user_id, notes, is_hot
            ) VALUES (?, ?, ?, 'Contacted', ?, 'carmeco note', 0)
            """,
            (CARMECO_ID, ids["multi_company_id"], ccr_rn, julie_id),
        )
        conn.execute(
            """
            INSERT INTO client_company_relationships (
                client_id, company_id, external_record_no, status, assigned_user_id, notes, is_hot
            ) VALUES (?, ?, ?, 'Qualified', ?, 'brown note', 1)
            """,
            (BROWN_ID, ids["multi_company_id"], master_rn, julie_id),
        )
        conn.execute(
            """
            INSERT INTO contacts (
                company_id, external_record_no, first_name, last_name, title,
                phone, alt_phone, email, source_row_index
            ) VALUES (?, ?, 'Pat', 'Export', 'Buyer', '555-0200', '555-0201',
                      'pat.export@example.test', 1)
            """,
            (ids["multi_company_id"], ccr_rn),
        )
        cur = conn.execute(
            """
            INSERT INTO companies (
                external_record_no, company_name, website, legacy_phone,
                address, city, state, zip
            ) VALUES (?, ?, 'empty.example.test', '555-0300',
                      '2 Empty Rd', 'Dayton', 'OH', '45402')
            """,
            (empty_rn, f"{MARKER} Empty {stamp}"),
        )
        ids["empty_company_id"] = int(cur.lastrowid)
        conn.execute(
            """
            INSERT INTO client_company_relationships (
                client_id, company_id, external_record_no, status, assigned_user_id
            ) VALUES (?, ?, ?, 'New', ?)
            """,
            (CARMECO_ID, ids["empty_company_id"], empty_rn, julie_id),
        )
        cur = conn.execute(
            """
            INSERT INTO companies (
                external_record_no, company_name, website, legacy_phone,
                address, city, state, zip
            ) VALUES (?, ?, 'solo.example.test', '555-0400',
                      '3 Solo St', 'Columbus', 'OH', '43215')
            """,
            (contact_rn, f"{MARKER} Solo {stamp}"),
        )
        ids["solo_company_id"] = int(cur.lastrowid)
        conn.execute(
            """
            INSERT INTO client_company_relationships (
                client_id, company_id, external_record_no, status, assigned_user_id
            ) VALUES (?, ?, ?, 'Left Message', ?)
            """,
            (BROWN_ID, ids["solo_company_id"], contact_rn, julie_id),
        )
        conn.execute(
            """
            INSERT INTO contacts (
                company_id, external_record_no, first_name, last_name, title,
                email, phone, source_row_index
            ) VALUES (?, ?, 'Alex', 'Solo', 'Engineer', 'alex.solo@example.test', '555-0500', 2)
            """,
            (ids["solo_company_id"], contact_rn),
        )
        cur = conn.execute(
            """
            INSERT INTO client_campaigns (
                client_id, campaign_name, description, is_active, is_default, status
            ) VALUES (?, ?, '', 1, 0, 'Active')
            """,
            (CARMECO_ID, campaign_name),
        )
        ids["campaign_id"] = int(cur.lastrowid)
        conn.execute(
            """
            INSERT INTO campaign_companies (
                campaign_id, client_id, company_id, notes, created_by
            ) VALUES (?, ?, ?, '', 'Master Export Test')
            """,
            (ids["campaign_id"], CARMECO_ID, ids["multi_company_id"]),
        )
        conn.commit()
    ids["master_rn"] = master_rn
    ids["ccr_rn"] = ccr_rn
    ids["empty_rn"] = empty_rn
    ids["contact_rn"] = contact_rn
    ids["campaign_name"] = campaign_name
    ids["stamp"] = stamp
    return ids


def _delete_fixtures(ids: dict[str, Any]) -> None:
    with get_connection() as conn:
        campaign_id = int(ids.get("campaign_id") or 0)
        if campaign_id:
            conn.execute("DELETE FROM campaign_companies WHERE campaign_id = ?", (campaign_id,))
            conn.execute("DELETE FROM campaign_contacts WHERE campaign_id = ?", (campaign_id,))
            conn.execute("DELETE FROM client_campaigns WHERE id = ?", (campaign_id,))
        for key in ("multi_company_id", "empty_company_id", "solo_company_id"):
            company_id = int(ids.get(key) or 0)
            if not company_id:
                continue
            conn.execute("DELETE FROM contacts WHERE company_id = ?", (company_id,))
            conn.execute(
                "DELETE FROM client_company_relationships WHERE company_id = ?",
                (company_id,),
            )
            conn.execute("DELETE FROM campaign_companies WHERE company_id = ?", (company_id,))
            conn.execute("DELETE FROM companies WHERE id = ?", (company_id,))
        conn.commit()


def _admin_session() -> tuple[TestClient, int]:
    password = _secret_password()
    email = f"mde.adm.{secrets.token_hex(4)}@example.test"
    user_id = _create_user(email=email, password=password, administrator=1)
    http = _client()
    login = _login(http, email, password)
    if login.status_code != 200:
        _delete_user(user_id)
        _fail(f"Administrator login failed: {login.status_code}")
    return http, user_id


def _csv_rows(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


def _find_company_row(rows: list[list[str]], company_id: int) -> list[str]:
    header = rows[0]
    id_idx = header.index("NorthStar Company ID")
    matches = [row for row in rows[1:] if str(row[id_idx]) == str(company_id)]
    if len(matches) != 1:
        _fail(f"Expected one Companies row for {company_id}, got {len(matches)}.")
    return matches[0]


def test_signed_out_export_is_401() -> None:
    os.environ.pop("NORTHSTAR_AUTH_ENFORCE", None)
    http = _client()
    resp = http.get(EXPORT, params={"mode": "companies", "format": "csv"})
    if resp.status_code != 401:
        _fail(f"Signed-out export must be 401, got {resp.status_code}.")
    if resp.json().get("detail") != AUTH_REQUIRED_DETAIL:
        _fail("Signed-out export 401 detail was not generic.")
    _assert_no_secrets(resp.json())


def test_non_admin_export_is_403() -> None:
    password = _secret_password()
    email = f"mde.staff.{secrets.token_hex(4)}@example.test"
    user_id = _create_user(email=email, password=password, administrator=0)
    http = _client()
    try:
        login = _login(http, email, password)
        if login.status_code != 200:
            _fail(f"Non-admin login failed: {login.status_code}")
        resp = http.get(EXPORT, params={"mode": "companies", "format": "csv"})
        if resp.status_code != 403:
            _fail(f"Non-admin export must be 403, got {resp.status_code}.")
        if resp.json().get("detail") != ADMIN_REQUIRED_DETAIL:
            _fail("Non-admin export 403 detail was not generic.")
        _assert_no_secrets(resp.json(), password)
    finally:
        _delete_user(user_id)


def test_admin_export_grain_and_xlsx_csv() -> None:
    http, user_id = _admin_session()
    ids = _insert_fixtures()
    try:
        before = snapshot_export_counts()
        company_count = before["companies"]
        relationship_count = before["client_company_relationships"]
        companies_without_contacts = 0
        with get_connection() as conn:
            companies_without_contacts = int(
                conn.execute(
                    """
                    SELECT COUNT(*) AS n FROM companies co
                    WHERE NOT EXISTS (
                        SELECT 1 FROM contacts ct WHERE ct.company_id = co.id
                    )
                    """
                ).fetchone()["n"]
            )
            contact_count = int(conn.execute("SELECT COUNT(*) AS n FROM contacts").fetchone()["n"])
        expected_contact_sheet_rows = contact_count + companies_without_contacts

        csv_companies = http.get(EXPORT, params={"mode": "companies", "format": "csv"})
        if csv_companies.status_code != 200:
            _fail(f"Admin companies CSV must be 200, got {csv_companies.status_code}.")
        if "text/csv" not in (csv_companies.headers.get("content-type") or ""):
            _fail("Companies CSV content-type must be CSV.")
        if csv_companies.content[:3] != b"\xef\xbb\xbf":
            _fail("CSV must be UTF-8-SIG.")
        company_rows = _csv_rows(csv_companies.content)
        if company_rows[0] != COMPANY_HEADERS:
            _fail(f"Companies CSV headers mismatch: {company_rows[0]}")
        if "Country" in company_rows[0]:
            _fail("Export must not invent a Country column.")
        data_rows = company_rows[1:]
        if len(data_rows) != company_count:
            _fail(
                f"Companies export must have one row per master company: "
                f"{len(data_rows)} != {company_count}."
            )
        id_idx = COMPANY_HEADERS.index("NorthStar Company ID")
        ids_in_export = [row[id_idx] for row in data_rows]
        if len(ids_in_export) != len(set(ids_in_export)):
            _fail("Companies export contains duplicate master company rows.")
        if ids_in_export.count(str(ids["multi_company_id"])) != 1:
            _fail("Multi-client company must appear once on Companies grain.")

        multi = _find_company_row(company_rows, int(ids["multi_company_id"]))
        rn_idx = COMPANY_HEADERS.index("LeadMaster Record Number")
        if multi[rn_idx] != ids["master_rn"]:
            _fail("LeadMaster Record Number must come from companies.external_record_no.")
        if multi[rn_idx] == ids["ccr_rn"]:
            _fail("LeadMaster Record Number must not use CCR record number.")
        linked = multi[COMPANY_HEADERS.index("Linked Clients")]
        if "Carmeco" not in linked or "Brown Industries" not in linked:
            _fail(f"Linked Clients must name both clients: {linked}")
        statuses = multi[COMPANY_HEADERS.index("Client Statuses")]
        if "Carmeco: Contacted" not in statuses or "Brown Industries: Qualified" not in statuses:
            _fail(f"Statuses must stay client-labeled: {statuses}")
        if statuses.strip() in {"Contacted", "Qualified"}:
            _fail("Statuses must not flatten into a fake global status.")
        reps = multi[COMPANY_HEADERS.index("Assigned Reps")]
        if "Carmeco:" not in reps or "Brown Industries:" not in reps:
            _fail(f"Assigned Reps must stay client-labeled: {reps}")
        campaigns = multi[COMPANY_HEADERS.index("Campaigns")]
        expected_campaign = f"Carmeco / {ids['campaign_name']}"
        if campaigns != expected_campaign:
            _fail(f"Campaigns must stay client-labeled, got {campaigns!r}.")

        csv_contacts = http.get(EXPORT, params={"mode": "companies-contacts", "format": "csv"})
        if csv_contacts.status_code != 200:
            _fail(f"Admin contacts CSV must be 200, got {csv_contacts.status_code}.")
        contact_rows = _csv_rows(csv_contacts.content)
        if contact_rows[0] != CONTACT_HEADERS:
            _fail(f"Contacts CSV headers mismatch: {contact_rows[0]}")
        if len(contact_rows) - 1 != expected_contact_sheet_rows:
            _fail(
                f"Contacts grain must keep zero-contact companies: "
                f"{len(contact_rows) - 1} != {expected_contact_sheet_rows}."
            )
        empty_matches = [
            row
            for row in contact_rows[1:]
            if row[CONTACT_HEADERS.index("NorthStar Company ID")] == str(ids["empty_company_id"])
        ]
        if len(empty_matches) != 1:
            _fail("Zero-contact company must still appear once on Contacts grain.")
        empty_row = empty_matches[0]
        if empty_row[CONTACT_HEADERS.index("NorthStar Contact ID")]:
            _fail("Zero-contact company row must have a blank NorthStar Contact ID.")
        if empty_row[CONTACT_HEADERS.index("LeadMaster Record Number")] != ids["empty_rn"]:
            _fail("Contacts grain LeadMaster Record Number must be companies.external_record_no.")

        multi_contact = [
            row
            for row in contact_rows[1:]
            if row[CONTACT_HEADERS.index("NorthStar Company ID")] == str(ids["multi_company_id"])
        ]
        if not multi_contact:
            _fail("Multi-client company contact row missing.")
        if multi_contact[0][CONTACT_HEADERS.index("LeadMaster Record Number")] != ids["master_rn"]:
            _fail("Contacts sheet must not treat contacts.external_record_no as LeadMaster ID.")
        if ids["ccr_rn"] == multi_contact[0][CONTACT_HEADERS.index("LeadMaster Record Number")]:
            _fail("Contact LeadMaster Record Number leaked the contact/CCR copy.")

        xlsx_companies = http.get(EXPORT, params={"mode": "companies", "format": "xlsx"})
        if xlsx_companies.status_code != 200:
            _fail(f"Admin companies XLSX must be 200, got {xlsx_companies.status_code}.")
        if "spreadsheetml" not in (xlsx_companies.headers.get("content-type") or ""):
            _fail("XLSX content-type must be spreadsheetml.")
        disposition = xlsx_companies.headers.get("content-disposition") or ""
        stamp = date.today().isoformat()
        if f"northstar-master-companies-{stamp}.xlsx" not in disposition:
            _fail(f"XLSX filename missing from Content-Disposition: {disposition}")
        wb = load_workbook(io.BytesIO(xlsx_companies.content), read_only=True, data_only=True)
        if wb.sheetnames != [SHEET_COMPANIES, SHEET_RELATIONSHIPS, SHEET_CAMPAIGNS]:
            _fail(f"Companies XLSX sheets mismatch: {wb.sheetnames}")
        companies_sheet = wb[SHEET_COMPANIES]
        company_xlsx_rows = list(companies_sheet.iter_rows(values_only=True))
        if list(company_xlsx_rows[0]) != COMPANY_HEADERS:
            _fail("Companies XLSX header mismatch.")
        if len(company_xlsx_rows) - 1 != company_count:
            _fail("Companies XLSX must have one row per master company.")
        rel_rows = list(wb[SHEET_RELATIONSHIPS].iter_rows(values_only=True))
        if list(rel_rows[0]) != RELATIONSHIP_HEADERS:
            _fail("Relationship sheet headers mismatch.")
        if len(rel_rows) - 1 != relationship_count:
            _fail("Relationship sheet must be one row per CCR.")
        multi_rels = [
            row
            for row in rel_rows[1:]
            if str(row[0]) == str(ids["multi_company_id"])
        ]
        if len(multi_rels) != 2:
            _fail("Relationship sheet should have two rows for the multi-client company.")
        camp_rows = list(wb[SHEET_CAMPAIGNS].iter_rows(values_only=True))
        if list(camp_rows[0]) != CAMPAIGN_HEADERS:
            _fail("Campaign sheet headers mismatch.")
        camp_hits = [
            row
            for row in camp_rows[1:]
            if str(row[0]) == str(ids["multi_company_id"])
        ]
        if not camp_hits:
            _fail("Campaign memberships sheet missing the fixture membership.")
        if str(camp_hits[0][3]) != "Carmeco" or str(camp_hits[0][4]) != ids["campaign_name"]:
            _fail(f"Campaign membership must stay client-labeled: {camp_hits[0]}")
        wb.close()

        xlsx_contacts = http.get(
            EXPORT, params={"mode": "companies-contacts", "format": "xlsx"}
        )
        if xlsx_contacts.status_code != 200:
            _fail(f"Admin contacts XLSX must be 200, got {xlsx_contacts.status_code}.")
        wb2 = load_workbook(io.BytesIO(xlsx_contacts.content), read_only=True, data_only=True)
        if SHEET_COMPANIES not in wb2.sheetnames or SHEET_CONTACTS not in wb2.sheetnames:
            _fail(f"Companies + Contacts XLSX missing sheets: {wb2.sheetnames}")
        if wb2.sheetnames != [
            SHEET_COMPANIES,
            SHEET_CONTACTS,
            SHEET_RELATIONSHIPS,
            SHEET_CAMPAIGNS,
        ]:
            _fail(f"Companies + Contacts sheet order mismatch: {wb2.sheetnames}")
        contact_xlsx_rows = list(wb2[SHEET_CONTACTS].iter_rows(values_only=True))
        if list(contact_xlsx_rows[0]) != CONTACT_HEADERS:
            _fail("Contacts XLSX header mismatch.")
        if len(contact_xlsx_rows) - 1 != expected_contact_sheet_rows:
            _fail("Contacts XLSX must retain companies with zero contacts.")
        wb2.close()

        after = snapshot_export_counts()
        if after != before:
            _fail(f"Export wrote to the database: {before} -> {after}")
    finally:
        _delete_fixtures(ids)
        _delete_user(user_id)


def test_invalid_params_are_400() -> None:
    http, user_id = _admin_session()
    try:
        bad_mode = http.get(EXPORT, params={"mode": "everything", "format": "csv"})
        if bad_mode.status_code != 400:
            _fail(f"Invalid mode must be 400, got {bad_mode.status_code}.")
        bad_format = http.get(EXPORT, params={"mode": "companies", "format": "pdf"})
        if bad_format.status_code != 400:
            _fail(f"Invalid format must be 400, got {bad_format.status_code}.")
    finally:
        _delete_user(user_id)


def main() -> None:
    tests = [
        test_signed_out_export_is_401,
        test_non_admin_export_is_403,
        test_admin_export_grain_and_xlsx_csv,
        test_invalid_params_are_400,
    ]
    failed = 0
    for test in tests:
        try:
            test()
            print(f"PASS: {test.__name__}")
        except Exception as exc:
            failed += 1
            print(f"FAIL: {test.__name__}: {exc}")
    if failed:
        sys.exit(1)
    print("PASS: master data export")


if __name__ == "__main__":
    main()
