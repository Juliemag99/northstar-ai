"""Phase 3A — READ-ONLY master-data standardization preview.

Does not write northstar.db. Does not rename, reformat, merge, or insert aliases.
"""

from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from contact_phone import format_us_phone_display, split_phone_extension
from import_brown_industries import digits_phone, norm_name
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parent.parent
DB = REPO / "database" / "northstar.db"
OUT_DIR = REPO / "working" / "master-data-std-phase3a"

SAFE = "SAFE AUTO-STANDARDIZE"
REVIEW = "REVIEW"
NO_CHANGE = "NO CHANGE"
CLASS_ORDER = {REVIEW: 0, SAFE: 1, NO_CHANGE: 2}

KEEP_WORDS = {
    "industries",
    "group",
    "works",
    "products",
    "manufacturing",
    "systems",
    "engineering",
    "international",
}

# Longest-first trailing legal suffixes only.
_SUFFIX_RE = re.compile(
    r"""(?ix)
    [\s,./-]+
    (
        incorporated | corporation | company | limited |
        l\.?\s*l\.?\s*c\.? | l\.?\s*l\.?\s*p\.? | l\.?\s*p\.? |
        plc | corp | ltd | llp | llc | inc | co | lp
    )
    \.?$
    """
)

_LOCATION_WORD_RE = re.compile(
    r"""(?ix)
    \b(
        plant | facility | facilities | division | div |
        headquarters | hq | operations | campus |
        off-?site | worldwide
    )\b
    """
)
_FALSE_PLANT_RE = re.compile(r"(?i)\bplant\s+food\b")
_SITE_SPLIT_RE = re.compile(r"\s*[-–—]\s*")
_PAREN_RE = re.compile(r"\(([^)]+)\)")
_SLASH_RE = re.compile(r"\s[/|]\s")
_DIVISION_OF_RE = re.compile(r"(?i)\ba\s+division\s+of\b")
_DBA_RE = re.compile(r"(?i)\b(d/?b/?a|aka|also known as|parent company)\b")

STREET_TYPES = {
    "street": "St",
    "st": "St",
    "road": "Rd",
    "rd": "Rd",
    "drive": "Dr",
    "dr": "Dr",
    "avenue": "Ave",
    "ave": "Ave",
    "boulevard": "Blvd",
    "blvd": "Blvd",
    "lane": "Ln",
    "ln": "Ln",
    "court": "Ct",
    "ct": "Ct",
    "circle": "Cir",
    "cir": "Cir",
    "parkway": "Pkwy",
    "pkwy": "Pkwy",
    "highway": "Hwy",
    "hwy": "Hwy",
    "place": "Pl",
    "pl": "Pl",
    "terrace": "Ter",
    "ter": "Ter",
    "trail": "Trl",
    "trl": "Trl",
}
UNIT_TYPES = {
    "suite": "Ste",
    "ste": "Ste",
    "building": "Bldg",
    "bldg": "Bldg",
    "floor": "Fl",
    "fl": "Fl",
    "apartment": "Apt",
    "apt": "Apt",
    "unit": "Unit",
}
DIR_FULL = {
    "north": "N",
    "south": "S",
    "east": "E",
    "west": "W",
    "northeast": "NE",
    "northwest": "NW",
    "southeast": "SE",
    "southwest": "SW",
}
DIR_ABBR = {"n", "s", "e", "w", "ne", "nw", "se", "sw"}
_PHONE_CHUNK_RE = re.compile(
    r"(?:\+?1[\s.\-]*)?(?:\(?\d{3}\)?[\s.\-]*)\d{3}[\s.\-]*\d{4}"
)
_EXPECTED = {
    "companies": 1792,
    "contacts": 4936,
    "ccr": 1889,
    "company_aliases": 174,
    "brown_sales_events": 97,
    "brown_quote_milestones": 26,
}


def _blank(value: object | None) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _snapshot(conn: sqlite3.Connection) -> dict:
    return {
        "integrity": conn.execute("PRAGMA integrity_check").fetchone()[0],
        "companies": conn.execute("SELECT COUNT(*) FROM companies").fetchone()[0],
        "contacts": conn.execute("SELECT COUNT(*) FROM contacts").fetchone()[0],
        "ccr": conn.execute("SELECT COUNT(*) FROM client_company_relationships").fetchone()[0],
        "brown_sales_events": conn.execute(
            "SELECT COUNT(*) FROM client_sales_events WHERE client_id=2"
        ).fetchone()[0],
        "brown_quote_milestones": conn.execute(
            """
            SELECT COUNT(*) FROM revenue_milestones
            WHERE client_id=2 AND milestone_type='Quote'
            """
        ).fetchone()[0],
        "company_aliases": conn.execute("SELECT COUNT(*) FROM company_aliases").fetchone()[0],
        "doc19": tuple(
            conn.execute(
                "SELECT id, processing_status, filename FROM client_documents WHERE id=19"
            ).fetchone()
            or ()
        ),
    }


def strip_trailing_legal_suffixes(name: str) -> str:
    text = re.sub(r"\s+", " ", _blank(name))
    if not text:
        return ""
    while True:
        match = _SUFFIX_RE.search(text)
        if not match:
            break
        token = re.sub(r"[^a-z]", "", match.group(1).lower())
        if token in KEEP_WORDS:
            break
        stripped = text[: match.start()].rstrip(" ,./-")
        if not stripped:
            break
        text = stripped
    return text


def _clean_punct_name(name: str) -> str:
    text = re.sub(r"\s+", " ", _blank(name))
    text = re.sub(r"\s+,", ",", text)
    text = re.sub(r",+", ",", text)
    text = text.strip(" ,")
    return text


_UNSPACED_SITE_RE = re.compile(
    r"""(?ix)
    ^(.+?)-
    (
        [^,-]+?
        \s+(?:plant|facility|division|div|hq|headquarters|site|campus|operations)
    )
    $
    """
)


def _looks_like_city(text: str, city: str) -> bool:
    right = re.sub(r"\s+", " ", _blank(text)).lower()
    city_n = re.sub(r"\s+", " ", _blank(city)).lower()
    return bool(city_n and (right == city_n or right.startswith(city_n + " ")))


def location_parts(name: str, city: str = "") -> tuple[str, str, list[str]]:
    """Return (base_candidate, site_candidate, reasons). Empty site if not location-like."""
    current = _clean_punct_name(name)
    reasons: list[str] = []
    base = current
    site = ""

    if _DIVISION_OF_RE.search(current):
        reasons.append("division-of wording")
        parts = _DIVISION_OF_RE.split(current, maxsplit=1)
        base = strip_trailing_legal_suffixes(_clean_punct_name(parts[0]))
        site = _clean_punct_name(current)
    paren = _PAREN_RE.search(current)
    if paren:
        inner = paren.group(1)
        if _LOCATION_WORD_RE.search(inner) or re.search(r"(?i)\bhq\b|headquarters", inner):
            reasons.append("HQ/location in parentheses")
            base = strip_trailing_legal_suffixes(_clean_punct_name(_PAREN_RE.sub("", current)))
            site = inner.strip()

    spaced = [p.strip() for p in re.split(r"\s+[-–—]\s+", current) if p.strip()]
    if len(spaced) >= 2:
        right = spaced[-1]
        left = " - ".join(spaced[:-1])
        if _LOCATION_WORD_RE.search(right) or _looks_like_city(right, city):
            reasons.append("dash-separated site/plant")
            base = strip_trailing_legal_suffixes(left)
            site = right
        else:
            reasons.append("possible division/DBA after dash")
            base = strip_trailing_legal_suffixes(left)
            site = right
    else:
        unspaced = _UNSPACED_SITE_RE.match(current)
        if unspaced:
            reasons.append("dash-separated site/plant")
            base = strip_trailing_legal_suffixes(unspaced.group(1))
            site = unspaced.group(2).strip()
        else:
            compact = current.rsplit("-", 1)
            if len(compact) == 2 and _looks_like_city(compact[1], city):
                reasons.append("dash-separated city/site")
                base = strip_trailing_legal_suffixes(compact[0])
                site = compact[1].strip()

    if _LOCATION_WORD_RE.search(current) and not _FALSE_PLANT_RE.search(current):
        if re.search(r"(?i)\b(plant|facility|division|headquarters|hq|campus)\b", current):
            if not any(
                key in r
                for r in reasons
                for key in ("site/plant", "parentheses", "division-of", "city/site")
            ):
                reasons.append("plant/facility/division/HQ wording")
                if not site:
                    site = current
                    base = strip_trailing_legal_suffixes(
                        _LOCATION_WORD_RE.sub(" ", current)
                    )
                    base = _clean_punct_name(base)

    if re.search(r"(?i)\bbranch\s+(office|location|plant|site)\b", current):
        reasons.append("branch site wording")

    reasons = list(dict.fromkeys(reasons))
    return _clean_punct_name(base), _clean_punct_name(site), reasons


def classify_company_name(name: str, city: str = "") -> dict:
    current = _clean_punct_name(name)
    proposed = strip_trailing_legal_suffixes(current)
    base, site, loc_reasons = location_parts(current, city=city)
    notes: list[str] = []

    if loc_reasons and (
        any(
            key in r
            for r in loc_reasons
            for key in ("plant", "facility", "division", "HQ", "parentheses", "site/plant", "branch site")
        )
        or "division-of" in " ".join(loc_reasons)
    ):
        return {
            "proposed": current,
            "classification": REVIEW,
            "reason": "REVIEW — LOCATION / SITE NAME. " + "; ".join(loc_reasons)
            + ". Do not auto-strip city/plant/HQ text.",
            "base": base or current,
            "site": site,
            "location_review": True,
            "notes": notes,
        }
    if loc_reasons and "possible division/DBA after dash" in loc_reasons:
        return {
            "proposed": current,
            "classification": REVIEW,
            "reason": "REVIEW — brand / DBA / division. Dash-separated name; do not auto-strip.",
            "base": base or current,
            "site": site,
            "location_review": True,
            "notes": notes,
        }
    if _SLASH_RE.search(current) or _DBA_RE.search(current) or _PAREN_RE.search(current):
        return {
            "proposed": current,
            "classification": REVIEW,
            "reason": "REVIEW — brand / DBA / acquired-name concern. Slash, parentheses, or parent-company wording.",
            "base": proposed,
            "site": "",
            "location_review": False,
            "notes": notes,
        }
    if proposed == current:
        return {
            "proposed": current,
            "classification": NO_CHANGE,
            "reason": "Already compliant (no trailing legal suffix to strip).",
            "base": "",
            "site": "",
            "location_review": False,
            "notes": notes,
        }
    if not proposed or proposed.lower() in {"inc", "llc", "corp", "co", "company"}:
        return {
            "proposed": current,
            "classification": REVIEW,
            "reason": "Unsafe suffix strip would leave an empty or suffix-only name.",
            "base": "",
            "site": "",
            "location_review": False,
            "notes": notes,
        }
    return {
        "proposed": proposed,
        "classification": SAFE,
        "reason": "Trailing legal suffix / suffix punctuation only. Interior brand words kept.",
        "base": "",
        "site": "",
        "location_review": False,
        "notes": notes,
    }


def _token_key(token: str) -> str:
    return re.sub(r"[^a-z0-9#]", "", token.lower())


def _is_street_type(token: str) -> bool:
    return _token_key(token) in STREET_TYPES


def _is_unit_type(token: str) -> bool:
    key = _token_key(token)
    return key in UNIT_TYPES or key in {"#"}


def _is_dir_full(token: str) -> bool:
    return _token_key(token) in DIR_FULL


def _is_dir_abbr(token: str) -> bool:
    return _token_key(token) in DIR_ABBR


def _is_directional(token: str) -> bool:
    return _is_dir_full(token) or _is_dir_abbr(token)


def _canon_dir(token: str, *, as_street_name: bool) -> str:
    key = _token_key(token)
    if as_street_name:
        if key in DIR_FULL:
            return key.title() if key not in {"northeast", "northwest", "southeast", "southwest"} else DIR_FULL[key]
        if key in DIR_ABBR:
            return key.upper()
        return token
    if key in DIR_FULL:
        return DIR_FULL[key]
    if key in DIR_ABBR:
        return key.upper()
    return token


def _canon_street_type(token: str) -> str:
    return STREET_TYPES.get(_token_key(token), token)


def _canon_unit(token: str) -> str:
    key = _token_key(token)
    if key == "#":
        return "#"
    return UNIT_TYPES.get(key, token)


def standardize_street_address(address: str) -> tuple[str, str, str]:
    current = _blank(address)
    if not current:
        return "", NO_CHANGE, "Blank address."
    if re.search(r"(?i)\b(p\.?\s*o\.?\s*box|post office box|rural route|\brr\b|\bhc\b)\b", current):
        return current, REVIEW, "PO Box / rural-route address; do not auto-standardize."
    if re.search(r"[;/]| and ", current, re.I) and re.search(r"\d", current):
        if current.count(",") > 3:
            return current, REVIEW, "Address looks like multiple lines jammed together."

    raw_tokens = re.split(r"\s+", current.replace(",", " , "))
    tokens = [t for t in raw_tokens if t and t != ","]
    if not tokens:
        return current, NO_CHANGE, "Blank after tokenize."

    out: list[str] = []
    i = 0
    seen_street_type = False
    grid_dir_count = 0
    while i < len(tokens):
        tok = tokens[i].strip(".,")
        if not tok:
            i += 1
            continue
        nxt = tokens[i + 1].strip(".,") if i + 1 < len(tokens) else ""
        if _is_unit_type(tok) or tok.startswith("#"):
            seen_street_type = True
            if tok.startswith("#") and len(tok) > 1:
                out.append("# " + tok[1:])
            else:
                out.append(_canon_unit(tok))
            i += 1
            continue
        if _is_street_type(tok):
            out.append(_canon_street_type(tok))
            seen_street_type = True
            i += 1
            continue
        if _is_directional(tok):
            grid_dir_count += 1
            next_is_type = bool(nxt) and _is_street_type(nxt)
            if seen_street_type:
                out.append(_canon_dir(tok, as_street_name=False))
            elif next_is_type:
                # Directional is the street name: "North St", "East Ave".
                out.append(_canon_dir(tok, as_street_name=True))
            else:
                out.append(_canon_dir(tok, as_street_name=False))
            i += 1
            continue
        out.append(tok)
        i += 1

    proposed = " ".join(out)
    proposed = re.sub(r"\s+", " ", proposed).strip()
    proposed = re.sub(r"\s+,", ",", proposed)
    proposed = re.sub(r",+", ",", proposed)
    proposed = re.sub(r"\s+#\s*", " # ", proposed)
    proposed = proposed.replace(" # #", " #").strip(" ,")

    if grid_dir_count >= 3:
        return (
            current,
            REVIEW,
            "Grid-style address with several directionals; confirm prefix vs street name before applying.",
        )
    if proposed == current:
        return current, NO_CHANGE, "Already matches the preferred street abbreviations and spacing."
    # Compare ignoring trailing periods on tokens only.
    compact_cur = re.sub(r"\.", "", current)
    compact_cur = re.sub(r"\s+", " ", compact_cur).replace(",", "").strip()
    compact_prop = re.sub(r"\.", "", proposed)
    compact_prop = re.sub(r"\s+", " ", compact_prop).replace(",", "").strip()
    if compact_cur.lower() == compact_prop.lower() and proposed == current:
        return current, NO_CHANGE, "Already compliant."
    return (
        proposed,
        SAFE,
        "Directional prefix vs street-name distinction preserved; street/unit types abbreviated; suite/unit kept.",
    )


def _digit_len(raw: str) -> int:
    digits = re.sub(r"\D", "", raw or "")
    if len(digits) >= 11 and digits.startswith("1"):
        digits = digits[1:]
    return len(digits)


def classify_phone(raw: str, *, field: str) -> dict:
    current = _blank(raw)
    if not current:
        return {
            "current": "",
            "main_digits": "",
            "extension": "",
            "proposed_main": "",
            "proposed_display": "",
            "classification": NO_CHANGE,
            "reason": "Blank.",
            "skip_counts": True,
        }
    main, ext = split_phone_extension(current)
    formatted = format_us_phone_display(current)
    chunks = _PHONE_CHUNK_RE.findall(current)
    all_digits = re.sub(r"\D", "", current)
    main_digits = digits_phone(main) if main else digits_phone(current)
    if main_digits and len(re.sub(r"\D", "", main)) >= 10:
        main_digits = re.sub(r"\D", "", main)
        if len(main_digits) >= 11 and main_digits.startswith("1"):
            main_digits = main_digits[1:]
        main_digits = main_digits[:10] if len(main_digits) >= 10 else main_digits
    else:
        tmp = all_digits[1:] if len(all_digits) >= 11 and all_digits.startswith("1") else all_digits
        main_digits = tmp[:10] if len(tmp) >= 10 else tmp

    proposed_main = ""
    if len(main_digits) == 10:
        proposed_main = f"({main_digits[:3]}) {main_digits[3:6]}-{main_digits[6:]}"

    ext_only = bool(
        re.fullmatch(r"(?i)\s*(ext\.?|extension|x|#)?\s*\d{2,6}\s*", current)
        and len(all_digits) <= 6
    )
    if ext_only and field.endswith("alt_phone"):
        return {
            "current": current,
            "main_digits": "",
            "extension": all_digits,
            "proposed_main": "",
            "proposed_display": current,
            "classification": REVIEW,
            "reason": "Alt field looks like an extension only. Do not format as a main number.",
            "skip_counts": False,
        }

    if len(chunks) >= 2:
        return {
            "current": current,
            "main_digits": main_digits,
            "extension": ext,
            "proposed_main": proposed_main,
            "proposed_display": current,
            "classification": REVIEW,
            "reason": "Multiple phone numbers in one field.",
            "skip_counts": False,
        }
    if re.match(r"^\s*\+" , current) and not re.match(r"^\s*\+1\b", current):
        return {
            "current": current,
            "main_digits": main_digits,
            "extension": ext,
            "proposed_main": proposed_main,
            "proposed_display": current,
            "classification": REVIEW,
            "reason": "International number (non-+1). Do not force US display.",
            "skip_counts": False,
        }
    leftover = all_digits[1:] if len(all_digits) >= 11 and all_digits.startswith("1") else all_digits
    lost_ext = (not ext) and len(leftover) > 10
    if lost_ext:
        return {
            "current": current,
            "main_digits": leftover[:10],
            "extension": leftover[10:],
            "proposed_main": f"({leftover[:10][:3]}) {leftover[:10][3:6]}-{leftover[:10][6:]}"
            if len(leftover) >= 10
            else "",
            "proposed_display": formatted,
            "classification": REVIEW,
            "reason": "Auto-format would keep 10 digits and drop leftover digits/extension (e.g. xt / trailing digits).",
            "skip_counts": False,
        }
    if len(main_digits) == 7 or (len(all_digits) == 7 and not ext):
        return {
            "current": current,
            "main_digits": all_digits,
            "extension": ext,
            "proposed_main": "",
            "proposed_display": current,
            "classification": REVIEW,
            "reason": "7-digit/local-only value. Do not assume an area code.",
            "skip_counts": False,
        }
    if len(main_digits) not in {0, 10} or (
        len(main_digits) != 10 and _digit_len(current) not in {10, 0}
    ):
        if len(main_digits) != 10:
            return {
                "current": current,
                "main_digits": main_digits or all_digits,
                "extension": ext,
                "proposed_main": proposed_main,
                "proposed_display": current,
                "classification": REVIEW,
                "reason": f"Unusual length ({len(all_digits)} digits) or ambiguous extension.",
                "skip_counts": False,
            }

    if proposed_main:
        proposed_display = proposed_main
        if ext:
            proposed_display = f"{proposed_main} x{ext}"
        if proposed_display == current or formatted == current and current.startswith("(") and " x" in current:
            if current == proposed_display:
                return {
                    "current": current,
                    "main_digits": main_digits,
                    "extension": ext,
                    "proposed_main": proposed_main,
                    "proposed_display": current,
                    "classification": NO_CHANGE,
                    "reason": "Already (xxx) xxx-xxxx with extension preserved.",
                    "skip_counts": False,
                }
        if current == proposed_main and not ext:
            return {
                "current": current,
                "main_digits": main_digits,
                "extension": "",
                "proposed_main": proposed_main,
                "proposed_display": current,
                "classification": NO_CHANGE,
                "reason": "Already (xxx) xxx-xxxx.",
                "skip_counts": False,
            }
        return {
            "current": current,
            "main_digits": main_digits,
            "extension": ext,
            "proposed_main": proposed_main,
            "proposed_display": proposed_display,
            "classification": SAFE,
            "reason": "Clear US/Canadian 10-digit main number"
            + ("; extension preserved." if ext else "."),
            "skip_counts": False,
        }

    return {
        "current": current,
        "main_digits": main_digits,
        "extension": ext,
        "proposed_main": proposed_main,
        "proposed_display": current,
        "classification": REVIEW,
        "reason": "Could not safely derive a 10-digit US/Canadian display value.",
        "skip_counts": False,
    }


def _self_check() -> None:
    assert strip_trailing_legal_suffixes("ABC Manufacturing, Inc.") == "ABC Manufacturing"
    assert strip_trailing_legal_suffixes("ABC Manufacturing LLC") == "ABC Manufacturing"
    assert strip_trailing_legal_suffixes("E J Ajax & Sons Co Inc") == "E J Ajax & Sons"
    assert strip_trailing_legal_suffixes("K Manufacturing Inc") == "K Manufacturing"
    assert standardize_street_address("1900 North St")[0] == "1900 North St"
    assert standardize_street_address("1900 North St")[1] == NO_CHANGE
    assert standardize_street_address("715 South St")[0] == "715 South St"
    assert standardize_street_address("93 East Ave")[0] == "93 East Ave"
    assert standardize_street_address("1525 E North St")[0] == "1525 E North St"
    proposed_walnut, wclass, _ = standardize_street_address("410 North Walnut Ave, Ste 105")
    assert proposed_walnut == "410 N Walnut Ave Ste 105", proposed_walnut
    assert wclass == SAFE
    assert format_us_phone_display("5155551212") == "(515) 555-1212"
    assert format_us_phone_display("515-555-1212") == "(515) 555-1212"
    assert format_us_phone_display("+1 515 555 1212") == "(515) 555-1212"


def _header_fill(ws, headers: list[str], fill: PatternFill, font: Font) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = font
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _write_rows(ws, rows: list[list], fills: dict[str, PatternFill], class_idx: int) -> None:
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append(row)
        klass = row[class_idx] if class_idx < len(row) else ""
        fill = fills.get(str(klass))
        for cell in ws[ws.max_row]:
            cell.alignment = wrap
            if fill:
                cell.fill = fill


def classify_live_dataset(conn: sqlite3.Connection) -> dict:
    """Regenerate Phase 3A classifications. Read-only; does not write."""
    aliases_by_company: dict[int, list[str]] = defaultdict(list)
    for row in conn.execute("SELECT company_id, alias_name FROM company_aliases"):
        aliases_by_company[int(row["company_id"])].append(_blank(row["alias_name"]))

    companies = list(
        conn.execute(
            """
            SELECT id, company_name, external_record_no, address, city, state, zip,
                   website, legacy_phone, legacy_alt_phone, legacy_mobile
            FROM companies
            ORDER BY id
            """
        )
    )
    contacts = list(
        conn.execute(
            """
            SELECT ct.id, ct.company_id, ct.first_name, ct.last_name, ct.phone, ct.alt_phone,
                   c.company_name
            FROM contacts ct
            JOIN companies c ON c.id = ct.company_id
            ORDER BY ct.id
            """
        )
    )

    name_rows = []
    loc_rows = []
    addr_rows = []
    company_phone_rows = []
    for row in companies:
        cid = int(row["id"])
        current = _blank(row["company_name"])
        classified = classify_company_name(current, city=_blank(row["city"]))
        proposed = classified["proposed"]
        name_rows.append(
            {
                "id": cid,
                "current": current,
                "proposed": proposed,
                "rn": _blank(row["external_record_no"]),
                "address": _blank(row["address"]),
                "city": _blank(row["city"]),
                "state": _blank(row["state"]),
                "zip": _blank(row["zip"]),
                "website": _blank(row["website"]),
                "phone": _blank(row["legacy_phone"]),
                "classification": classified["classification"],
                "reason": classified["reason"],
                "base": classified["base"],
                "site": classified["site"],
                "location_review": classified["location_review"],
                "aliases": aliases_by_company.get(cid, []),
            }
        )
        proposed_addr, addr_class, addr_reason = standardize_street_address(_blank(row["address"]))
        addr_rows.append(
            {
                "id": cid,
                "name": current,
                "current": _blank(row["address"]),
                "proposed": proposed_addr,
                "city": _blank(row["city"]),
                "state": _blank(row["state"]),
                "zip": _blank(row["zip"]),
                "classification": addr_class,
                "reason": addr_reason,
            }
        )
        for field, label in (
            ("legacy_phone", "legacy_phone"),
            ("legacy_alt_phone", "legacy_alt_phone"),
            ("legacy_mobile", "legacy_mobile"),
        ):
            phone_info = classify_phone(_blank(row[field]), field=label)
            if not phone_info["current"]:
                continue
            company_phone_rows.append(
                {
                    "id": cid,
                    "name": current,
                    "field": label,
                    **phone_info,
                }
            )

    collision_map: dict[str, list[dict]] = defaultdict(list)
    for item in name_rows:
        key_src = item["proposed"] if item["classification"] == SAFE else strip_trailing_legal_suffixes(item["current"])
        if item["classification"] == REVIEW and item["proposed"] == item["current"]:
            key_src = strip_trailing_legal_suffixes(item["current"]) or item["current"]
        key = re.sub(r"\s+", " ", key_src).strip().lower()
        if not key:
            continue
        collision_map[key].append(item)

    collision_groups = {k: v for k, v in collision_map.items() if len(v) >= 2}
    collision_ids = {item["id"] for group in collision_groups.values() for item in group}

    for item in name_rows:
        if item["classification"] == SAFE and item["id"] in collision_ids:
            others = [
                str(peer["id"])
                for peer in collision_map.get(item["proposed"].strip().lower(), [])
                if peer["id"] != item["id"]
            ]
            if not others:
                key = strip_trailing_legal_suffixes(item["current"]).strip().lower()
                others = [str(peer["id"]) for peer in collision_map.get(key, []) if peer["id"] != item["id"]]
            item["classification"] = REVIEW
            item["reason"] = (
                "REVIEW — POSSIBLE DUPLICATE. Proposed canonical name collides with another master. "
                "Do not auto-rename."
            )
            item["collision_ids"] = ", ".join(others)
        else:
            key = item["proposed"].strip().lower()
            others = [str(peer["id"]) for peer in collision_map.get(key, []) if peer["id"] != item["id"]]
            if not others:
                key2 = strip_trailing_legal_suffixes(item["current"]).strip().lower()
                others = [str(peer["id"]) for peer in collision_map.get(key2, []) if peer["id"] != item["id"]]
            item["collision_ids"] = ", ".join(others) if item["id"] in collision_ids else ""

        current_l = item["current"].strip().lower()
        preserved = any(_blank(alias).strip().lower() == current_l for alias in item["aliases"])
        item["alias_preserved"] = "Yes" if preserved else "No"
        item["alias_needed"] = (
            "Yes"
            if item["classification"] == SAFE and not preserved
            else ("No" if item["classification"] == SAFE else "")
        )
        if item["location_review"]:
            loc_rows.append(item)

    contact_phone_rows = []
    for row in contacts:
        display = " ".join(p for p in (_blank(row["first_name"]), _blank(row["last_name"])) if p) or "(unnamed)"
        for field in ("phone", "alt_phone"):
            info = classify_phone(_blank(row[field]), field=field)
            if not info["current"]:
                continue
            contact_phone_rows.append(
                {
                    "contact_id": int(row["id"]),
                    "company_id": int(row["company_id"]),
                    "company_name": _blank(row["company_name"]),
                    "name": display,
                    "field": field,
                    **info,
                }
            )

    def _count(rows, key="classification"):
        out = defaultdict(int)
        for row in rows:
            out[row[key]] += 1
        return out

    return {
        "name_rows": name_rows,
        "addr_rows": addr_rows,
        "company_phone_rows": company_phone_rows,
        "contact_phone_rows": contact_phone_rows,
        "loc_rows": loc_rows,
        "collision_groups": collision_groups,
        "name_counts": _count(name_rows),
        "addr_counts": _count(addr_rows),
        "cphone_counts": _count(company_phone_rows),
        "tphone_counts": _count(contact_phone_rows),
        "alias_needed_n": sum(1 for row in name_rows if row.get("alias_needed") == "Yes"),
    }


def main() -> int:
    _self_check()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    sha_before = _sha256(DB)
    conn = sqlite3.connect(f"file:{DB.as_posix()}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA query_only = ON")
    before = _snapshot(conn)
    classified = classify_live_dataset(conn)
    after_query = _snapshot(conn)
    conn.close()

    name_rows = classified["name_rows"]
    loc_rows = classified["loc_rows"]
    addr_rows = classified["addr_rows"]
    company_phone_rows = classified["company_phone_rows"]
    contact_phone_rows = classified["contact_phone_rows"]
    collision_groups = classified["collision_groups"]
    name_counts = classified["name_counts"]
    addr_counts = classified["addr_counts"]
    cphone_counts = classified["cphone_counts"]
    tphone_counts = classified["tphone_counts"]
    alias_needed_n = classified["alias_needed_n"]
    ext_company = sum(1 for row in company_phone_rows if row.get("extension"))
    ext_contact = sum(1 for row in contact_phone_rows if row.get("extension"))
    unusual_company = sum(1 for row in company_phone_rows if row["classification"] == REVIEW)
    unusual_contact = sum(1 for row in contact_phone_rows if row["classification"] == REVIEW)

    def _sort_rows(rows, name_key):
        return sorted(
            rows,
            key=lambda r: (
                CLASS_ORDER.get(r["classification"], 9),
                str(r.get(name_key) or "").lower(),
                r.get("id") or r.get("contact_id") or 0,
            ),
        )

    name_rows = _sort_rows(name_rows, "current")
    addr_rows = _sort_rows(addr_rows, "name")
    company_phone_rows = _sort_rows(company_phone_rows, "name")
    contact_phone_rows = _sort_rows(contact_phone_rows, "name")
    loc_rows = _sort_rows(loc_rows, "current")

    header_font = Font(bold=True, color="FFFFFF")
    fills = {
        REVIEW: PatternFill("solid", fgColor="C65911"),
        SAFE: PatternFill("solid", fgColor="548235"),
        NO_CHANGE: PatternFill("solid", fgColor="7F7F7F"),
        "header": PatternFill("solid", fgColor="1F4E79"),
    }
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary.append(["NorthStar Master Data Standardization — Phase 3A Preview"])
    summary.append(["Generated", datetime.now().replace(microsecond=0).isoformat(sep=" ")])
    summary.append(["Read-only. No live rows updated."])
    summary.append([])
    summary.append(["Metric", "SAFE AUTO-STANDARDIZE", "REVIEW", "NO CHANGE", "Total"])
    summary.append(
        [
            "Company names",
            name_counts[SAFE],
            name_counts[REVIEW],
            name_counts[NO_CHANGE],
            len(name_rows),
        ]
    )
    summary.append(
        [
            "Company addresses",
            addr_counts[SAFE],
            addr_counts[REVIEW],
            addr_counts[NO_CHANGE],
            len(addr_rows),
        ]
    )
    summary.append(
        [
            "Company phones (non-blank fields)",
            cphone_counts[SAFE],
            cphone_counts[REVIEW],
            cphone_counts[NO_CHANGE],
            len(company_phone_rows),
        ]
    )
    summary.append(
        [
            "Contact phones (non-blank fields)",
            tphone_counts[SAFE],
            tphone_counts[REVIEW],
            tphone_counts[NO_CHANGE],
            len(contact_phone_rows),
        ]
    )
    summary.append([])
    summary.append(["Canonical-name collision groups", len(collision_groups)])
    summary.append(["Location/site review companies", len(loc_rows)])
    summary.append(["SAFE renames needing alias preservation first", alias_needed_n])
    summary.append(["Company phones with detected extension", ext_company])
    summary.append(["Contact phones with detected extension", ext_contact])
    summary.append(["Company phone REVIEW (unusual/international/multi)", unusual_company])
    summary.append(["Contact phone REVIEW (unusual/international/multi)", unusual_contact])
    summary.append([])
    summary.append(["Live snapshot", json.dumps(before, default=str)])
    for col, width in enumerate([48, 24, 16, 16, 12], start=1):
        summary.column_dimensions[get_column_letter(col)].width = width

    names_ws = wb.create_sheet("Company Names")
    name_headers = [
        "Company ID",
        "Current Company Name",
        "Proposed Company Name",
        "Master RN",
        "Address",
        "City",
        "State",
        "Classification",
        "Reason",
        "Alias Already Preserved?",
        "Alias Needed Before Apply?",
        "Collision Company ID(s)",
        "Notes",
    ]
    _header_fill(names_ws, name_headers, fills["header"], header_font)
    _write_rows(
        names_ws,
        [
            [
                r["id"],
                r["current"],
                r["proposed"],
                r["rn"],
                r["address"],
                r["city"],
                r["state"],
                r["classification"],
                r["reason"],
                r["alias_preserved"],
                r["alias_needed"],
                r.get("collision_ids") or "",
                "; ".join(r["aliases"][:5]),
            ]
            for r in name_rows
        ],
        fills,
        7,
    )
    for i, width in enumerate([12, 40, 40, 14, 32, 16, 10, 22, 48, 18, 22, 18, 36], start=1):
        names_ws.column_dimensions[get_column_letter(i)].width = width

    addr_ws = wb.create_sheet("Company Addresses")
    addr_headers = [
        "Company ID",
        "Company Name",
        "Current Address",
        "Proposed Address",
        "City",
        "State",
        "ZIP",
        "Classification",
        "Reason",
    ]
    _header_fill(addr_ws, addr_headers, fills["header"], header_font)
    _write_rows(
        addr_ws,
        [
            [
                r["id"],
                r["name"],
                r["current"],
                r["proposed"],
                r["city"],
                r["state"],
                r["zip"],
                r["classification"],
                r["reason"],
            ]
            for r in addr_rows
        ],
        fills,
        7,
    )
    for i, width in enumerate([12, 36, 36, 36, 16, 10, 12, 22, 48], start=1):
        addr_ws.column_dimensions[get_column_letter(i)].width = width

    cph_ws = wb.create_sheet("Company Phones")
    phone_headers = [
        "Company ID",
        "Name",
        "Field",
        "Current Phone",
        "Proposed Main Phone",
        "Extension",
        "Main Phone Digits",
        "Classification",
        "Reason",
    ]
    _header_fill(cph_ws, phone_headers, fills["header"], header_font)
    _write_rows(
        cph_ws,
        [
            [
                r["id"],
                r["name"],
                r["field"],
                r["current"],
                r["proposed_main"],
                r["extension"],
                r["main_digits"],
                r["classification"],
                r["reason"],
            ]
            for r in company_phone_rows
        ],
        fills,
        7,
    )
    for i, width in enumerate([12, 36, 18, 28, 22, 14, 16, 22, 48], start=1):
        cph_ws.column_dimensions[get_column_letter(i)].width = width

    tph_ws = wb.create_sheet("Contact Phones")
    tph_headers = [
        "Contact ID",
        "Company ID",
        "Name",
        "Company Name",
        "Field",
        "Current Phone",
        "Proposed Main Phone",
        "Extension",
        "Classification",
        "Reason",
    ]
    _header_fill(tph_ws, tph_headers, fills["header"], header_font)
    _write_rows(
        tph_ws,
        [
            [
                r["contact_id"],
                r["company_id"],
                r["name"],
                r["company_name"],
                r["field"],
                r["current"],
                r["proposed_main"],
                r["extension"],
                r["classification"],
                r["reason"],
            ]
            for r in contact_phone_rows
        ],
        fills,
        8,
    )
    for i, width in enumerate([12, 12, 28, 32, 12, 28, 22, 14, 22, 48], start=1):
        tph_ws.column_dimensions[get_column_letter(i)].width = width

    col_ws = wb.create_sheet("Name Collisions")
    col_headers = [
        "Collision Key",
        "Company ID",
        "Current Company Name",
        "Proposed Company Name",
        "Address",
        "City",
        "State",
        "Phone",
        "Master RN",
        "Classification",
    ]
    _header_fill(col_ws, col_headers, fills["header"], header_font)
    collision_out = []
    for key, group in sorted(collision_groups.items(), key=lambda kv: kv[0]):
        for item in sorted(group, key=lambda r: r["id"]):
            collision_out.append(
                [
                    key,
                    item["id"],
                    item["current"],
                    item["proposed"],
                    item["address"],
                    item["city"],
                    item["state"],
                    item["phone"],
                    item["rn"],
                    item["classification"],
                ]
            )
    _write_rows(col_ws, collision_out, fills, 9)
    for i, width in enumerate([28, 12, 40, 40, 32, 16, 10, 18, 14, 22], start=1):
        col_ws.column_dimensions[get_column_letter(i)].width = width

    loc_ws = wb.create_sheet("Location-Site Review")
    loc_headers = [
        "Company ID",
        "Current Company Name",
        "Base Company Candidate",
        "Location/Site Candidate",
        "City",
        "State",
        "Address",
        "Master RN",
        "Alias/source evidence",
        "Reason",
    ]
    _header_fill(loc_ws, loc_headers, fills["header"], header_font)
    _write_rows(
        loc_ws,
        [
            [
                r["id"],
                r["current"],
                r["base"],
                r["site"],
                r["city"],
                r["state"],
                r["address"],
                r["rn"],
                "; ".join(r["aliases"][:8]),
                r["reason"],
            ]
            for r in loc_rows
        ],
        fills,
        9,
    )
    for i, width in enumerate([12, 44, 32, 32, 16, 10, 32, 14, 36, 48], start=1):
        loc_ws.column_dimensions[get_column_letter(i)].width = width

    xlsx_path = OUT_DIR / "NorthStar_Master_Data_Standardization_Preview.xlsx"
    json_path = OUT_DIR / "phase3a_preview_summary.json"
    wb.save(xlsx_path)

    examples = {
        "ajax": next((r for r in name_rows if r["id"] == 6), None),
        "spx": next((r for r in name_rows if r["id"] == 2), None),
        "heat_control": next((r for r in name_rows if r["id"] == 158), None),
        "jc_hq": next((r for r in name_rows if r["id"] == 757), None),
        "american_plant_food": next((r for r in name_rows if r["id"] == 1342), None),
        "north_st": next((r for r in addr_rows if r["id"] == 230), None),
        "south_st": next((r for r in addr_rows if r["id"] == 256), None),
        "east_ave": next((r for r in addr_rows if r["id"] == 554), None),
        "e_north_st": next((r for r in addr_rows if r["id"] == 172), None),
        "parsons": [r for r in name_rows if "parsons" in r["current"].lower()],
    }
    summary_payload = {
        "generated_at": datetime.now().replace(microsecond=0).isoformat(sep=" "),
        "company_names": {
            SAFE: name_counts[SAFE],
            REVIEW: name_counts[REVIEW],
            NO_CHANGE: name_counts[NO_CHANGE],
            "total": len(name_rows),
        },
        "addresses": {
            SAFE: addr_counts[SAFE],
            REVIEW: addr_counts[REVIEW],
            NO_CHANGE: addr_counts[NO_CHANGE],
            "total": len(addr_rows),
        },
        "company_phones": {
            SAFE: cphone_counts[SAFE],
            REVIEW: cphone_counts[REVIEW],
            NO_CHANGE: cphone_counts[NO_CHANGE],
            "total": len(company_phone_rows),
        },
        "contact_phones": {
            SAFE: tphone_counts[SAFE],
            REVIEW: tphone_counts[REVIEW],
            NO_CHANGE: tphone_counts[NO_CHANGE],
            "total": len(contact_phone_rows),
        },
        "collision_groups": len(collision_groups),
        "location_site_review": len(loc_rows),
        "safe_renames_needing_alias": alias_needed_n,
        "company_phones_with_extension": ext_company,
        "contact_phones_with_extension": ext_contact,
        "examples": {
            key: (
                {
                    "id": val["id"],
                    "current": val.get("current") or val.get("name"),
                    "proposed": val.get("proposed"),
                    "classification": val.get("classification"),
                    "reason": val.get("reason"),
                }
                if isinstance(val, dict)
                else [
                    {
                        "id": item["id"],
                        "current": item["current"],
                        "proposed": item["proposed"],
                        "classification": item["classification"],
                    }
                    for item in val
                ]
            )
            for key, val in examples.items()
            if val
        },
        "xlsx": str(xlsx_path),
        "db_before": before,
        "db_after_queries": after_query,
        "sha256_before": sha_before,
    }
    json_path.write_text(json.dumps(summary_payload, indent=2, default=str), encoding="utf-8")
    sha_after = _sha256(DB)
    report = {
        **summary_payload,
        "sha256_after": sha_after,
        "sha_unchanged": sha_before == sha_after,
        "snapshot_unchanged": before == after_query,
        "expected_ok": all(before.get(k) == v for k, v in _EXPECTED.items())
        and tuple(before.get("doc19") or ())[:2] == (19, "uploaded"),
    }
    print(json.dumps(report, indent=2, default=str))
    if sha_before != sha_after:
        raise SystemExit("Live DB SHA-256 changed; aborting preview as unsafe.")
    if before != after_query:
        raise SystemExit("Live DB counts changed during preview.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
