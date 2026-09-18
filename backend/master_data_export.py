"""Administrator Master Data Export of NorthStar companies and contacts.

Read-only. One Companies row per companies.id. Client statuses, reps, and
campaign memberships stay labeled by client — never flattened into a fake
global status. LeadMaster Record Number is companies.external_record_no.
"""

from __future__ import annotations

import csv
import io
from collections import defaultdict
from datetime import date
from typing import Any, Iterable

from db import get_connection

MODE_COMPANIES = "companies"
MODE_COMPANIES_CONTACTS = "companies-contacts"
FORMAT_XLSX = "xlsx"
FORMAT_CSV = "csv"

VALID_MODES = frozenset({MODE_COMPANIES, MODE_COMPANIES_CONTACTS})
VALID_FORMATS = frozenset({FORMAT_XLSX, FORMAT_CSV})

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA = "text/csv; charset=utf-8"

COMPANY_HEADERS = [
    "NorthStar Company ID",
    "LeadMaster Record Number",
    "Company Name",
    "Website",
    "Phone",
    "Phone Extension",
    "Alt Phone",
    "Alt Phone Extension",
    "Address",
    "City",
    "State",
    "ZIP",
    "Industry",
    "Source",
    "Created At",
    "Linked Clients",
    "Client Statuses",
    "Assigned Reps",
    "Campaigns",
]

RELATIONSHIP_HEADERS = [
    "NorthStar Company ID",
    "LeadMaster Record Number",
    "Company Name",
    "Client",
    "Client Code",
    "Status",
    "Assigned Rep",
    "CCR Record Number",
    "Hot",
    "Notes",
]

CAMPAIGN_HEADERS = [
    "NorthStar Company ID",
    "LeadMaster Record Number",
    "Company Name",
    "Client",
    "Campaign",
]

CONTACT_HEADERS = [
    "NorthStar Company ID",
    "LeadMaster Record Number",
    "Company Name",
    "Website",
    "Company Phone",
    "Company Phone Extension",
    "Address",
    "City",
    "State",
    "ZIP",
    "NorthStar Contact ID",
    "First Name",
    "Last Name",
    "Title",
    "Email",
    "Phone",
    "Phone Extension",
    "Alt Phone",
    "Alt Phone Extension",
    "Linked Clients",
    "Client Statuses",
    "Assigned Reps",
    "Campaigns",
]

SHEET_COMPANIES = "Companies"
SHEET_CONTACTS = "Contacts"
SHEET_RELATIONSHIPS = "Company-Client Relationships"
SHEET_CAMPAIGNS = "Campaign Memberships"

EXPORT_COUNT_TABLES = (
    "companies",
    "contacts",
    "client_company_relationships",
    "client_campaigns",
    "campaign_companies",
    "campaign_contacts",
    "users",
    "clients",
    "company_identity_keys",
    "activities",
)


def _blank(value: object | None) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _format_date(value: object | None) -> str:
    text = _blank(value)
    if not text:
        return ""
    if "T" in text:
        text = text.replace("T", " ", 1)
    if text.endswith("Z"):
        text = text[:-1]
    return text.strip()


def _yes_no(value: object | None) -> str:
    try:
        return "Yes" if int(value or 0) else "No"
    except (TypeError, ValueError):
        return "No"


def _join_clients(names: list[str]) -> str:
    seen: set[str] = set()
    out: list[str] = []
    for name in names:
        label = _blank(name)
        if not label or label in seen:
            continue
        seen.add(label)
        out.append(label)
    return "; ".join(out)


def _join_labeled(pairs: list[tuple[str, str]], *, sep: str = " | ") -> str:
    parts: list[str] = []
    seen: set[tuple[str, str]] = set()
    for client, value in pairs:
        key = (_blank(client), _blank(value))
        if not key[0] or key in seen:
            continue
        seen.add(key)
        parts.append(f"{key[0]}: {key[1]}" if key[1] else f"{key[0]}:")
    return sep.join(parts)


def _join_campaigns(pairs: list[tuple[str, str]]) -> str:
    parts: list[str] = []
    seen: set[tuple[str, str]] = set()
    for client, campaign in pairs:
        key = (_blank(client), _blank(campaign))
        if not key[0] or not key[1] or key in seen:
            continue
        seen.add(key)
        parts.append(f"{key[0]} / {key[1]}")
    return " | ".join(parts)


def _table_exists(conn, name: str) -> bool:
    row = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = ? LIMIT 1",
        (name,),
    ).fetchone()
    return row is not None


class _CompanyBundle:
    __slots__ = ("linked", "statuses", "reps", "campaigns")

    def __init__(self) -> None:
        self.linked: list[str] = []
        self.statuses: list[tuple[str, str]] = []
        self.reps: list[tuple[str, str]] = []
        self.campaigns: list[tuple[str, str]] = []


def _load_relationship_bundles(conn) -> dict[int, _CompanyBundle]:
    bundles: dict[int, _CompanyBundle] = defaultdict(_CompanyBundle)
    for row in conn.execute(
        """
        SELECT
            ccr.company_id,
            cl.name AS client_name,
            ccr.status,
            u.full_name AS assigned_rep
        FROM client_company_relationships ccr
        JOIN clients cl ON cl.id = ccr.client_id
        LEFT JOIN users u ON u.id = ccr.assigned_user_id
        ORDER BY cl.id, ccr.id
        """
    ).fetchall():
        company_id = int(row["company_id"])
        bundle = bundles[company_id]
        client_name = _blank(row["client_name"])
        bundle.linked.append(client_name)
        bundle.statuses.append((client_name, _blank(row["status"])))
        bundle.reps.append((client_name, _blank(row["assigned_rep"])))
    if _table_exists(conn, "campaign_companies") and _table_exists(conn, "client_campaigns"):
        for row in conn.execute(
            """
            SELECT
                cc.company_id,
                cl.name AS client_name,
                camp.campaign_name
            FROM campaign_companies cc
            JOIN client_campaigns camp ON camp.id = cc.campaign_id
            JOIN clients cl ON cl.id = cc.client_id
            ORDER BY cl.id, camp.campaign_name COLLATE NOCASE, cc.id
            """
        ).fetchall():
            bundles[int(row["company_id"])].campaigns.append(
                (_blank(row["client_name"]), _blank(row["campaign_name"]))
            )
    return bundles


def _labels_for(bundles: dict[int, _CompanyBundle], company_id: int) -> tuple[str, str, str, str]:
    bundle = bundles.get(int(company_id))
    if bundle is None:
        return "", "", "", ""
    return (
        _join_clients(bundle.linked),
        _join_labeled(bundle.statuses),
        _join_labeled(bundle.reps),
        _join_campaigns(bundle.campaigns),
    )


def _company_row(row, bundles: dict[int, _CompanyBundle]) -> list[Any]:
    company_id = int(row["id"])
    linked, statuses, reps, campaigns = _labels_for(bundles, company_id)
    return [
        company_id,
        _blank(row["external_record_no"]),
        _blank(row["company_name"]),
        _blank(row["website"]),
        _blank(row["legacy_phone"]),
        _blank(row["legacy_phone_extension"]),
        _blank(row["legacy_alt_phone"]),
        _blank(row["legacy_alt_phone_extension"]),
        _blank(row["address"]),
        _blank(row["city"]),
        _blank(row["state"]),
        _blank(row["zip"]),
        _blank(row["type_of_industry"]),
        _blank(row["source"]),
        _format_date(row["created_at"]),
        linked,
        statuses,
        reps,
        campaigns,
    ]


def _load_company_rows(conn, bundles: dict[int, _CompanyBundle]) -> list[list[Any]]:
    rows = conn.execute(
        """
        SELECT
            id, external_record_no, company_name, website, legacy_phone, legacy_phone_extension,
            legacy_alt_phone, legacy_alt_phone_extension,
            address, city, state, zip, type_of_industry, source, created_at
        FROM companies
        ORDER BY company_name COLLATE NOCASE, id
        """
    ).fetchall()
    return [_company_row(row, bundles) for row in rows]


def _load_relationship_rows(conn) -> list[list[Any]]:
    return [
        [
            int(row["company_id"]),
            _blank(row["master_record_no"]),
            _blank(row["company_name"]),
            _blank(row["client_name"]),
            _blank(row["client_code"]),
            _blank(row["status"]),
            _blank(row["assigned_rep"]),
            _blank(row["ccr_record_no"]),
            _yes_no(row["is_hot"]),
            _blank(row["notes"]),
        ]
        for row in conn.execute(
            """
            SELECT
                ccr.company_id,
                co.external_record_no AS master_record_no,
                co.company_name,
                cl.name AS client_name,
                cl.code AS client_code,
                ccr.status,
                u.full_name AS assigned_rep,
                ccr.external_record_no AS ccr_record_no,
                ccr.is_hot,
                ccr.notes
            FROM client_company_relationships ccr
            JOIN companies co ON co.id = ccr.company_id
            JOIN clients cl ON cl.id = ccr.client_id
            LEFT JOIN users u ON u.id = ccr.assigned_user_id
            ORDER BY co.company_name COLLATE NOCASE, co.id, cl.id, ccr.id
            """
        ).fetchall()
    ]


def _load_campaign_rows(conn) -> list[list[Any]]:
    if not _table_exists(conn, "campaign_companies") or not _table_exists(
        conn, "client_campaigns"
    ):
        return []
    return [
        [
            int(row["company_id"]),
            _blank(row["master_record_no"]),
            _blank(row["company_name"]),
            _blank(row["client_name"]),
            _blank(row["campaign_name"]),
        ]
        for row in conn.execute(
            """
            SELECT
                cc.company_id,
                co.external_record_no AS master_record_no,
                co.company_name,
                cl.name AS client_name,
                camp.campaign_name
            FROM campaign_companies cc
            JOIN companies co ON co.id = cc.company_id
            JOIN clients cl ON cl.id = cc.client_id
            JOIN client_campaigns camp ON camp.id = cc.campaign_id
            ORDER BY co.company_name COLLATE NOCASE, co.id, cl.id, camp.campaign_name COLLATE NOCASE
            """
        ).fetchall()
    ]


def _load_contact_rows(conn, bundles: dict[int, _CompanyBundle]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for row in conn.execute(
        """
        SELECT
            co.id AS company_id,
            co.external_record_no,
            co.company_name,
            co.website,
            co.legacy_phone,
            co.legacy_phone_extension,
            co.address,
            co.city,
            co.state,
            co.zip,
            ct.id AS contact_id,
            ct.first_name,
            ct.last_name,
            ct.title,
            ct.email,
            ct.phone,
            ct.phone_extension,
            ct.alt_phone,
            ct.alt_phone_extension
        FROM companies co
        LEFT JOIN contacts ct ON ct.company_id = co.id
        ORDER BY
            co.company_name COLLATE NOCASE,
            co.id,
            ct.last_name COLLATE NOCASE,
            ct.first_name COLLATE NOCASE,
            ct.id
        """
    ).fetchall():
        company_id = int(row["company_id"])
        linked, statuses, reps, campaigns = _labels_for(bundles, company_id)
        contact_id = row["contact_id"]
        rows.append(
            [
                company_id,
                _blank(row["external_record_no"]),
                _blank(row["company_name"]),
                _blank(row["website"]),
                _blank(row["legacy_phone"]),
                _blank(row["legacy_phone_extension"]),
                _blank(row["address"]),
                _blank(row["city"]),
                _blank(row["state"]),
                _blank(row["zip"]),
                int(contact_id) if contact_id is not None else "",
                _blank(row["first_name"]),
                _blank(row["last_name"]),
                _blank(row["title"]),
                _blank(row["email"]),
                _blank(row["phone"]),
                _blank(row["phone_extension"]),
                _blank(row["alt_phone"]),
                _blank(row["alt_phone_extension"]),
                linked,
                statuses,
                reps,
                campaigns,
            ]
        )
    return rows


def snapshot_export_counts(conn=None) -> dict[str, int]:
    """Read-only table counts used by tests to prove the export does not write."""
    owns = conn is None
    if owns:
        conn = get_connection()
    try:
        conn.execute("PRAGMA query_only = ON")
        out: dict[str, int] = {}
        for table in EXPORT_COUNT_TABLES:
            if _table_exists(conn, table):
                out[table] = int(conn.execute(f"SELECT COUNT(*) AS n FROM {table}").fetchone()["n"])
        return out
    finally:
        if owns:
            conn.close()


def _csv_bytes(headers: list[str], rows: Iterable[list[Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if cell is None else cell for cell in row])
    return buf.getvalue().encode("utf-8-sig")


def _style_sheet(ws, headers: list[str], data_row_count: int) -> None:
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font
    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(headers))
    last_row = max(1, data_row_count + 1)
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    for idx, header in enumerate(headers, start=1):
        width = min(42, max(14, len(header) + 4))
        ws.column_dimensions[get_column_letter(idx)].width = width


def _append_sheet(wb, title: str, headers: list[str], rows: list[list[Any]], *, active: bool):
    if active:
        ws = wb.active
        ws.title = title
    else:
        ws = wb.create_sheet(title=title)
    ws.append(headers)
    for row in rows:
        ws.append(["" if cell is None else cell for cell in row])
    _style_sheet(ws, headers, len(rows))
    return ws


def _xlsx_bytes(sheets: list[tuple[str, list[str], list[list[Any]]]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    first = True
    for title, headers, rows in sheets:
        _append_sheet(wb, title, headers, rows, active=first)
        first = False
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _normalize_mode(mode: str) -> str:
    key = _blank(mode).lower()
    if key not in VALID_MODES:
        raise ValueError("mode must be companies or companies-contacts.")
    return key


def _normalize_format(file_format: str) -> str:
    key = _blank(file_format).lower()
    if key not in VALID_FORMATS:
        raise ValueError("format must be xlsx or csv.")
    return key


def build_master_data_export(*, mode: str, file_format: str) -> tuple[str, bytes, str]:
    """Return (filename, content_bytes, media_type). Read-only."""
    mode_key = _normalize_mode(mode)
    format_key = _normalize_format(file_format)
    stamp = date.today().isoformat()
    with get_connection() as conn:
        conn.execute("PRAGMA query_only = ON")
        bundles = _load_relationship_bundles(conn)
        company_rows = _load_company_rows(conn, bundles)
        if format_key == FORMAT_CSV:
            if mode_key == MODE_COMPANIES:
                filename = f"northstar-master-companies-{stamp}.csv"
                return filename, _csv_bytes(COMPANY_HEADERS, company_rows), CSV_MEDIA
            contact_rows = _load_contact_rows(conn, bundles)
            filename = f"northstar-master-companies-contacts-{stamp}.csv"
            return filename, _csv_bytes(CONTACT_HEADERS, contact_rows), CSV_MEDIA

        relationship_rows = _load_relationship_rows(conn)
        campaign_rows = _load_campaign_rows(conn)
        sheets: list[tuple[str, list[str], list[list[Any]]]] = [
            (SHEET_COMPANIES, COMPANY_HEADERS, company_rows),
        ]
        if mode_key == MODE_COMPANIES_CONTACTS:
            sheets.append((SHEET_CONTACTS, CONTACT_HEADERS, _load_contact_rows(conn, bundles)))
        sheets.append((SHEET_RELATIONSHIPS, RELATIONSHIP_HEADERS, relationship_rows))
        sheets.append((SHEET_CAMPAIGNS, CAMPAIGN_HEADERS, campaign_rows))
        suffix = "companies-contacts" if mode_key == MODE_COMPANIES_CONTACTS else "companies"
        filename = f"northstar-master-{suffix}-{stamp}.xlsx"
        return filename, _xlsx_bytes(sheets), XLSX_MEDIA
