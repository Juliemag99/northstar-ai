"""Isolated tests for consolidated shared note-history import.

Run: python test_shared_note_history_import.py

Uses isolated testdb copies only — never writes production northstar.db,
never restarts services, never imports live Dawson into production.
"""

from __future__ import annotations

import csv
import hashlib
import io
import os
import secrets
from pathlib import Path

import testdb
from auth_http import CSRF_HEADER
from auth_passwords import hash_password
from db import DB_PATH, PRODUCTION_DB_PATH, get_connection, migrate_schema
from fastapi.testclient import TestClient
from main import app
from models import NorthStarUser
from shared_history_data import list_shared_history
from shared_note_history_import import (
    CLEANUP_COMPLETED,
    CLEANUP_PENDING,
    UNATTRIBUTED_LABEL,
    build_preview_counts,
    confirm_shared_note_history_import,
    ensure_shared_note_history_schema,
    normalize_record_no,
    parse_history_csv,
    parse_prospects_csv,
    plan_contact_actions,
    retry_shared_note_history_staging_cleanup,
    upload_shared_note_history_import,
    _company_by_record_no,
)

UPLOAD = "/api/clients/{client_id}/admin/shared-note-history-imports"
CONFIRM = "/api/clients/{client_id}/admin/shared-note-history-imports/{batch_id}/confirm"

DAWSON_DIR = Path(r"C:\Users\julie\Downloads")
DAWSON_PROSPECTS = DAWSON_DIR / "Dawson_NorthStar_Prospects_No_Closed.csv"
DAWSON_HISTORY = DAWSON_DIR / "Dawson_Shared_Note_History_No_Closed.csv"
DAWSON_VALIDATION_XLSX = DAWSON_DIR / "Dawson_Event_History_Validation.xlsx"
DAWSON_PREIMPORT_BACKUP = (
    PRODUCTION_DB_PATH.parent
    / "northstar.db.backup-shared-history-activate-20260907-152255"
)

FREEZE_TABLES = (
    "activities",
    "legacy_notes",
    "appointments",
    "work_queue_items",
    "company_research_jobs",
    "field_change_audit",
)


def _fail(msg: str) -> None:
    raise AssertionError(msg)


def _prove_isolated() -> None:
    opened = Path(os.fspath(DB_PATH)).resolve()
    prod = PRODUCTION_DB_PATH.resolve()
    if opened == prod:
        _fail("Testdb isolation failed: opened production DB.")
    env = Path(os.environ.get("NORTHSTAR_TEST_DB", "")).resolve()
    if env != opened:
        _fail("NORTHSTAR_TEST_DB does not match the opened database.")
    print(f"ISOLATED {opened}")


def _client() -> TestClient:
    return TestClient(app)


def _secret_password() -> str:
    return f"NsTest9{secrets.token_hex(10)}"


def _csrf(response) -> str:
    return str(response.json().get("csrf_token") or "")


def _create_user(*, email: str, password: str, administrator: int = 1) -> int:
    digest = hash_password(password, email=email)
    with get_connection() as conn:
        conn.execute(
            """
            INSERT INTO users (
                email, full_name, is_administrator, is_internal_northstar, active,
                password_hash, failed_login_count, locked_until
            ) VALUES (?, ?, ?, 1, 1, ?, 0, '')
            """,
            (email, "SNH Import Test User", int(administrator), digest),
        )
        user_id = int(
            conn.execute("SELECT id FROM users WHERE email = ?", (email,)).fetchone()["id"]
        )
        conn.commit()
        return user_id


def _assign(user_id: int, client_id: int) -> None:
    with get_connection() as conn:
        conn.execute(
            """
            INSERT OR REPLACE INTO user_client_assignments
                (user_id, client_id, role, active, assigned_at)
            VALUES (?, ?, 'staff', 1, datetime('now'))
            """,
            (user_id, client_id),
        )
        conn.commit()


def _delete_user(user_id: int) -> None:
    with get_connection() as conn:
        conn.execute("DELETE FROM staff_sessions WHERE user_id = ?", (user_id,))
        conn.execute("DELETE FROM user_client_assignments WHERE user_id = ?", (user_id,))
        conn.execute("DELETE FROM users WHERE id = ?", (user_id,))
        conn.commit()


def _client_ids() -> tuple[int, int]:
    with get_connection() as conn:
        rows = conn.execute("SELECT id FROM clients ORDER BY id LIMIT 2").fetchall()
        if len(rows) < 2:
            if len(rows) < 1:
                conn.execute(
                    "INSERT INTO clients (code, name) VALUES ('snh_a', 'SNH Client A')"
                )
            conn.execute(
                "INSERT INTO clients (code, name) VALUES (?, ?)",
                (f"snh_b_{secrets.token_hex(3)}", "SNH Client B"),
            )
            conn.commit()
            rows = conn.execute("SELECT id FROM clients ORDER BY id LIMIT 2").fetchall()
    return int(rows[0]["id"]), int(rows[1]["id"])


def _actor(user_id: int) -> NorthStarUser:
    with get_connection() as conn:
        row = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    return NorthStarUser(
        id=int(row["id"]),
        email=str(row["email"] or ""),
        full_name=str(row["full_name"] or ""),
        is_administrator=bool(row["is_administrator"]),
        is_internal_northstar=bool(row["is_internal_northstar"])
        if "is_internal_northstar" in row.keys()
        else True,
        active=bool(row["active"]),
        created_at=str(row["created_at"] or ""),
    )


def _csv_bytes(headers: list[str], rows: list[list[object]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


def _long_note(n: int = 2500) -> str:
    return ("LONG-NOTE-" + ("x" * 80) + "\n") * ((n // 90) + 1)


def _fixture_csvs() -> tuple[bytes, bytes]:
    """Fixtures using the real Dawson prepared-file header names."""
    long_body = _long_note(2500)
    prospects = _csv_bytes(
        [
            "FirstName",
            "LastName",
            "Title",
            "Phone",
            "Alt Phone",
            "Email",
            "Record No.",
            "Company",
            "Address1",
            "City",
            "State",
            "Zip",
            "Web Address",
            "Customer Campaign",
            "Mobile",
            "Status",
        ],
        [
            [
                "Ann",
                "Alpha",
                "Buyer",
                "5551112222",
                "",
                "ann@alpha.test",
                "1001",
                "Alpha Fab",
                "1 Main",
                "Peoria",
                "IL",
                "61602",
                "",
                "Dawson CAPS[]Brown Cap",
                "",
                "New",
            ],
            [
                "Bob",
                "Alpha",
                "Eng",
                "5551113333",
                "",
                "bob@alpha.test",
                "1001",
                "Alpha Fab",
                "1 Main",
                "Peoria",
                "IL",
                "61602",
                "",
                "Dawson CAPS[]Brown Cap",
                "",
                "New",
            ],
            [
                "Bea",
                "Beta",
                "",
                "",
                "",
                "bea@beta.test",
                "1002.0",
                "Beta Works",
                "",
                "",
                "",
                "",
                "",
                "Carm CAPS",
                "",
                "Contacted",
            ],
            [
                "Carl",
                "Closed",
                "",
                "",
                "",
                "carl@closed.test",
                "1003",
                "Closed Co",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "Closed",
            ],
            [
                "Cara",
                "Closed",
                "",
                "",
                "",
                "cara@closed.test",
                "1003",
                "Closed Co",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "Closed",
            ],
        ],
    )
    history = _csv_bytes(
        [
            "LeadMaster Record No.",
            "Company",
            "Current Dawson Status",
            "History Order",
            "Event Sequence",
            "Event Timestamp",
            "Author",
            "Event Type",
            "Attributed Client/Campaign",
            "Attribution Evidence",
            "Note Text",
            "Source File",
            "Event Hash",
        ],
        [
            [
                "1001",
                "Alpha Fab",
                "New",
                "Newest to oldest",
                "1",
                "2-Sep-2026 5:00 PM EDT",
                "Tyler Sullivan",
                "Left Message",
                UNATTRIBUTED_LABEL,
                "",
                "Left Message - Ann",
                "Notes dawson.csv",
                "hash-alpha-1",
            ],
            [
                "1001",
                "Alpha Fab",
                "New",
                "Newest to oldest",
                "2",
                "1-Sep-2026 4:00 PM EDT",
                "Ian Andrews",
                "Note",
                "Brown Industries",
                "Brown Cap campaign tag",
                "Brown - follow-up note.",
                "Notes dawson.csv",
                "hash-alpha-2",
            ],
            [
                "1001",
                "Alpha Fab",
                "New",
                "Newest to oldest",
                "3",
                "31-Aug-2026 3:00 PM CDT",
                "Haley Knudsen",
                "Call",
                "Dawson Fabrication",
                "Dawson -",
                "Dawson - earlier call.",
                "Notes dawson.csv",
                "hash-alpha-3",
            ],
            [
                "1002",
                "Beta Works",
                "Contacted",
                "Newest to oldest",
                "1",
                "3-Jan-2026 12:00 PM CST",
                "Haley Knudsen",
                "Note",
                "Carmeco",
                "Other-client campaign tag",
                long_body,
                "Notes dawson.csv",
                "hash-beta-long",
            ],
            [
                "1003",
                "Closed Co",
                "Closed",
                "Newest to oldest",
                "1",
                "4-Jan-2026 13:00:00",
                "Someone",
                "Note",
                "Should exclude",
                "Closed company",
                "Closed history must not import.",
                "Notes dawson.csv",
                "hash-closed-1",
            ],
            [
                "9999",
                "Orphan Co",
                "New",
                "Newest to oldest",
                "1",
                "5-Jan-2026 14:00:00",
                "Orphan",
                "Note",
                "Orphan",
                "Not in prospects",
                "Orphan history skipped.",
                "Notes dawson.csv",
                "hash-orphan",
            ],
        ],
    )
    return prospects, history


def _table_fingerprint(conn, table: str) -> str:
    if not conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",
        (table,),
    ).fetchone():
        return f"{table}:missing"
    rows = conn.execute(f"SELECT * FROM {table} ORDER BY rowid").fetchall()
    cols = [d[0] for d in conn.execute(f"SELECT * FROM {table} LIMIT 0").description]
    payload = [
        tuple("" if row[c] is None else str(row[c]) for c in cols) for row in rows
    ]
    digest = hashlib.sha256(repr(payload).encode("utf-8")).hexdigest()
    return f"{table}:{len(rows)}:{digest}"


def _freeze_snapshot() -> dict[str, str]:
    with get_connection() as conn:
        return {
            table: _table_fingerprint(conn, table)
            for table in FREEZE_TABLES
            if conn.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",
                (table,),
            ).fetchone()
        }


def test_normalize_record_no() -> None:
    assert normalize_record_no(" 1002.0 ") == "1002"
    assert normalize_record_no("4636") == "4636"
    assert normalize_record_no("01001") == "01001"
    # Exact string match after normalize — never by company name.
    assert normalize_record_no("4636") != normalize_record_no("46360")


def test_closed_exclusion_and_exact_record_match() -> None:
    prospects_bytes, history_bytes = _fixture_csvs()
    prospects = parse_prospects_csv(prospects_bytes)
    assert prospects.excluded_closed_companies == 1
    assert "1003" not in prospects.included_company_record_nos
    assert set(prospects.included_company_record_nos) == {"1001", "1002"}
    assert len(prospects.contacts) == 3

    history = parse_history_csv(
        history_bytes,
        allowed_record_nos=set(prospects.included_company_record_nos),
        closed_record_nos={"1003"},
    )
    assert history.excluded_closed_events >= 1
    assert len(history.events) == 4
    assert history.long_note_count == 1
    assert history.unattributed_count == 1
    assert {e.record_no for e in history.events} == {"1001", "1002"}
    # Attribution column mapped from Attributed Client/Campaign
    attrs = {e.attribution for e in history.events}
    assert "Brown Industries" in attrs
    assert "Dawson Fabrication" in attrs
    assert UNATTRIBUTED_LABEL in attrs


def test_dawson_prepared_files_read_only_counts() -> None:
    """Validate prepared Downloads files without writing any database."""
    if not DAWSON_PROSPECTS.is_file() or not DAWSON_HISTORY.is_file():
        print("SKIP dawson prepared file counts (files not present)")
        return
    prospects = parse_prospects_csv(DAWSON_PROSPECTS.read_bytes())
    assert not prospects.errors, prospects.errors
    assert prospects.excluded_closed_companies == 0
    assert len(prospects.included_company_record_nos) == 243
    assert len(prospects.contacts) == 997

    history = parse_history_csv(
        DAWSON_HISTORY.read_bytes(),
        allowed_record_nos=set(prospects.included_company_record_nos),
    )
    assert not history.errors, history.errors
    assert len(history.events) == 13319
    assert history.long_note_count == 11
    assert history.excluded_closed_events == 0
    hashes = [e.event_hash for e in history.events]
    assert len(hashes) == len(set(hashes))
    assert all(h for h in hashes)
    # Empty source-note rows (Apollo / Eaton-Boyd) correctly have no history events.
    with_hist = {e.record_no for e in history.events}
    without = set(prospects.included_company_record_nos) - with_hist
    expected_missing = len(prospects.included_company_record_nos) - len(with_hist)
    assert len(without) == expected_missing == 11
    assert {"1053898", "938403"}.issubset(without)

    if DAWSON_VALIDATION_XLSX.is_file():
        summary_without, missing_sheet_rns = _read_validation_missing_counts(
            DAWSON_VALIDATION_XLSX
        )
        assert summary_without == expected_missing
        assert len(missing_sheet_rns) == expected_missing
        assert set(missing_sheet_rns) == without
    print("PASS dawson prepared file read-only counts + workbook reconciliation")


def _read_validation_missing_counts(path: Path) -> tuple[int, list[str]]:
    """Read Summary 'records without note history' and Missing Note History RNs."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    summary = wb["Summary"]
    without_count = None
    for row in summary.iter_rows(min_row=4, max_col=2, values_only=True):
        label = str(row[0] or "").strip().casefold()
        if label == "records without note history":
            without_count = int(row[1])
            break
    if without_count is None:
        _fail("Validation workbook Summary missing 'records without note history'.")
    miss = wb["Missing Note History"]
    rns: list[str] = []
    for row in miss.iter_rows(min_row=2, max_col=1, values_only=True):
        rn = normalize_record_no(row[0])
        if rn:
            rns.append(rn)
    return without_count, rns

def test_preview_confirm_visibility_order_idempotency_freeze() -> None:
    _prove_isolated()
    with get_connection() as conn:
        ensure_shared_note_history_schema(conn)
        migrate_schema(conn)
        conn.commit()

    client_a, client_b = _client_ids()
    password = _secret_password()
    email = f"snh.admin.{secrets.token_hex(6)}@example.test"
    user_id = _create_user(email=email, password=password, administrator=1)
    _assign(user_id, client_a)
    _assign(user_id, client_b)
    actor = _actor(user_id)
    prospects_bytes, history_bytes = _fixture_csvs()

    with get_connection() as conn:
        ensure_shared_note_history_schema(conn)
        # Seed existing Alpha company linked to both clients; preserve Hot Prospect.
        conn.execute(
            """
            INSERT INTO companies (
                external_record_no, company_name, created_at, last_updated_at
            ) VALUES ('1001', 'Alpha Fab Existing', datetime('now'), datetime('now'))
            """
        )
        company_id = int(
            conn.execute(
                "SELECT id FROM companies WHERE external_record_no = '1001'"
            ).fetchone()["id"]
        )
        conn.execute(
            """
            INSERT INTO client_company_relationships (
                client_id, company_id, status, assigned_user_id, notes,
                is_hot, created_at, updated_at, external_record_no
            ) VALUES (?, ?, 'Hot Prospect', ?, 'keep-me', 1, datetime('now'), datetime('now'), '1001')
            """,
            (client_a, company_id, user_id),
        )
        rel_a = int(
            conn.execute(
                """
                SELECT id FROM client_company_relationships
                WHERE client_id = ? AND company_id = ?
                """,
                (client_a, company_id),
            ).fetchone()["id"]
        )
        conn.execute(
            """
            INSERT INTO client_company_relationships (
                client_id, company_id, status, assigned_user_id, notes,
                is_hot, created_at, updated_at, external_record_no
            ) VALUES (?, ?, 'Send Information', ?, '', 0, datetime('now'), datetime('now'), '1001')
            """,
            (client_b, company_id, user_id),
        )
        # Existing CRM artifacts that must freeze.
        if conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='activities'"
        ).fetchone():
            conn.execute(
                """
                INSERT INTO activities (
                    client_id, company_id, relationship_id, external_record_no,
                    activity_type, notes, activity_at, created_at, created_by
                ) VALUES (?, ?, ?, '1001', 'Note', 'pre-existing activity',
                          datetime('now'), datetime('now'), 'tester')
                """,
                (client_a, company_id, rel_a),
            )
        if conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='legacy_notes'"
        ).fetchone():
            conn.execute(
                """
                INSERT INTO legacy_notes (client_id, company_id, note_text, source_field)
                VALUES (?, ?, 'pre-existing legacy', 'Sales Rep Comments/Notes')
                """,
                (client_a, company_id),
            )
        conn.commit()

    before_freeze = _freeze_snapshot()

    upload = upload_shared_note_history_import(
        client_id=client_a,
        actor=actor,
        prospects_filename="Dawson_NorthStar_Prospects_No_Closed.csv",
        prospects_content=prospects_bytes,
        history_filename="Dawson_Shared_Note_History_No_Closed.csv",
        history_content=history_bytes,
    )
    assert upload.preview.ok, upload.preview.errors
    counts = upload.preview.counts
    assert counts.companies_included == 2
    assert counts.companies_excluded_closed == 1
    assert counts.history_events_total == 4
    assert counts.history_long_notes == 1
    assert counts.history_unattributed == 1
    assert counts.status_preserved >= 1

    result = confirm_shared_note_history_import(
        client_id=client_a, batch_id=upload.batch_id, actor=actor
    )
    assert result.ok
    assert result.history_inserted == 4
    assert result.status_preserved >= 1

    after_freeze = _freeze_snapshot()
    assert after_freeze == before_freeze, "Import must not mutate existing CRM history tables"

    with get_connection() as conn:
        status = conn.execute(
            """
            SELECT status FROM client_company_relationships
            WHERE client_id = ? AND external_record_no = '1001'
            """,
            (client_a,),
        ).fetchone()["status"]
        assert status == "Hot Prospect"

        # Company-level history once (not per contact).
        hist_alpha = list(
            conn.execute(
                """
                SELECT * FROM company_shared_history_events
                WHERE external_record_no = '1001'
                ORDER BY
                    CASE
                        WHEN TRIM(COALESCE(event_sequence, '')) GLOB '[0-9]*'
                         AND TRIM(COALESCE(event_sequence, '')) != ''
                        THEN CAST(TRIM(event_sequence) AS INTEGER)
                        ELSE 999999999
                    END ASC,
                    id DESC
                """
            )
        )
        assert len(hist_alpha) == 3
        contacts = int(
            conn.execute(
                "SELECT COUNT(*) AS c FROM contacts WHERE external_record_no = '1001'"
            ).fetchone()["c"]
        )
        assert contacts == 2
        # Chronological newest-first via Event Sequence (1 = newest)
        assert [r["event_sequence"] for r in hist_alpha] == ["1", "2", "3"]
        assert hist_alpha[0]["event_hash"] == "hash-alpha-1"
        assert hist_alpha[2]["event_hash"] == "hash-alpha-3"
        # Attribution preserved
        attrs = {r["attribution"] for r in hist_alpha}
        assert "Brown Industries" in attrs
        assert "Dawson Fabrication" in attrs
        assert UNATTRIBUTED_LABEL in attrs
        # Long note intact
        long_row = conn.execute(
            """
            SELECT note_text FROM company_shared_history_events
            WHERE event_hash = 'hash-beta-long'
            """
        ).fetchone()
        assert long_row is not None
        assert len(long_row["note_text"]) > 2000
        assert int(
            conn.execute(
                "SELECT COUNT(*) AS c FROM company_shared_history_events WHERE event_hash='hash-closed-1'"
            ).fetchone()["c"]
        ) == 0
        # No duplicate hashes in DB
        dup = conn.execute(
            """
            SELECT event_hash, COUNT(*) AS c
            FROM company_shared_history_events
            GROUP BY company_id, event_hash
            HAVING c > 1
            """
        ).fetchall()
        assert dup == []

    # Cross-client visibility + attribution + unattributed
    shared = list_shared_history("1001", user_id=user_id, working_client_id=client_b)
    assert shared.company_id > 0
    assert len(shared.shared_company_history_items) == 3
    # Newest first via Event Sequence
    assert shared.shared_company_history_items[0].event_hash == "hash-alpha-1"
    assert shared.shared_company_history_items[2].event_hash == "hash-alpha-3"
    shared_attrs = {i.attribution for i in shared.shared_company_history_items}
    assert UNATTRIBUTED_LABEL in shared_attrs
    assert "Brown Industries" in shared_attrs

    # Idempotent rerun: returns stored result; does not double-insert.
    retry = confirm_shared_note_history_import(
        client_id=client_a, batch_id=upload.batch_id, actor=actor
    )
    assert retry.ok
    assert retry.history_inserted == 4
    assert retry.staging_cleanup_status == "completed"
    with get_connection() as conn:
        total = int(
            conn.execute(
                """
                SELECT COUNT(*) AS c FROM company_shared_history_events
                WHERE import_batch_id = ?
                """,
                (upload.batch_id,),
            ).fetchone()["c"]
        )
        assert total == 4
        batch = conn.execute(
            "SELECT status, prospects_path, history_path, staging_cleanup_status, "
            "prospects_filename, history_filename, prospects_sha256, history_sha256, "
            "plan_fingerprint, result_json FROM shared_note_history_import_batches WHERE id = ?",
            (upload.batch_id,),
        ).fetchone()
        assert batch["status"] == "confirmed"
        assert batch["staging_cleanup_status"] == "completed"
        assert batch["prospects_path"] == ""
        assert batch["history_path"] == ""
        assert batch["prospects_filename"]
        assert batch["history_filename"]
        assert batch["prospects_sha256"]
        assert batch["history_sha256"]
        assert batch["plan_fingerprint"]
        assert batch["result_json"]

    assert _freeze_snapshot() == before_freeze
    _delete_user(user_id)
    print("PASS preview/confirm/visibility/order/idempotency/freeze")


def test_authorization_and_csrf() -> None:
    _prove_isolated()
    with get_connection() as conn:
        ensure_shared_note_history_schema(conn)
        conn.commit()
    client_a, _ = _client_ids()
    password = _secret_password()
    email = f"snh.staff.{secrets.token_hex(6)}@example.test"
    user_id = _create_user(email=email, password=password, administrator=0)
    _assign(user_id, client_a)
    actor = _actor(user_id)
    prospects_bytes, history_bytes = _fixture_csvs()
    try:
        upload_shared_note_history_import(
            client_id=client_a,
            actor=actor,
            prospects_filename="p.csv",
            prospects_content=prospects_bytes,
            history_filename="h.csv",
            history_content=history_bytes,
        )
        _fail("Non-admin upload should be rejected.")
    except PermissionError:
        pass

    http = _client()
    login = http.post("/api/auth/login", json={"email": email, "password": password})
    assert login.status_code == 200
    csrf = _csrf(login)
    # Non-admin rejected
    resp = http.post(
        UPLOAD.format(client_id=client_a),
        headers={CSRF_HEADER: csrf},
        files={
            "prospects_file": ("p.csv", prospects_bytes, "text/csv"),
            "history_file": ("h.csv", history_bytes, "text/csv"),
        },
    )
    if resp.status_code not in (401, 403):
        _fail(f"Expected 401/403 for non-admin, got {resp.status_code}")

    # Admin without CSRF rejected
    admin_pw = _secret_password()
    admin_email = f"snh.csrf.{secrets.token_hex(6)}@example.test"
    admin_id = _create_user(email=admin_email, password=admin_pw, administrator=1)
    _assign(admin_id, client_a)
    login2 = http.post("/api/auth/login", json={"email": admin_email, "password": admin_pw})
    assert login2.status_code == 200
    no_csrf = http.post(
        UPLOAD.format(client_id=client_a),
        files={
            "prospects_file": ("p.csv", prospects_bytes, "text/csv"),
            "history_file": ("h.csv", history_bytes, "text/csv"),
        },
    )
    if no_csrf.status_code not in (401, 403):
        _fail(f"Expected CSRF rejection, got {no_csrf.status_code}")

    _delete_user(user_id)
    _delete_user(admin_id)
    print("PASS authorization + CSRF")


def test_atomic_rollback_on_failure() -> None:
    _prove_isolated()
    with get_connection() as conn:
        ensure_shared_note_history_schema(conn)
        migrate_schema(conn)
        conn.commit()

    client_a, _ = _client_ids()
    password = _secret_password()
    email = f"snh.rollback.{secrets.token_hex(6)}@example.test"
    user_id = _create_user(email=email, password=password, administrator=1)
    _assign(user_id, client_a)
    actor = _actor(user_id)
    rn = f"88{secrets.token_hex(3)}"
    prospects = _csv_bytes(
        ["Record No.", "Company", "Status", "FirstName", "LastName", "Email"],
        [[rn, "Rollback Co", "New", "Ray", "Back", "ray@rollback.test"]],
    )
    history = _csv_bytes(
        [
            "LeadMaster Record No.",
            "Company",
            "Current Dawson Status",
            "Event Sequence",
            "Event Timestamp",
            "Author",
            "Event Type",
            "Attributed Client/Campaign",
            "Attribution Evidence",
            "Note Text",
            "Source File",
            "Event Hash",
        ],
        [
            [
                rn,
                "Rollback Co",
                "New",
                "1",
                "2026-02-01T10:00:00Z",
                "Author",
                "Note",
                UNATTRIBUTED_LABEL,
                "",
                "Should roll back",
                "fixture.csv",
                f"hash-rollback-{rn}",
            ]
        ],
    )
    upload = upload_shared_note_history_import(
        client_id=client_a,
        actor=actor,
        prospects_filename="p.csv",
        prospects_content=prospects,
        history_filename="h.csv",
        history_content=history,
    )
    with get_connection() as conn:
        before_companies = int(
            conn.execute(
                "SELECT COUNT(*) AS c FROM companies WHERE external_record_no = ?",
                (rn,),
            ).fetchone()["c"]
        )
        before_events = int(
            conn.execute("SELECT COUNT(*) AS c FROM company_shared_history_events").fetchone()["c"]
        )

    try:
        confirm_shared_note_history_import(
            client_id=client_a,
            batch_id=upload.batch_id,
            actor=actor,
            fail_after="after_history_insert",
        )
        _fail("Expected injected failure.")
    except RuntimeError:
        pass

    with get_connection() as conn:
        after_companies = int(
            conn.execute(
                "SELECT COUNT(*) AS c FROM companies WHERE external_record_no = ?",
                (rn,),
            ).fetchone()["c"]
        )
        after_events = int(
            conn.execute("SELECT COUNT(*) AS c FROM company_shared_history_events").fetchone()["c"]
        )
        batch_status = conn.execute(
            "SELECT status FROM shared_note_history_import_batches WHERE id = ?",
            (upload.batch_id,),
        ).fetchone()["status"]
    assert after_companies == before_companies
    assert after_events == before_events
    assert batch_status == "previewed"
    _delete_user(user_id)
    print("PASS rollback")


def test_http_admin_preview_and_confirm() -> None:
    _prove_isolated()
    with get_connection() as conn:
        ensure_shared_note_history_schema(conn)
        conn.commit()
    client_a, _ = _client_ids()
    password = _secret_password()
    email = f"snh.http.{secrets.token_hex(6)}@example.test"
    user_id = _create_user(email=email, password=password, administrator=1)
    _assign(user_id, client_a)
    http = _client()
    login = http.post("/api/auth/login", json={"email": email, "password": password})
    assert login.status_code == 200
    csrf = _csrf(login)
    rn = f"77{secrets.token_hex(3)}"
    prospects = _csv_bytes(
        ["Record No.", "Company", "Status", "FirstName", "LastName"],
        [[rn, "Http Co", "Left Message", "Hal", "Http"]],
    )
    history = _csv_bytes(
        [
            "LeadMaster Record No.",
            "Company",
            "Current Dawson Status",
            "Event Sequence",
            "Event Timestamp",
            "Author",
            "Event Type",
            "Attributed Client/Campaign",
            "Attribution Evidence",
            "Note Text",
            "Source File",
            "Event Hash",
        ],
        [
            [
                rn,
                "Http Co",
                "Left Message",
                "1",
                "2026-03-01T09:00:00Z",
                "Author",
                "Call",
                "Dawson Fabrication",
                "Dawson -",
                "HTTP import event",
                "fixture.csv",
                f"hash-http-{rn}",
            ]
        ],
    )
    upload = http.post(
        UPLOAD.format(client_id=client_a),
        headers={CSRF_HEADER: csrf},
        files={
            "prospects_file": (
                "Dawson_NorthStar_Prospects_No_Closed.csv",
                prospects,
                "text/csv",
            ),
            "history_file": (
                "Dawson_Shared_Note_History_No_Closed.csv",
                history,
                "text/csv",
            ),
        },
    )
    assert upload.status_code == 200, upload.text
    payload = upload.json()
    assert payload["preview"]["ok"] is True
    assert payload["preview"]["counts"]["companies_included"] == 1
    batch_id = int(payload["batch_id"])
    confirm = http.post(
        CONFIRM.format(client_id=client_a, batch_id=batch_id),
        headers={CSRF_HEADER: csrf},
        params={"confirm": "true"},
    )
    assert confirm.status_code == 200, confirm.text
    body = confirm.json()
    assert body["history_inserted"] == 1
    assert body["companies_created"] == 1
    _delete_user(user_id)
    print("PASS http preview/confirm")


def _history_csv_for_rn(rn: str, company: str, event_hash: str) -> bytes:
    return _csv_bytes(
        [
            "LeadMaster Record No.",
            "Company",
            "Current Dawson Status",
            "Event Sequence",
            "Event Timestamp",
            "Author",
            "Event Type",
            "Attributed Client/Campaign",
            "Attribution Evidence",
            "Note Text",
            "Source File",
            "Event Hash",
        ],
        [
            [
                rn,
                company,
                "New",
                "1",
                "2026-01-01T12:00:00Z",
                "Author",
                "Note",
                "Dawson Fabrication",
                "Dawson -",
                "note",
                "fixture.csv",
                event_hash,
            ]
        ],
    )


def test_preview_confirm_contact_counters_reconcile_with_in_batch_registry() -> None:
    """Preview must count in-batch contact identity collisions the same as confirm."""
    _prove_isolated()
    with get_connection() as conn:
        ensure_shared_note_history_schema(conn)
        migrate_schema(conn)
        conn.commit()
    client_a, _ = _client_ids()
    password = _secret_password()
    email = f"snh.contact.reconcile.{secrets.token_hex(6)}@example.test"
    user_id = _create_user(email=email, password=password, administrator=1)
    _assign(user_id, client_a)
    actor = _actor(user_id)
    rn = f"91{secrets.token_hex(3)}"
    prospects = _csv_bytes(
        [
            "Record No.",
            "Company",
            "Status",
            "FirstName",
            "LastName",
            "Email",
            "Phone",
        ],
        [
            [rn, "Dup Co", "New", "Ann", "Alpha", "ann@dup.example", ""],
            [rn, "Dup Co", "New", "Ann", "Alpha", "ann@dup.example", ""],
            [rn, "Dup Co", "New", "Bob", "Beta", "", "555-0100"],
            [rn, "Dup Co", "New", "Bob", "Beta", "", "555-0199"],
        ],
    )
    history = _history_csv_for_rn(rn, "Dup Co", f"hash-dup-{rn}")
    upload = upload_shared_note_history_import(
        client_id=client_a,
        actor=actor,
        prospects_filename="p.csv",
        prospects_content=prospects,
        history_filename="h.csv",
        history_content=history,
    )
    assert upload.preview.ok
    assert upload.preview.counts.contacts_to_create == 2
    assert upload.preview.counts.contacts_to_reuse == 2
    result = confirm_shared_note_history_import(
        client_id=client_a, batch_id=upload.batch_id, actor=actor
    )
    assert result.contacts_created == upload.preview.counts.contacts_to_create == 2
    assert result.contacts_reused == upload.preview.counts.contacts_to_reuse == 2
    assert result.staging_cleanup_status == CLEANUP_COMPLETED
    with get_connection() as conn:
        n = int(
            conn.execute(
                "SELECT COUNT(*) AS c FROM contacts WHERE external_record_no = ?",
                (rn,),
            ).fetchone()["c"]
        )
        assert n == 2
        batch = conn.execute(
            "SELECT prospects_path, history_path, staging_cleanup_status FROM "
            "shared_note_history_import_batches WHERE id = ?",
            (upload.batch_id,),
        ).fetchone()
        assert batch["staging_cleanup_status"] == CLEANUP_COMPLETED
        assert batch["prospects_path"] == ""
        assert batch["history_path"] == ""
    _delete_user(user_id)
    print("PASS preview/confirm contact counter reconcile")


def test_staging_cleanup_failure_pending_and_retry() -> None:
    _prove_isolated()
    with get_connection() as conn:
        ensure_shared_note_history_schema(conn)
        conn.commit()
    client_a, _ = _client_ids()
    password = _secret_password()
    email = f"snh.cleanup.{secrets.token_hex(6)}@example.test"
    user_id = _create_user(email=email, password=password, administrator=1)
    _assign(user_id, client_a)
    actor = _actor(user_id)
    rn = f"92{secrets.token_hex(3)}"
    prospects = _csv_bytes(
        ["Record No.", "Company", "Status", "FirstName", "LastName"],
        [[rn, "Cleanup Co", "New", "Cara", "Clean"]],
    )
    history = _history_csv_for_rn(rn, "Cleanup Co", f"hash-clean-{rn}")
    upload = upload_shared_note_history_import(
        client_id=client_a,
        actor=actor,
        prospects_filename="p.csv",
        prospects_content=prospects,
        history_filename="h.csv",
        history_content=history,
    )
    with get_connection() as conn:
        paths = conn.execute(
            "SELECT prospects_path, history_path FROM shared_note_history_import_batches "
            "WHERE id = ?",
            (upload.batch_id,),
        ).fetchone()
    assert Path(paths["prospects_path"]).is_file()
    assert Path(paths["history_path"]).is_file()

    def boom(_path: Path) -> None:
        raise OSError("simulated staging delete failure")

    result = confirm_shared_note_history_import(
        client_id=client_a,
        batch_id=upload.batch_id,
        actor=actor,
        unlink_fn=boom,
    )
    assert result.ok
    assert result.history_inserted == 1
    assert result.staging_cleanup_status == CLEANUP_PENDING
    assert "simulated staging delete failure" in result.staging_cleanup_error
    with get_connection() as conn:
        batch = conn.execute(
            "SELECT status, prospects_path, history_path, staging_cleanup_status, "
            "prospects_filename, prospects_sha256, result_json, plan_fingerprint "
            "FROM shared_note_history_import_batches WHERE id = ?",
            (upload.batch_id,),
        ).fetchone()
        assert batch["status"] == "confirmed"
        assert batch["staging_cleanup_status"] == CLEANUP_PENDING
        assert batch["prospects_path"] == paths["prospects_path"]
        assert batch["history_path"] == paths["history_path"]
        assert batch["prospects_filename"]
        assert batch["prospects_sha256"]
        assert batch["result_json"]
        assert batch["plan_fingerprint"]
        assert Path(batch["prospects_path"]).is_file()

    retry = retry_shared_note_history_staging_cleanup(
        client_id=client_a, batch_id=upload.batch_id, actor=actor
    )
    assert retry.staging_cleanup_status == CLEANUP_COMPLETED
    with get_connection() as conn:
        batch = conn.execute(
            "SELECT status, prospects_path, history_path, staging_cleanup_status, "
            "prospects_filename, prospects_sha256 FROM shared_note_history_import_batches "
            "WHERE id = ?",
            (upload.batch_id,),
        ).fetchone()
        events = int(
            conn.execute(
                "SELECT COUNT(*) FROM company_shared_history_events WHERE import_batch_id = ?",
                (upload.batch_id,),
            ).fetchone()[0]
        )
        assert batch["status"] == "confirmed"
        assert batch["staging_cleanup_status"] == CLEANUP_COMPLETED
        assert batch["prospects_path"] == ""
        assert batch["history_path"] == ""
        assert batch["prospects_filename"]
        assert batch["prospects_sha256"]
        assert events == 1
        assert not Path(paths["prospects_path"]).exists()
        assert not Path(paths["history_path"]).exists()

    again = confirm_shared_note_history_import(
        client_id=client_a, batch_id=upload.batch_id, actor=actor
    )
    assert again.history_inserted == 1
    assert again.staging_cleanup_status == CLEANUP_COMPLETED
    with get_connection() as conn:
        events = int(
            conn.execute(
                "SELECT COUNT(*) FROM company_shared_history_events WHERE import_batch_id = ?",
                (upload.batch_id,),
            ).fetchone()[0]
        )
        assert events == 1
    _delete_user(user_id)
    print("PASS staging cleanup failure/pending/retry/idempotency")


def test_dawson_batch1_retrospective_contact_plan_read_only() -> None:
    """Against pre-import backup, corrected planner yields 783/214 and lists 10 in-batch reuses."""
    if not DAWSON_PROSPECTS.is_file() or not DAWSON_HISTORY.is_file():
        print("SKIP dawson retrospective (prepared CSVs missing)")
        return
    if not DAWSON_PREIMPORT_BACKUP.is_file():
        print("SKIP dawson retrospective (pre-import backup missing)")
        return
    import shutil
    import sqlite3
    import tempfile

    prospects = parse_prospects_csv(DAWSON_PROSPECTS.read_bytes())
    history = parse_history_csv(
        DAWSON_HISTORY.read_bytes(),
        allowed_record_nos=set(prospects.included_company_record_nos),
    )
    # Copy pre-import backup to a throwaway file; never open production for writes.
    tmp = tempfile.NamedTemporaryFile(prefix="snh-dawson-retro-", suffix=".db", delete=False)
    tmp_path = Path(tmp.name)
    tmp.close()
    shutil.copy2(DAWSON_PREIMPORT_BACKUP, tmp_path)
    conn = sqlite3.connect(str(tmp_path))
    conn.row_factory = sqlite3.Row
    try:
        ensure_shared_note_history_schema(conn)
        counts = build_preview_counts(
            conn, client_id=3, prospects=prospects, history=history
        )
        assert counts.companies_to_create == 208
        assert counts.companies_to_reuse == 35
        assert counts.contacts_to_create == 783
        assert counts.contacts_to_reuse == 214
        assert counts.history_to_insert == 13319

        company_ids: dict[str, int] = {}
        for rn in prospects.included_company_record_nos:
            existing = _company_by_record_no(conn, rn)
            if existing is not None:
                company_ids[rn] = int(existing["id"])
        _create, _reuse, in_batch = plan_contact_actions(
            conn, prospects=prospects, company_ids=company_ids
        )
        assert len(in_batch) == 10
        print("DAWSON_IN_BATCH_CONTACT_REUSES")
        for item in in_batch:
            print(
                f"  RN={item.record_no} company={item.company_name!r} "
                f"row={item.source_row} matched_row={item.matched_source_row} "
                f"name={item.first_name} {item.last_name} email={item.email}"
            )
    finally:
        conn.close()
        try:
            tmp_path.unlink(missing_ok=True)
        except OSError:
            pass
    print("PASS dawson retrospective 783/214 + 10 in-batch reuses")


def main() -> None:
    _prove_isolated()
    test_normalize_record_no()
    print("PASS normalize_record_no")
    test_closed_exclusion_and_exact_record_match()
    print("PASS closed exclusion + Record No. match + attribution mapping")
    test_dawson_prepared_files_read_only_counts()
    test_preview_confirm_visibility_order_idempotency_freeze()
    test_authorization_and_csrf()
    test_atomic_rollback_on_failure()
    test_http_admin_preview_and_confirm()
    test_preview_confirm_contact_counters_reconcile_with_in_batch_registry()
    test_staging_cleanup_failure_pending_and_retry()
    test_dawson_batch1_retrospective_contact_plan_read_only()
    print("ALL SHARED NOTE-HISTORY IMPORT TESTS PASSED")


if __name__ == "__main__":
    main()
