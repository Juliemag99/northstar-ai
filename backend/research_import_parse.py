"""Parse research-prospect CSV/XLSX uploads. Does not write master data."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from typing import Any

from research_import_mapping import blank

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_COLUMNS = 40
MAX_SOURCE_ROWS = 10_000
MAX_CELL_CHARS = 20_000
GENERIC_PARSE_ERROR = "The spreadsheet could not be read. Check the file and try again."
UNSUPPORTED_TYPE = "Upload a CSV or XLSX file."
NEED_WORKSHEET = "This workbook has multiple sheets. Choose one visible sheet to continue."
HIDDEN_SHEET = "Hidden worksheets cannot be imported. Choose a visible sheet."
TOO_MANY_COLUMNS = f"This spreadsheet has more than {MAX_COLUMNS} columns."
TOO_MANY_ROWS = f"This spreadsheet has more than {MAX_SOURCE_ROWS} data rows."
FILE_TOO_LARGE = "This file is too large to upload."
JANCO_PROSPECT_SHEET = "Qualified Prospects"


def file_sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def detect_file_type(filename: str, content: bytes) -> str:
    name = blank(filename).lower()
    if name.endswith(".csv"):
        return "csv"
    if name.endswith(".xlsx"):
        return "xlsx"
    if content.startswith(b"PK"):
        return "xlsx"
    return "csv" if b"," in content[:200] else ""


def _cell(value: object | None) -> str:
    text = blank(value)
    # URLs are provenance; never truncate them even if a cell is unusually long.
    if re.match(r"^https?://", text, re.I):
        return text
    if len(text) > MAX_CELL_CHARS:
        return text[:MAX_CELL_CHARS]
    return text


def parse_csv(content: bytes) -> tuple[list[str], list[dict[str, str]], list[str]]:
    warnings: list[str] = []
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
        warnings.append("File was not UTF-8; Latin-1 was used to read it.")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        raise ValueError(GENERIC_PARSE_ERROR)
    headers = [_cell(h) for h in rows[0]]
    if len(headers) > MAX_COLUMNS:
        raise ValueError(TOO_MANY_COLUMNS)
    data: list[dict[str, str]] = []
    for raw in rows[1:]:
        if len(data) >= MAX_SOURCE_ROWS:
            raise ValueError(TOO_MANY_ROWS)
        values = {_cell(headers[i]): _cell(raw[i] if i < len(raw) else "") for i in range(len(headers))}
        if not any(values.values()):
            continue
        data.append(values)
    return headers, data, warnings


def parse_xlsx(
    content: bytes, *, worksheet: str = ""
) -> tuple[str, list[str], list[dict[str, str]], list[str], list[str]]:
    from openpyxl import load_workbook

    warnings: list[str] = []
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_vba=False)
    except Exception as exc:
        raise ValueError(GENERIC_PARSE_ERROR) from exc
    visible = []
    for ws in wb.worksheets:
        hidden = bool(getattr(ws, "sheet_state", "visible") and str(ws.sheet_state) != "visible")
        if hidden:
            continue
        visible.append(ws.title)
    if not visible:
        wb.close()
        raise ValueError(HIDDEN_SHEET)
    requested = blank(worksheet)
    if not requested:
        if JANCO_PROSPECT_SHEET in visible:
            requested = JANCO_PROSPECT_SHEET
        elif len(visible) == 1:
            requested = visible[0]
        else:
            wb.close()
            return "", [], [], warnings, visible
    if requested not in visible:
        wb.close()
        raise ValueError(NEED_WORKSHEET if requested not in [ws.title for ws in wb.worksheets] else HIDDEN_SHEET)
    ws = wb[requested]
    raw_rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not raw_rows:
        raise ValueError(GENERIC_PARSE_ERROR)
    headers = [_cell(c) for c in raw_rows[0]]
    while headers and not headers[-1]:
        headers.pop()
    if len(headers) > MAX_COLUMNS:
        raise ValueError(TOO_MANY_COLUMNS)
    data: list[dict[str, str]] = []
    for rec in raw_rows[1:]:
        if len(data) >= MAX_SOURCE_ROWS:
            raise ValueError(TOO_MANY_ROWS)
        values = {
            headers[i]: _cell(rec[i] if rec is not None and i < len(rec) else "")
            for i in range(len(headers))
        }
        if not any(values.values()):
            continue
        data.append(values)
    return requested, headers, data, warnings, visible


def parse_upload(
    filename: str, content: bytes, *, worksheet: str = ""
) -> dict[str, Any]:
    if len(content) > MAX_FILE_BYTES:
        raise ValueError(FILE_TOO_LARGE)
    file_type = detect_file_type(filename, content)
    if file_type not in {"csv", "xlsx"}:
        raise ValueError(UNSUPPORTED_TYPE)
    if file_type == "xlsx":
        try:
            zipfile.ZipFile(io.BytesIO(content))
        except zipfile.BadZipFile as exc:
            raise ValueError(GENERIC_PARSE_ERROR) from exc
        sheet, headers, rows, warnings, visible = parse_xlsx(content, worksheet=worksheet)
        if not sheet and visible:
            return {
                "kind": "needs_worksheet",
                "needs_worksheet": True,
                "visible_sheets": visible,
                "filename": blank(filename),
                "file_type": file_type,
                "message": NEED_WORKSHEET,
                "headers": [],
                "rows": [],
                "warnings": warnings,
                "sha256": file_sha256(content),
                "worksheet_name": "",
            }
    else:
        headers, rows, warnings = parse_csv(content)
        sheet = "CSV"
        visible = ["CSV"]
    return {
        "kind": "ok",
        "needs_worksheet": False,
        "visible_sheets": visible,
        "filename": blank(filename),
        "file_type": file_type,
        "message": "",
        "headers": headers,
        "rows": rows,
        "warnings": warnings,
        "sha256": file_sha256(content),
        "worksheet_name": sheet,
        "file_size_bytes": len(content),
    }


def rows_to_json(rows: list[dict[str, str]]) -> list[str]:
    return [json.dumps(row, ensure_ascii=False) for row in rows]
