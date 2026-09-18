"""Phase 3 — Appointment / Engagement workbook importer.

Preview and match without writing CRM until Confirm Import.
Never auto-creates companies/contacts unless user chooses Create New on a row.
"""

from __future__ import annotations

from staff_context import resolve_staff_actor

import hashlib
import json
import os
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from access import get_default_user, get_user_by_id, user_can_access_client
from client_engagement_appt_grid import (
    already_imported_by_fingerprint,
    already_imported_by_provenance,
    appointment_grid_fingerprint,
    build_delta_caller_notes,
    classify_already_imported,
    classify_appointment_narrative,
    classify_event_type,
    classify_field_vs_history,
    compose_forecast_quoted_won_block,
    corroborate_year_from_history,
    is_footer_or_total_row,
    match_company_for_appt_grid,
    match_contact_for_appt_grid,
    merge_sales_notes_with_fqw,
    norm_text,
    parse_datetime_text,
    propose_row_action,
    row_blocks_confirm,
)
from client_setup_data import user_can_edit_client_setup
from db import DATABASE_DIR, get_connection
from models import (
    EngagementImportBatchView,
    EngagementImportConfirmRequest,
    EngagementImportPreview,
    EngagementImportRowView,
    EngagementImportSheetView,
    EngagementImportSummary,
    SalesEventView,
    SheetClassificationUpdate,
)

DOCUMENTS_ROOT = DATABASE_DIR / "client_documents"

SHEET_TYPES = (
    "Appointments",
    "Send Information",
    "Engagement",
    "Quote",
    "Purchase Order",
    "WebLead",
    "Other Sales Activity",
    "Ignore",
)

TARGET_FIELDS = (
    "list_name",
    "external_record_no",
    "northstar_client",
    "revenue_specialist",
    "appointment_date_time",
    "company",
    "contact",
    "title",
    "address",
    "city_state_zip",
    "phone",
    "email",
    "caller_notes",
    "appointment_grade",
    "sales_notes",
    "dollars_quoted",
    "outcome",
)

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "list_name": ("list",),
    "external_record_no": ("record number", "record no", "record #"),
    "northstar_client": ("northstar client", "client", "ns client"),
    "revenue_specialist": ("rev spec", "revenue specialist", "revspec"),
    "appointment_date_time": (
        "appointment date/time",
        "appointment date time",
        "appointment datetime",
        "date/time",
        "date time",
    ),
    "company": ("company name", "company"),
    "contact": ("contact person", "contact", "contact name"),
    "title": ("title",),
    "address": ("address",),
    "city_state_zip": ("city/state/zip code", "city/state/zip", "city state zip"),
    "phone": ("phone number", "phone"),
    "email": ("email address", "email"),
    "caller_notes": ("northstar caller notes", "caller notes"),
    "appointment_grade": ("appt grade", "appointment grade"),
    "sales_notes": ("sales notes", "client notes"),
    "dollars_quoted": ("dollars quoted", "quoted", "quote amount"),
    "outcome": (
        "account won, lost or in progress",
        "won lost in progress",
        "outcome",
    ),
}


def _now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _blank(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _norm_name(value: str) -> str:
    t = _blank(value).lower()
    t = re.sub(r"[^a-z0-9]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def _norm_client(value: str) -> str:
    return re.sub(r"\s+", " ", _blank(value)).lower()


def _digits(value: str) -> str:
    return re.sub(r"\D", "", _blank(value))


def ensure_engagement_import_schema(conn=None) -> None:
    owns = conn is None
    if owns:
        conn = get_connection()
    try:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS client_engagement_import_batches (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                client_id INTEGER NOT NULL,
                document_id INTEGER,
                filename TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL DEFAULT 'uploaded',
                uploaded_at TEXT NOT NULL DEFAULT '',
                uploaded_by_user_id INTEGER,
                uploaded_by_name TEXT NOT NULL DEFAULT '',
                confirmed_at TEXT NOT NULL DEFAULT '',
                confirmed_by_name TEXT NOT NULL DEFAULT '',
                notes TEXT NOT NULL DEFAULT '',
                summary_json TEXT NOT NULL DEFAULT '{}',
                FOREIGN KEY (client_id) REFERENCES clients(id) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS client_engagement_import_sheets (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                batch_id INTEGER NOT NULL,
                client_id INTEGER NOT NULL,
                sheet_name TEXT NOT NULL,
                detected_type TEXT NOT NULL DEFAULT 'Other Sales Activity',
                user_type TEXT NOT NULL DEFAULT '',
                headers_json TEXT NOT NULL DEFAULT '[]',
                column_mapping_json TEXT NOT NULL DEFAULT '{}',
                row_count INTEGER NOT NULL DEFAULT 0,
                is_empty INTEGER NOT NULL DEFAULT 0,
                FOREIGN KEY (batch_id) REFERENCES client_engagement_import_batches(id) ON DELETE CASCADE,
                FOREIGN KEY (client_id) REFERENCES clients(id) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS client_engagement_import_rows (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                batch_id INTEGER NOT NULL,
                sheet_id INTEGER NOT NULL,
                client_id INTEGER NOT NULL,
                source_row_number INTEGER NOT NULL DEFAULT 0,
                raw_json TEXT NOT NULL DEFAULT '{}',
                mapped_json TEXT NOT NULL DEFAULT '{}',
                sheet_name TEXT NOT NULL DEFAULT '',
                sheet_type TEXT NOT NULL DEFAULT '',
                event_type TEXT NOT NULL DEFAULT 'Other',
                source_date_time_text TEXT NOT NULL DEFAULT '',
                event_date TEXT NOT NULL DEFAULT '',
                event_time TEXT NOT NULL DEFAULT '',
                timezone TEXT NOT NULL DEFAULT '',
                meeting_type TEXT NOT NULL DEFAULT '',
                datetime_needs_review INTEGER NOT NULL DEFAULT 0,
                company_match_status TEXT NOT NULL DEFAULT 'NEW',
                company_id INTEGER,
                relationship_id INTEGER,
                company_name TEXT NOT NULL DEFAULT '',
                company_record_no TEXT NOT NULL DEFAULT '',
                contact_match_status TEXT NOT NULL DEFAULT 'NEW',
                contact_id INTEGER,
                contact_name TEXT NOT NULL DEFAULT '',
                client_validation TEXT NOT NULL DEFAULT 'OK',
                rev_spec_user_id INTEGER,
                source_rev_spec_text TEXT NOT NULL DEFAULT '',
                rev_spec_needs_review INTEGER NOT NULL DEFAULT 0,
                source_appointment_grade TEXT NOT NULL DEFAULT '',
                quoted_amount REAL,
                source_quoted_value TEXT NOT NULL DEFAULT '',
                potential_quote_signal INTEGER NOT NULL DEFAULT 0,
                outcome_normalized TEXT NOT NULL DEFAULT '',
                source_outcome TEXT NOT NULL DEFAULT '',
                thread_key TEXT NOT NULL DEFAULT '',
                source_row_fingerprint TEXT NOT NULL DEFAULT '',
                duplicate_status TEXT NOT NULL DEFAULT '',
                import_status TEXT NOT NULL DEFAULT 'staged',
                selected INTEGER NOT NULL DEFAULT 1,
                needs_review INTEGER NOT NULL DEFAULT 0,
                review_flags_json TEXT NOT NULL DEFAULT '[]',
                event_id INTEGER,
                FOREIGN KEY (batch_id) REFERENCES client_engagement_import_batches(id) ON DELETE CASCADE,
                FOREIGN KEY (sheet_id) REFERENCES client_engagement_import_sheets(id) ON DELETE CASCADE,
                FOREIGN KEY (client_id) REFERENCES clients(id) ON DELETE CASCADE
            );

            CREATE INDEX IF NOT EXISTS idx_eng_import_rows_batch
                ON client_engagement_import_rows(batch_id, sheet_id, source_row_number);

            CREATE TABLE IF NOT EXISTS client_sales_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                client_id INTEGER NOT NULL,
                relationship_id INTEGER,
                company_id INTEGER,
                contact_id INTEGER,
                thread_key TEXT NOT NULL DEFAULT '',
                parent_event_id INTEGER,
                event_type TEXT NOT NULL DEFAULT 'Other',
                event_family TEXT NOT NULL DEFAULT '',
                source_date_time_text TEXT NOT NULL DEFAULT '',
                event_date TEXT NOT NULL DEFAULT '',
                event_time TEXT NOT NULL DEFAULT '',
                timezone TEXT NOT NULL DEFAULT '',
                meeting_type TEXT NOT NULL DEFAULT '',
                company_name TEXT NOT NULL DEFAULT '',
                contact_name TEXT NOT NULL DEFAULT '',
                contact_title TEXT NOT NULL DEFAULT '',
                phone TEXT NOT NULL DEFAULT '',
                email TEXT NOT NULL DEFAULT '',
                address TEXT NOT NULL DEFAULT '',
                city_state_zip TEXT NOT NULL DEFAULT '',
                caller_notes TEXT NOT NULL DEFAULT '',
                sales_notes TEXT NOT NULL DEFAULT '',
                source_appointment_grade TEXT NOT NULL DEFAULT '',
                source_rev_spec_text TEXT NOT NULL DEFAULT '',
                rev_spec_user_id INTEGER,
                quoted_amount REAL,
                source_quoted_value TEXT NOT NULL DEFAULT '',
                potential_quote_signal INTEGER NOT NULL DEFAULT 0,
                outcome_normalized TEXT NOT NULL DEFAULT '',
                source_outcome TEXT NOT NULL DEFAULT '',
                source_document_id INTEGER,
                source_batch_id INTEGER,
                source_sheet TEXT NOT NULL DEFAULT '',
                source_row INTEGER NOT NULL DEFAULT 0,
                source_file_name TEXT NOT NULL DEFAULT '',
                source_record_number TEXT NOT NULL DEFAULT '',
                source_row_fingerprint TEXT NOT NULL DEFAULT '',
                imported_at TEXT NOT NULL DEFAULT '',
                imported_by_user_id INTEGER,
                imported_by_name TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL DEFAULT '',
                location_id INTEGER,
                FOREIGN KEY (client_id) REFERENCES clients(id) ON DELETE CASCADE
            );

            CREATE INDEX IF NOT EXISTS idx_sales_events_client
                ON client_sales_events(client_id, event_date DESC, id DESC);
            CREATE INDEX IF NOT EXISTS idx_sales_events_company
                ON client_sales_events(client_id, company_id, event_date DESC);
            CREATE INDEX IF NOT EXISTS idx_sales_events_contact
                ON client_sales_events(client_id, contact_id, event_date DESC);
            CREATE UNIQUE INDEX IF NOT EXISTS idx_sales_events_fingerprint
                ON client_sales_events(client_id, source_row_fingerprint)
                WHERE source_row_fingerprint <> '';
            CREATE INDEX IF NOT EXISTS idx_client_sales_events_location_id
                ON client_sales_events(location_id)
                WHERE location_id IS NOT NULL;
            """
        )
        if owns:
            conn.commit()
    finally:
        if owns:
            conn.close()


def _require_access(user_id: int, client_id: int) -> None:
    from access import require_write_client_id

    require_write_client_id(client_id, user_id=user_id)


def _require_edit(user_id: int, client_id: int) -> None:
    if not user_can_edit_client_setup(user_id, client_id):
        raise PermissionError("Not authorized to import for this client.")


def _documents_root() -> Path:
    """Keep Appointment Grid test uploads off the live client_documents tree."""
    test_db = os.environ.get("NORTHSTAR_TEST_DB", "").strip()
    if test_db:
        return Path(test_db + ".docs")
    return DOCUMENTS_ROOT


def _client_dir(client_id: int) -> Path:
    path = _documents_root() / str(int(client_id))
    path.mkdir(parents=True, exist_ok=True)
    return path


def _document_file_sha256(conn, client_id: int, document_id: int | None) -> str:
    if not document_id:
        return ""
    drow = conn.execute(
        """
        SELECT stored_filename FROM client_documents
        WHERE id = ? AND client_id = ?
        """,
        (int(document_id), int(client_id)),
    ).fetchone()
    stored = _blank(drow["stored_filename"]) if drow else ""
    if not stored:
        return ""
    path = _documents_root() / str(int(client_id)) / stored
    if not path.is_file():
        return ""
    return hashlib.sha256(path.read_bytes()).hexdigest().lower()


def _same_source_document_ids(
    conn, client_id: int, document_id: int | None, document_sha256: str
) -> set[int]:
    """Document ids for this client that share identity with the current file.

    Identity is document_id and/or immutable SHA-256 of stored bytes.
    Never includes another client's documents.
    """
    ids: set[int] = set()
    if document_id:
        ids.add(int(document_id))
    sha = _blank(document_sha256).lower()
    if not sha:
        return ids
    rows = conn.execute(
        """
        SELECT DISTINCT source_document_id AS id
        FROM client_sales_events
        WHERE client_id = ? AND source_document_id IS NOT NULL
        """,
        (int(client_id),),
    ).fetchall()
    cache: dict[int, str] = {}
    if document_id:
        cache[int(document_id)] = sha
    for row in rows:
        did = int(row["id"])
        if did in ids:
            continue
        if did not in cache:
            cache[did] = _document_file_sha256(conn, client_id, did)
        if cache[did] == sha:
            ids.add(did)
    return ids


def _detect_sheet_type(sheet_name: str, headers: list[str], rows: list[list[str]]) -> str:
    name = _blank(sheet_name).lower()
    if not rows or not any(any(_blank(c) for c in r) for r in rows):
        return "Ignore"
    if "send" in name and "mail" in name:
        return "Send Information"
    if "appointment" in name:
        return "Appointments"
    if "quote" in name:
        return "Quote"
    if "purchase" in name or name in {"po", "p.o."}:
        return "Purchase Order"
    if "weblead" in name or "web lead" in name:
        return "WebLead"
    if "sheet3" in name or re.match(r"^sheet\s*\d+$", name):
        # empty already handled; non-empty unknown sheet stays Other
        return "Other Sales Activity"
    return "Other Sales Activity"


def _auto_map_columns(headers: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    norm_headers = {_norm_name(h): h for h in headers if _blank(h)}
    for target, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            key = _norm_name(alias)
            if key in norm_headers:
                mapping[target] = norm_headers[key]
                break
    return mapping


def _parse_workbook(path: Path) -> list[dict[str, Any]]:
    ext = path.suffix.lower()
    sheets: list[dict[str, Any]] = []
    if ext == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
        for name in wb.sheetnames:
            ws = wb[name]
            raw_rows: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                raw_rows.append([_blank(c) for c in row])
            # trim fully empty trailing rows
            while raw_rows and not any(raw_rows[-1]):
                raw_rows.pop()
            headers = raw_rows[0] if raw_rows else []
            data = raw_rows[1:] if len(raw_rows) > 1 else []
            # drop empty data rows and Appointment Grid footer/total rows
            filtered: list[list[str]] = []
            for r in data:
                if not any(_blank(c) for c in r):
                    continue
                if is_footer_or_total_row(headers, r):
                    continue
                filtered.append(r)
            data = filtered
            sheets.append(
                {
                    "sheet_name": name,
                    "headers": headers,
                    "rows": data,
                    "is_empty": 0 if data else 1,
                }
            )
        wb.close()
        return sheets
    if ext == ".csv":
        import csv
        import io

        text = path.read_text(encoding="utf-8-sig", errors="ignore")
        rows = [[_blank(c) for c in r] for r in csv.reader(io.StringIO(text))]
        headers = rows[0] if rows else []
        data = [r for r in rows[1:] if any(r)]
        return [
            {
                "sheet_name": "CSV",
                "headers": headers,
                "rows": data,
                "is_empty": 0 if data else 1,
            }
        ]
    raise ValueError("Only .xlsx and .csv workbooks are supported for engagement import.")


def _split_city_state_zip(value: str) -> tuple[str, str, str]:
    t = _blank(value)
    m = re.match(r"^(.*?),\s*([A-Za-z]{2})\s+(\d{5}(?:-\d{4})?)$", t)
    if m:
        return _blank(m.group(1)), _blank(m.group(2)).upper(), _blank(m.group(3))
    return t, "", ""


def _parse_datetime_text(
    text: str, *, corroborated_year: int | None = None
) -> dict[str, Any]:
    """Delegate to Phase 3A safe parser (no invented year)."""
    return parse_datetime_text(text, corroborated_year=corroborated_year)


def _classify_event(sheet_type: str, date_time_text: str) -> str:
    st = _blank(sheet_type)
    if st == "Quote":
        return "Quote"
    if st == "Purchase Order":
        return "Purchase Order"
    if st == "WebLead":
        return "WebLead"
    return classify_event_type(sheet_type, date_time_text)


def _fingerprint(parts: dict[str, str]) -> str:
    """Legacy helper retained for non-grid callers; Appointment Grid uses
    appointment_grid_fingerprint() instead."""
    payload = "|".join(
        [
            _norm_name(parts.get("client", "")),
            _norm_name(parts.get("sheet", "")),
            _blank(parts.get("record_no", "")),
            _norm_name(parts.get("contact", "")),
            _blank(parts.get("email", "")).lower(),
            _blank(parts.get("datetime", "")),
            _norm_name(parts.get("caller_notes", ""))[:200],
            _norm_name(parts.get("sales_notes", ""))[:200],
            _blank(parts.get("event_type", "")),
        ]
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _match_rev_spec(conn, text: str) -> tuple[int | None, int]:
    raw = _blank(text)
    if not raw:
        return None, 0
    # Rev Spec/Tyler
    name = re.sub(r"(?i)^rev\s*spec\s*/?\s*", "", raw).strip()
    if not name:
        return None, 1
    rows = conn.execute(
        "SELECT id, full_name, email FROM users WHERE full_name IS NOT NULL AND trim(full_name) <> ''"
    ).fetchall()
    hits = []
    n = _norm_name(name)
    for r in rows:
        fn = _norm_name(r["full_name"])
        if n == fn or n in fn.split() or fn.startswith(n) or n in fn:
            hits.append(int(r["id"]))
    if len(hits) == 1:
        return hits[0], 0
    return None, 1 if name else 0


def _parse_money(value: str) -> tuple[float | None, str, int]:
    raw = _blank(value)
    if not raw:
        return None, "", 0
    m = re.search(r"\$?\s*([0-9]{1,3}(?:,[0-9]{3})*(?:\.[0-9]+)?|[0-9]+(?:\.[0-9]+)?)", raw)
    if not m:
        return None, raw, 0
    try:
        amount = float(m.group(1).replace(",", ""))
        return amount, raw, 1
    except ValueError:
        return None, raw, 0


def _normalize_outcome(value: str) -> str:
    t = _blank(value).lower()
    if not t:
        return ""
    if "won" in t:
        return "Won"
    if "lost" in t:
        return "Lost"
    if "progress" in t or "open" in t:
        return "In Progress"
    return ""


def _match_company(
    conn, client_id: int, record_no: str, company_name: str, phone: str, address: str, city: str
) -> tuple[str, int | None, int | None, str, str]:
    """Appointment-grid-safe company match (foreign RN ignored; never overwrite RN)."""
    status, company_id, rel_id, name, brown_rn, _flags = match_company_for_appt_grid(
        conn, client_id, record_no, company_name
    )
    return status, company_id, rel_id, name, brown_rn


def _match_contact(
    conn,
    company_id: int | None,
    email: str,
    contact_name: str,
    phone: str,
    title: str,
) -> tuple[str, int | None, str]:
    status, contact_id, display, _cands, _note = match_contact_for_appt_grid(
        conn, company_id, email, contact_name, phone, title
    )
    return status, contact_id, display


def _load_company_history_notes(conn, client_id: int, company_id: int | None) -> list[dict[str, Any]]:
    if not company_id:
        return []
    rows = conn.execute(
        """
        SELECT h.id, h.event_at, h.event_type, h.note_text
        FROM company_shared_history_events h
        JOIN client_company_relationships ccr ON ccr.company_id = h.company_id
        WHERE ccr.client_id = ? AND h.company_id = ?
        """,
        (client_id, int(company_id)),
    ).fetchall()
    return [dict(r) for r in rows]


def _header_value(raw: dict[str, Any], *names: str) -> str:
    for n in names:
        for k, v in raw.items():
            if _blank(k).lower().rstrip() == n.lower():
                return _blank(v)
    return ""


def _corroborated_year_from_mapped(mapped: dict[str, str]) -> int | None:
    """Optional deterministic year from mapped review resolution only.

    Accepts mapped['_corroborated_year'] when an upstream review process set it.
    Does not invent years from clock or defaults.
    """
    raw = mapped.get("_corroborated_year") or mapped.get("corroborated_year") or ""
    if not _blank(raw):
        return None
    try:
        y = int(str(raw).strip())
    except ValueError:
        return None
    if 1990 <= y <= 2100:
        return y
    return None


def _client_name(conn, client_id: int) -> str:
    row = conn.execute("SELECT name FROM clients WHERE id = ?", (client_id,)).fetchone()
    return _blank(row["name"]) if row else ""


def start_engagement_import(
    client_id: int,
    *,
    filename: str,
    content: bytes,
    user_id: int | None = None,
) -> EngagementImportBatchView:
    user = resolve_staff_actor(user_id)
    if user is None:
        raise PermissionError("User not found.")
    _require_edit(user.id, client_id)
    ensure_engagement_import_schema()

    name = Path(_blank(filename)).name
    ext = Path(name).suffix.lower()
    if ext not in {".xlsx", ".csv"}:
        raise ValueError("Only .xlsx and .csv files are supported.")

    stored = f"{uuid.uuid4().hex}{ext}"
    dest = _client_dir(client_id) / stored
    dest.write_bytes(content)
    now = _now()

    # Also register in client_documents library
    from client_knowledge_data import ensure_client_knowledge_schema

    ensure_client_knowledge_schema()
    with get_connection() as conn:
        cur = conn.execute(
            """
            INSERT INTO client_documents (
                client_id, filename, stored_filename, document_type, mime_type,
                file_size, uploaded_at, uploaded_by_user_id, uploaded_by_name,
                processing_status, created_at, updated_at
            ) VALUES (?, ?, ?, 'Appointment Grid', ?, ?, ?, ?, ?, 'uploaded', ?, ?)
            """,
            (
                client_id,
                name,
                stored,
                ext,
                len(content),
                now,
                user.id,
                _blank(user.full_name) or _blank(user.email),
                now,
                now,
            ),
        )
        doc_id = int(cur.lastrowid)
        cur = conn.execute(
            """
            INSERT INTO client_engagement_import_batches (
                client_id, document_id, filename, status, uploaded_at,
                uploaded_by_user_id, uploaded_by_name
            ) VALUES (?, ?, ?, 'uploaded', ?, ?, ?)
            """,
            (
                client_id,
                doc_id,
                name,
                now,
                user.id,
                _blank(user.full_name) or _blank(user.email),
            ),
        )
        batch_id = int(cur.lastrowid)
        conn.commit()

    sheets = _parse_workbook(dest)
    with get_connection() as conn:
        for sh in sheets:
            detected = _detect_sheet_type(sh["sheet_name"], sh["headers"], sh["rows"])
            if sh["is_empty"]:
                detected = "Ignore"
            mapping = _auto_map_columns(sh["headers"]) if not sh["is_empty"] else {}
            cur = conn.execute(
                """
                INSERT INTO client_engagement_import_sheets (
                    batch_id, client_id, sheet_name, detected_type, user_type,
                    headers_json, column_mapping_json, row_count, is_empty
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    batch_id,
                    client_id,
                    sh["sheet_name"],
                    detected,
                    detected,
                    json.dumps(sh["headers"]),
                    json.dumps(mapping),
                    len(sh["rows"]),
                    int(sh["is_empty"]),
                ),
            )
            sheet_id = int(cur.lastrowid)
            for i, row in enumerate(sh["rows"], start=2):  # Excel-like (header=1)
                raw = {
                    _blank(h) or f"col_{idx}": _blank(row[idx]) if idx < len(row) else ""
                    for idx, h in enumerate(sh["headers"])
                }
                conn.execute(
                    """
                    INSERT INTO client_engagement_import_rows (
                        batch_id, sheet_id, client_id, source_row_number, raw_json,
                        sheet_name, sheet_type, import_status, selected
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, 'staged', 1)
                    """,
                    (
                        batch_id,
                        sheet_id,
                        client_id,
                        i,
                        json.dumps(raw, ensure_ascii=False),
                        sh["sheet_name"],
                        detected,
                    ),
                )
        conn.execute(
            "UPDATE client_engagement_import_batches SET status = 'sheets_ready' WHERE id = ? AND client_id = ?",
            (batch_id, client_id),
        )
        conn.commit()
    return get_engagement_import_batch(client_id, batch_id, user_id=user.id)


def get_engagement_import_batch(
    client_id: int, batch_id: int, *, user_id: int | None = None
) -> EngagementImportBatchView:
    user = resolve_staff_actor(user_id)
    if user is None:
        raise PermissionError("User not found.")
    _require_access(user.id, client_id)
    ensure_engagement_import_schema()
    with get_connection() as conn:
        b = conn.execute(
            "SELECT * FROM client_engagement_import_batches WHERE id = ? AND client_id = ?",
            (batch_id, client_id),
        ).fetchone()
        if not b:
            raise LookupError("Import batch not found for this client.")
        sheets = [
            EngagementImportSheetView(
                sheet_id=int(s["id"]),
                sheet_name=_blank(s["sheet_name"]),
                detected_type=_blank(s["detected_type"]),
                user_type=_blank(s["user_type"]) or _blank(s["detected_type"]),
                headers=json.loads(s["headers_json"] or "[]"),
                column_mapping=json.loads(s["column_mapping_json"] or "{}"),
                row_count=int(s["row_count"] or 0),
                is_empty=bool(s["is_empty"]),
            )
            for s in conn.execute(
                "SELECT * FROM client_engagement_import_sheets WHERE batch_id = ? AND client_id = ? ORDER BY id",
                (batch_id, client_id),
            ).fetchall()
        ]
        try:
            summary = json.loads(b["summary_json"] or "{}")
        except Exception:
            summary = {}
        return EngagementImportBatchView(
            batch_id=int(b["id"]),
            client_id=client_id,
            document_id=int(b["document_id"]) if b["document_id"] else None,
            filename=_blank(b["filename"]),
            status=_blank(b["status"]),
            uploaded_at=_blank(b["uploaded_at"]),
            uploaded_by=_blank(b["uploaded_by_name"]),
            sheets=sheets,
            target_fields=list(TARGET_FIELDS),
            sheet_type_options=list(SHEET_TYPES),
            summary=summary if isinstance(summary, dict) else {},
        )


def update_sheet_classifications(
    client_id: int,
    batch_id: int,
    sheets: list[SheetClassificationUpdate],
    *,
    user_id: int | None = None,
) -> EngagementImportBatchView:
    user = resolve_staff_actor(user_id)
    if user is None:
        raise PermissionError("User not found.")
    _require_edit(user.id, client_id)
    ensure_engagement_import_schema()
    with get_connection() as conn:
        for s in sheets:
            st = _blank(s.user_type)
            if st not in SHEET_TYPES:
                raise ValueError(f"Invalid sheet type: {st}")
            conn.execute(
                """
                UPDATE client_engagement_import_sheets
                SET user_type = ?
                WHERE id = ? AND batch_id = ? AND client_id = ?
                """,
                (st, s.sheet_id, batch_id, client_id),
            )
            conn.execute(
                """
                UPDATE client_engagement_import_rows
                SET sheet_type = ?
                WHERE sheet_id = ? AND batch_id = ? AND client_id = ?
                """,
                (st, s.sheet_id, batch_id, client_id),
            )
        conn.execute(
            "UPDATE client_engagement_import_batches SET status = 'classified' WHERE id = ? AND client_id = ?",
            (batch_id, client_id),
        )
        conn.commit()
    return get_engagement_import_batch(client_id, batch_id, user_id=user.id)


def map_engagement_import(
    client_id: int,
    batch_id: int,
    sheet_mappings: dict[int, dict[str, str]] | None = None,
    *,
    user_id: int | None = None,
) -> EngagementImportPreview:
    """Map columns, match companies/contacts, classify events — preview only."""
    user = resolve_staff_actor(user_id)
    if user is None:
        raise PermissionError("User not found.")
    _require_edit(user.id, client_id)
    ensure_engagement_import_schema()

    with get_connection() as conn:
        b = conn.execute(
            "SELECT * FROM client_engagement_import_batches WHERE id = ? AND client_id = ?",
            (batch_id, client_id),
        ).fetchone()
        if not b:
            raise LookupError("Import batch not found for this client.")
        expected_client = _client_name(conn, client_id)
        sheets = conn.execute(
            "SELECT * FROM client_engagement_import_sheets WHERE batch_id = ? AND client_id = ?",
            (batch_id, client_id),
        ).fetchall()
        sheet_map = {int(s["id"]): dict(s) for s in sheets}

        if sheet_mappings:
            for sid, mapping in sheet_mappings.items():
                conn.execute(
                    """
                    UPDATE client_engagement_import_sheets
                    SET column_mapping_json = ?
                    WHERE id = ? AND batch_id = ? AND client_id = ?
                    """,
                    (json.dumps(mapping), int(sid), batch_id, client_id),
                )
                sheet_map[int(sid)]["column_mapping_json"] = json.dumps(mapping)

        rows = conn.execute(
            """
            SELECT * FROM client_engagement_import_rows
            WHERE batch_id = ? AND client_id = ?
            ORDER BY sheet_id, source_row_number
            """,
            (batch_id, client_id),
        ).fetchall()

        batch_document_id = int(b["document_id"]) if b["document_id"] else None
        doc_sha = _document_file_sha256(conn, client_id, batch_document_id)
        same_document_ids = _same_source_document_ids(
            conn, client_id, batch_document_id, doc_sha
        )

        preview_rows: list[EngagementImportRowView] = []
        for r in rows:
            sh = sheet_map.get(int(r["sheet_id"]))
            if not sh:
                continue
            sheet_type = _blank(sh["user_type"]) or _blank(sh["detected_type"])
            if sheet_type == "Ignore" or int(sh["is_empty"] or 0):
                conn.execute(
                    """
                    UPDATE client_engagement_import_rows
                    SET sheet_type = 'Ignore', import_status = 'ignored', selected = 0,
                        event_type = 'Other', needs_review = 0
                    WHERE id = ? AND client_id = ?
                    """,
                    (int(r["id"]), client_id),
                )
                continue

            try:
                raw = json.loads(r["raw_json"] or "{}")
            except Exception:
                raw = {}
            try:
                mapping = json.loads(sh["column_mapping_json"] or "{}")
            except Exception:
                mapping = {}
            mapped: dict[str, Any] = {}
            for target, header in mapping.items():
                mapped[target] = _blank(raw.get(header, ""))
            # Unmapped Forecast / QUOTED / WON (header trailing-space tolerant)
            forecast = _header_value(raw, "Forecast")
            quoted_label = _header_value(raw, "QUOTED", "Quoted")
            won_val = _header_value(raw, "WON", "Won")
            mapped["forecast"] = forecast
            mapped["quoted_label"] = quoted_label
            mapped["won"] = won_val

            c_status, company_id, rel_id, c_name, c_rn, co_flags = match_company_for_appt_grid(
                conn,
                client_id,
                mapped.get("external_record_no", ""),
                mapped.get("company", ""),
            )
            t_status, contact_id, t_name, _cands, contact_note = match_contact_for_appt_grid(
                conn,
                company_id,
                mapped.get("email", ""),
                mapped.get("contact", ""),
                mapped.get("phone", ""),
                mapped.get("title", ""),
            )
            hist = _load_company_history_notes(conn, client_id, company_id)
            # Deterministic year resolution: mapped override OR unique history/caller corroboration
            corroborated_year = _corroborated_year_from_mapped(mapped)
            year_evidence = ""
            if corroborated_year is None:
                corroborated_year, year_evidence = corroborate_year_from_history(
                    appointment_date_time=mapped.get("appointment_date_time", ""),
                    caller_notes=mapped.get("caller_notes", ""),
                    history_events=hist,
                )
            else:
                year_evidence = "mapped_review_resolution"
            if corroborated_year is not None:
                mapped["_corroborated_year"] = str(corroborated_year)
                mapped["_year_evidence"] = year_evidence
            dt = _parse_datetime_text(
                mapped.get("appointment_date_time", ""),
                corroborated_year=corroborated_year,
            )
            event_type = _classify_event(sheet_type, mapped.get("appointment_date_time", ""))
            client_val = mapped.get("northstar_client", "")
            if client_val and _norm_client(client_val) != _norm_client(expected_client):
                client_validation = "MISMATCH"
            else:
                client_validation = "OK"

            rev_uid, rev_review = _match_rev_spec(conn, mapped.get("revenue_specialist", ""))
            amount, src_amt, quote_sig = _parse_money(mapped.get("dollars_quoted", ""))
            outcome_n = _normalize_outcome(mapped.get("outcome", ""))

            hist_blob = norm_text("\n".join(_blank(h.get("note_text")) for h in hist))
            narr_class, narr_why = classify_appointment_narrative(
                event_type=event_type,
                contact_name=mapped.get("contact", ""),
                caller_notes=mapped.get("caller_notes", ""),
                appointment_date_time=mapped.get("appointment_date_time", ""),
                resolved_date=dt.get("event_date") or "",
                history_events=hist,
            )
            field_classes = {
                "appointment_narrative": narr_class,
                "appointment_date_time": (
                    "AMBIGUOUS"
                    if dt.get("datetime_needs_review")
                    and event_type.startswith("Appointment")
                    else classify_field_vs_history(
                        dt.get("event_date") or mapped.get("appointment_date_time", ""),
                        hist_blob,
                    )
                ),
                "appointment_grade": classify_field_vs_history(
                    mapped.get("appointment_grade", ""), hist_blob
                ),
                "caller_notes": classify_field_vs_history(
                    mapped.get("caller_notes", ""), hist_blob
                ),
                "sales_notes": classify_field_vs_history(
                    mapped.get("sales_notes", ""), hist_blob
                ),
                "dollars_quoted": classify_field_vs_history(
                    mapped.get("dollars_quoted", ""), hist_blob
                ),
                "outcome": classify_field_vs_history(mapped.get("outcome", ""), hist_blob),
                "forecast": classify_field_vs_history(forecast, hist_blob),
                "quoted_label": classify_field_vs_history(quoted_label, hist_blob),
                "won": classify_field_vs_history(won_val, hist_blob),
                "revenue_specialist": classify_field_vs_history(
                    mapped.get("revenue_specialist", ""), hist_blob
                ),
            }
            proposed_action, action_reason = propose_row_action(
                field_classes,
                company_status=c_status,
                contact_status=t_status,
                datetime_needs_review=bool(dt.get("datetime_needs_review")),
                event_type=event_type,
            )
            fqw_block = compose_forecast_quoted_won_block(
                forecast=forecast, quoted=quoted_label, won=won_val
            )
            composed_sales = merge_sales_notes_with_fqw(
                mapped.get("sales_notes", ""), fqw_block
            )
            caller_for_event = build_delta_caller_notes(
                narrative_class=narr_class,
                caller_notes=mapped.get("caller_notes", ""),
            )
            mapped["_field_classes"] = json.dumps(field_classes, ensure_ascii=False)
            mapped["_proposed_action"] = proposed_action
            mapped["_action_reason"] = action_reason
            mapped["_composed_sales_notes"] = composed_sales
            mapped["_caller_notes_for_event"] = caller_for_event
            mapped["_narrative_why"] = narr_why

            thread_key = (
                f"{client_id}:"
                f"{_blank(c_rn) or _blank(mapped.get('external_record_no')) or _norm_name(mapped.get('company',''))}:"
                f"{_norm_name(mapped.get('contact',''))}"
            )
            fp = appointment_grid_fingerprint(
                client_id=client_id,
                source_document_sha256=doc_sha,
                source_sheet=_blank(sh["sheet_name"]),
                source_row=int(r["source_row_number"] or 0),
                source_rn=mapped.get("external_record_no", ""),
                company_name=mapped.get("company", ""),
                contact_name=mapped.get("contact", ""),
                email=mapped.get("email", ""),
                event_type=event_type,
                event_date=dt.get("event_date") or "",
                source_datetime_text=mapped.get("appointment_date_time", ""),
                grade=mapped.get("appointment_grade", ""),
                dollars_quoted=str(mapped.get("dollars_quoted", "")),
                outcome=mapped.get("outcome", ""),
                forecast=forecast,
                quoted_label=quoted_label,
                won=won_val,
            )
            fp_event_id = already_imported_by_fingerprint(conn, client_id, fp)
            prov_event_id = already_imported_by_provenance(
                conn,
                client_id=client_id,
                source_sheet=_blank(sh["sheet_name"]),
                source_row=int(r["source_row_number"] or 0),
                same_document_ids=same_document_ids,
            )
            dup, dup_kind, dup_event_id = classify_already_imported(
                fingerprint_event_id=fp_event_id,
                provenance_event_id=prov_event_id,
            )
            mapped["_duplicate_kind"] = dup_kind
            if dup_event_id:
                mapped["_already_imported_event_id"] = str(dup_event_id)

            flags: list[str] = list(co_flags)
            if contact_note and t_status == "REVIEW":
                flags.append(contact_note)
            if client_validation == "MISMATCH":
                flags.append("Client mismatch")
            if dt["datetime_needs_review"]:
                flags.append("Date/Time Needs Review")
            if rev_review:
                flags.append("Rev Spec Needs Review")
            if c_status in {"POSSIBLE MATCH", "CONFLICT", "NEW", "REVIEW"}:
                flags.append(f"Company {c_status}")
            if t_status in {"POSSIBLE MATCH", "NEW", "REVIEW"}:
                flags.append(f"Contact {t_status}")
            if quote_sig:
                flags.append("Potential Quote Signal")
            if proposed_action:
                flags.append(f"Plan:{proposed_action}")
            if dup:
                flags.append(dup)

            import_status = "duplicate" if dup else (
                "skipped" if proposed_action in {"SKIP", "SKIP_FULL_DUPLICATE"} else "previewed"
            )
            blocks = row_blocks_confirm(
                proposed_action=proposed_action,
                company_status=c_status,
                contact_status=t_status,
                datetime_needs_review=bool(dt.get("datetime_needs_review")),
                event_type=event_type,
                client_validation=client_validation,
                import_status=import_status,
            )
            selected = 1
            if (
                client_validation == "MISMATCH"
                or dup
                or sheet_type == "Ignore"
                or proposed_action in {"SKIP", "SKIP_FULL_DUPLICATE"}
            ):
                selected = 0
            if blocks:
                selected = 0
            # Already-imported rows never block a rerun. Informational flags stay
            # on the row for diagnostics.
            needs_review = 0 if dup else (1 if blocks else 0)

            conn.execute(
                """
                UPDATE client_engagement_import_rows SET
                    mapped_json = ?, sheet_type = ?, event_type = ?,
                    source_date_time_text = ?, event_date = ?, event_time = ?,
                    timezone = ?, meeting_type = ?, datetime_needs_review = ?,
                    company_match_status = ?, company_id = ?, relationship_id = ?,
                    company_name = ?, company_record_no = ?,
                    contact_match_status = ?, contact_id = ?, contact_name = ?,
                    client_validation = ?, rev_spec_user_id = ?, source_rev_spec_text = ?,
                    rev_spec_needs_review = ?, source_appointment_grade = ?,
                    quoted_amount = ?, source_quoted_value = ?, potential_quote_signal = ?,
                    outcome_normalized = ?, source_outcome = ?,
                    thread_key = ?, source_row_fingerprint = ?, duplicate_status = ?,
                    import_status = ?, selected = ?, needs_review = ?, review_flags_json = ?
                WHERE id = ? AND client_id = ?
                """,
                (
                    json.dumps(mapped, ensure_ascii=False),
                    sheet_type,
                    event_type,
                    dt["source_date_time_text"],
                    dt["event_date"],
                    dt["event_time"],
                    dt["timezone"],
                    dt["meeting_type"],
                    int(dt["datetime_needs_review"]),
                    c_status,
                    company_id,
                    rel_id,
                    c_name or mapped.get("company", ""),
                    # Always Brown RN when matched — never write grid RN over Brown identity
                    c_rn,
                    t_status,
                    contact_id,
                    t_name or mapped.get("contact", ""),
                    client_validation,
                    rev_uid,
                    mapped.get("revenue_specialist", ""),
                    rev_review,
                    mapped.get("appointment_grade", ""),
                    amount,
                    src_amt,
                    quote_sig,
                    outcome_n,
                    mapped.get("outcome", ""),
                    thread_key,
                    fp,
                    dup,
                    import_status,
                    selected,
                    needs_review,
                    json.dumps(flags),
                    int(r["id"]),
                    client_id,
                ),
            )

        conn.execute(
            "UPDATE client_engagement_import_batches SET status = 'previewed' WHERE id = ? AND client_id = ?",
            (batch_id, client_id),
        )
        conn.commit()

    return get_engagement_import_preview(client_id, batch_id, user_id=user.id)


def _row_to_view(r: Any) -> EngagementImportRowView:
    d = dict(r)
    try:
        mapped = json.loads(d.get("mapped_json") or "{}")
    except Exception:
        mapped = {}
    try:
        flags = json.loads(d.get("review_flags_json") or "[]")
    except Exception:
        flags = []
    return EngagementImportRowView(
        row_id=int(d["id"]),
        sheet_name=_blank(d.get("sheet_name")),
        sheet_type=_blank(d.get("sheet_type")),
        source_row_number=int(d.get("source_row_number") or 0),
        mapped=mapped if isinstance(mapped, dict) else {},
        event_type=_blank(d.get("event_type")),
        source_date_time_text=_blank(d.get("source_date_time_text")),
        event_date=_blank(d.get("event_date")),
        event_time=_blank(d.get("event_time")),
        timezone=_blank(d.get("timezone")),
        meeting_type=_blank(d.get("meeting_type")),
        datetime_needs_review=bool(d.get("datetime_needs_review")),
        company_match_status=_blank(d.get("company_match_status")) or "NEW",
        company_id=int(d["company_id"]) if d.get("company_id") else None,
        relationship_id=int(d["relationship_id"]) if d.get("relationship_id") else None,
        company_name=_blank(d.get("company_name")),
        company_record_no=_blank(d.get("company_record_no")),
        contact_match_status=_blank(d.get("contact_match_status")) or "NEW",
        contact_id=int(d["contact_id"]) if d.get("contact_id") else None,
        contact_name=_blank(d.get("contact_name")),
        client_validation=_blank(d.get("client_validation")) or "OK",
        source_rev_spec_text=_blank(d.get("source_rev_spec_text")),
        rev_spec_user_id=int(d["rev_spec_user_id"]) if d.get("rev_spec_user_id") else None,
        rev_spec_needs_review=bool(d.get("rev_spec_needs_review")),
        source_appointment_grade=_blank(d.get("source_appointment_grade")),
        quoted_amount=float(d["quoted_amount"]) if d.get("quoted_amount") is not None else None,
        source_quoted_value=_blank(d.get("source_quoted_value")),
        potential_quote_signal=bool(d.get("potential_quote_signal")),
        outcome_normalized=_blank(d.get("outcome_normalized")),
        source_outcome=_blank(d.get("source_outcome")),
        thread_key=_blank(d.get("thread_key")),
        source_row_fingerprint=_blank(d.get("source_row_fingerprint")),
        duplicate_status=_blank(d.get("duplicate_status")),
        import_status=_blank(d.get("import_status")),
        selected=bool(d.get("selected")),
        needs_review=bool(d.get("needs_review")),
        review_flags=flags if isinstance(flags, list) else [],
    )


def _build_summary(rows: list[EngagementImportRowView]) -> EngagementImportSummary:
    active = [r for r in rows if r.sheet_type != "Ignore" and r.import_status != "ignored"]
    return EngagementImportSummary(
        rows_detected=len(active),
        rows_selected=sum(1 for r in active if r.selected),
        matched_companies=sum(1 for r in active if r.company_match_status == "MATCHED"),
        possible_company_matches=sum(
            1 for r in active if r.company_match_status == "POSSIBLE MATCH"
        ),
        new_companies=sum(1 for r in active if r.company_match_status == "NEW"),
        company_conflicts=sum(1 for r in active if r.company_match_status == "CONFLICT"),
        matched_contacts=sum(1 for r in active if r.contact_match_status == "MATCHED"),
        possible_contact_matches=sum(
            1 for r in active if r.contact_match_status == "POSSIBLE MATCH"
        ),
        new_contacts=sum(1 for r in active if r.contact_match_status == "NEW"),
        duplicates_skipped=sum(1 for r in active if r.duplicate_status),
        rows_needing_review=sum(1 for r in active if r.needs_review),
        appointment_events=sum(
            1 for r in active if r.event_type.startswith("Appointment")
        ),
        send_information_events=sum(1 for r in active if r.event_type == "Send Information"),
        datetime_parse_success=sum(
            1 for r in active if r.event_date and not r.datetime_needs_review
        ),
        datetime_needs_review=sum(1 for r in active if r.datetime_needs_review),
        client_mismatches=sum(1 for r in active if r.client_validation == "MISMATCH"),
    )


def get_engagement_import_preview(
    client_id: int,
    batch_id: int,
    *,
    user_id: int | None = None,
    filter_status: str | None = None,
) -> EngagementImportPreview:
    user = resolve_staff_actor(user_id)
    if user is None:
        raise PermissionError("User not found.")
    _require_access(user.id, client_id)
    ensure_engagement_import_schema()
    batch = get_engagement_import_batch(client_id, batch_id, user_id=user.id)
    with get_connection() as conn:
        rows = [
            _row_to_view(r)
            for r in conn.execute(
                """
                SELECT * FROM client_engagement_import_rows
                WHERE batch_id = ? AND client_id = ?
                ORDER BY sheet_id, source_row_number
                """,
                (batch_id, client_id),
            ).fetchall()
        ]
    summary = _build_summary(rows)
    with get_connection() as conn:
        conn.execute(
            "UPDATE client_engagement_import_batches SET summary_json = ? WHERE id = ? AND client_id = ?",
            (json.dumps(summary.model_dump()), batch_id, client_id),
        )
        conn.commit()
    batch.summary = summary.model_dump()

    def _match_filter(r: EngagementImportRowView) -> bool:
        if not filter_status:
            return True
        f = filter_status.lower()
        if f == "matched":
            return r.company_match_status == "MATCHED"
        if f == "possible match":
            return r.company_match_status == "POSSIBLE MATCH"
        if f == "new":
            return r.company_match_status == "NEW"
        if f == "conflict":
            return r.company_match_status == "CONFLICT"
        if f == "duplicate":
            return bool(r.duplicate_status)
        if f == "needs review":
            return r.needs_review
        return True

    filtered = [r for r in rows if r.sheet_type != "Ignore" and _match_filter(r)]
    return EngagementImportPreview(batch=batch, rows=filtered, summary=summary)


def confirm_engagement_import(
    client_id: int,
    batch_id: int,
    body: EngagementImportConfirmRequest,
    *,
    user_id: int | None = None,
) -> EngagementImportBatchView:
    """Write selected non-duplicate rows into client_sales_events, then project
    company-level Quote milestones from structured quote evidence.

    Does NOT create companies/contacts unless body.create_new_ids explicitly lists rows
    (Phase 3 default: empty — skip NEW company/contact creation).
    Does NOT change statuses or overwrite legacy notes.
    Does NOT invent Purchase Order milestones from Won/PO narrative.
    """
    user = resolve_staff_actor(user_id)
    if user is None:
        raise PermissionError("User not found.")
    _require_edit(user.id, client_id)
    if not body.confirm:
        raise ValueError("Import not confirmed.")
    ensure_engagement_import_schema()
    now = _now()
    user_name = _blank(user.full_name) or _blank(user.email)

    with get_connection() as conn:
        b = conn.execute(
            "SELECT * FROM client_engagement_import_batches WHERE id = ? AND client_id = ?",
            (batch_id, client_id),
        ).fetchone()
        if not b:
            raise LookupError("Import batch not found for this client.")
        if _blank(b["status"]) == "imported":
            return get_engagement_import_batch(client_id, batch_id, user_id=user.id)

        batch_document_id = int(b["document_id"]) if b["document_id"] else None
        doc_sha = _document_file_sha256(conn, client_id, batch_document_id)
        same_document_ids = _same_source_document_ids(
            conn, client_id, batch_document_id, doc_sha
        )

        selected_ids = set(int(x) for x in (body.row_ids or []))
        create_new = set(int(x) for x in (body.create_new_ids or []))
        rows = conn.execute(
            """
            SELECT * FROM client_engagement_import_rows
            WHERE batch_id = ? AND client_id = ?
            ORDER BY sheet_id, source_row_number
            """,
            (batch_id, client_id),
        ).fetchall()

        # REVIEW gate: never silently drop unresolved critical rows
        blocking: list[int] = []
        for r in rows:
            d = dict(r)
            if _blank(d.get("sheet_type")) == "Ignore":
                continue
            if _blank(d.get("import_status")) in {"ignored", "skipped", "duplicate"}:
                continue
            if _blank(d.get("duplicate_status")):
                continue
            try:
                mapped_chk = json.loads(d.get("mapped_json") or "{}")
            except Exception:
                mapped_chk = {}
            action = _blank(mapped_chk.get("_proposed_action"))
            if action in {"SKIP", "SKIP_FULL_DUPLICATE"}:
                continue
            if row_blocks_confirm(
                proposed_action=action
                or ("REVIEW" if int(d.get("needs_review") or 0) else "CREATE_STRUCTURED_EVENT"),
                company_status=_blank(d.get("company_match_status")),
                contact_status=_blank(d.get("contact_match_status")),
                datetime_needs_review=bool(d.get("datetime_needs_review")),
                event_type=_blank(d.get("event_type")),
                import_status=_blank(d.get("import_status")),
                client_validation=_blank(d.get("client_validation")),
            ) or int(d.get("needs_review") or 0):
                blocking.append(int(d["id"]))
        if blocking:
            raise ValueError(
                f"Cannot confirm engagement import: {len(blocking)} row(s) require REVIEW "
                f"(ids={blocking[:20]}{'…' if len(blocking) > 20 else ''}). "
                "Resolve company/contact/date issues or mark those rows skipped first."
            )

        # Snapshot relationship RNs — confirm must never overwrite external_record_no
        rn_before = {
            int(x["id"]): _blank(x["external_record_no"])
            for x in conn.execute(
                "SELECT id, external_record_no FROM client_company_relationships WHERE client_id = ?",
                (client_id,),
            ).fetchall()
        }

        imported = 0
        for r in rows:
            d = dict(r)
            if _blank(d.get("sheet_type")) == "Ignore":
                continue
            if _blank(d.get("duplicate_status")):
                continue
            if _blank(d.get("client_validation")) == "MISMATCH":
                continue
            if _blank(d.get("import_status")) in {"skipped", "ignored"}:
                continue
            rid = int(d["id"])
            if selected_ids and rid not in selected_ids:
                continue
            if not selected_ids and not int(d.get("selected") or 0):
                continue
            # Skip NEW company/contact rows unless explicitly allowed
            if d.get("company_match_status") == "NEW" and rid not in create_new:
                continue
            if d.get("contact_match_status") == "NEW" and rid not in create_new:
                # Still allow event with null contact if company matched
                if d.get("company_match_status") != "MATCHED":
                    continue
            if d.get("company_match_status") not in {"MATCHED"} and rid not in create_new:
                continue

            fp = _blank(d.get("source_row_fingerprint"))
            fp_event_id = already_imported_by_fingerprint(conn, client_id, fp)
            prov_event_id = already_imported_by_provenance(
                conn,
                client_id=client_id,
                source_sheet=_blank(d.get("sheet_name")),
                source_row=int(d.get("source_row_number") or 0),
                same_document_ids=same_document_ids,
            )
            if fp_event_id or prov_event_id:
                continue

            try:
                mapped = json.loads(d.get("mapped_json") or "{}")
            except Exception:
                mapped = {}

            sales_notes = _blank(mapped.get("_composed_sales_notes")) or _blank(
                mapped.get("sales_notes")
            )
            caller_notes = (
                mapped.get("_caller_notes_for_event")
                if "_caller_notes_for_event" in mapped
                else mapped.get("caller_notes")
            )

            cur = conn.execute(
                """
                INSERT INTO client_sales_events (
                    client_id, relationship_id, company_id, contact_id, thread_key,
                    event_type, event_family, source_date_time_text, event_date, event_time,
                    timezone, meeting_type, company_name, contact_name, contact_title,
                    phone, email, address, city_state_zip, caller_notes, sales_notes,
                    source_appointment_grade, source_rev_spec_text, rev_spec_user_id,
                    quoted_amount, source_quoted_value, potential_quote_signal,
                    outcome_normalized, source_outcome,
                    source_document_id, source_batch_id, source_sheet, source_row,
                    source_file_name, source_record_number, source_row_fingerprint,
                    imported_at, imported_by_user_id, imported_by_name, created_at
                ) VALUES (
                    ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?, ?,
                    ?, ?, ?,
                    ?, ?, ?,
                    ?, ?,
                    ?, ?, ?, ?,
                    ?, ?, ?,
                    ?, ?, ?, ?
                )
                """,
                (
                    client_id,
                    d.get("relationship_id"),
                    d.get("company_id"),
                    d.get("contact_id"),
                    _blank(d.get("thread_key")),
                    _blank(d.get("event_type")),
                    "Appointment"
                    if _blank(d.get("event_type")).startswith("Appointment")
                    else _blank(d.get("event_type")),
                    _blank(d.get("source_date_time_text")),
                    _blank(d.get("event_date")),
                    _blank(d.get("event_time")),
                    _blank(d.get("timezone")),
                    _blank(d.get("meeting_type")),
                    _blank(d.get("company_name")),
                    _blank(d.get("contact_name")),
                    _blank(mapped.get("title")),
                    _blank(mapped.get("phone")),
                    _blank(mapped.get("email")),
                    _blank(mapped.get("address")),
                    _blank(mapped.get("city_state_zip")),
                    _blank(caller_notes),
                    _blank(sales_notes),
                    _blank(d.get("source_appointment_grade")),
                    _blank(d.get("source_rev_spec_text")),
                    d.get("rev_spec_user_id"),
                    d.get("quoted_amount"),
                    _blank(d.get("source_quoted_value")),
                    int(d.get("potential_quote_signal") or 0),
                    _blank(d.get("outcome_normalized")),
                    _blank(d.get("source_outcome")),
                    b["document_id"],
                    batch_id,
                    _blank(d.get("sheet_name")),
                    int(d.get("source_row_number") or 0),
                    _blank(b["filename"]),
                    _blank(d.get("company_record_no")) or _blank(mapped.get("external_record_no")),
                    fp,
                    now,
                    user.id,
                    user_name,
                    now,
                ),
            )
            event_id = int(cur.lastrowid)
            conn.execute(
                """
                UPDATE client_engagement_import_rows
                SET import_status = 'imported', event_id = ?
                WHERE id = ? AND client_id = ?
                """,
                (event_id, rid, client_id),
            )
            imported += 1

        # Link reschedule threads: same thread_key, appointment family
        thread_rows = conn.execute(
            """
            SELECT id, thread_key, event_type, event_date, id
            FROM client_sales_events
            WHERE client_id = ? AND source_batch_id = ?
              AND event_type LIKE 'Appointment%'
            ORDER BY thread_key, event_date, id
            """,
            (client_id, batch_id),
        ).fetchall()
        by_thread: dict[str, list[Any]] = {}
        for tr in thread_rows:
            by_thread.setdefault(_blank(tr["thread_key"]), []).append(tr)
        for items in by_thread.values():
            if len(items) < 2:
                continue
            parent_id = int(items[0]["id"])
            for child in items[1:]:
                if _blank(child["event_type"]) == "Appointment Rescheduled":
                    conn.execute(
                        "UPDATE client_sales_events SET parent_event_id = ? WHERE id = ? AND client_id = ?",
                        (parent_id, int(child["id"]), client_id),
                    )

        conn.execute(
            """
            UPDATE client_engagement_import_batches
            SET status = 'imported', confirmed_at = ?, confirmed_by_name = ?,
                notes = ?
            WHERE id = ? AND client_id = ?
            """,
            (now, user_name, f"Imported {imported} events", batch_id, client_id),
        )
        rn_after = {
            int(x["id"]): _blank(x["external_record_no"])
            for x in conn.execute(
                "SELECT id, external_record_no FROM client_company_relationships WHERE client_id = ?",
                (client_id,),
            ).fetchall()
        }
        if rn_before != rn_after:
            raise RuntimeError(
                "Safety abort: engagement confirm attempted to change "
                "client_company_relationships.external_record_no"
            )

        # Project company-level Quote milestones from structured quote evidence.
        # Does not change relationship status; does not invent Purchase Orders.
        from engagement_quote_milestones import project_quote_milestones_for_batch

        project_quote_milestones_for_batch(
            conn,
            client_id,
            batch_id,
            created_by=user_name or "engagement_import",
        )

        conn.commit()
    return get_engagement_import_batch(client_id, batch_id, user_id=user.id)


def list_sales_events(
    client_id: int,
    *,
    user_id: int | None = None,
    company_id: int | None = None,
    contact_id: int | None = None,
    event_family: str | None = None,
    rev_spec: str | None = None,
    outcome: str | None = None,
    appointment_grade: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    company_name: str | None = None,
    contact_name: str | None = None,
    limit: int = 200,
) -> list[SalesEventView]:
    user = resolve_staff_actor(user_id)
    if user is None:
        raise PermissionError("User not found.")
    _require_access(user.id, client_id)
    ensure_engagement_import_schema()
    sql = "SELECT * FROM client_sales_events WHERE client_id = ?"
    params: list[Any] = [client_id]
    if company_id is not None:
        sql += " AND company_id = ?"
        params.append(company_id)
    if contact_id is not None:
        sql += " AND contact_id = ?"
        params.append(contact_id)
    if event_family == "appointment":
        sql += " AND event_type LIKE 'Appointment%'"
    elif event_family == "engagement":
        sql += " AND event_type = 'Send Information'"
    if rev_spec:
        sql += " AND lower(COALESCE(source_rev_spec_text,'')) LIKE ?"
        params.append(f"%{_blank(rev_spec).lower()}%")
    if outcome:
        sql += " AND (lower(COALESCE(outcome_normalized,'')) = ? OR lower(COALESCE(source_outcome,'')) LIKE ?)"
        o = _blank(outcome).lower()
        params.extend([o, f"%{o}%"])
    if appointment_grade:
        sql += " AND lower(COALESCE(source_appointment_grade,'')) = ?"
        params.append(_blank(appointment_grade).lower())
    if date_from:
        sql += " AND COALESCE(event_date,'') >= ?"
        params.append(_blank(date_from))
    if date_to:
        sql += " AND COALESCE(event_date,'') <= ?"
        params.append(_blank(date_to))
    if company_name:
        sql += " AND lower(COALESCE(company_name,'')) LIKE ?"
        params.append(f"%{_blank(company_name).lower()}%")
    if contact_name:
        sql += " AND lower(COALESCE(contact_name,'')) LIKE ?"
        params.append(f"%{_blank(contact_name).lower()}%")
    sql += " ORDER BY COALESCE(event_date,'' ) DESC, id DESC LIMIT ?"
    params.append(max(1, min(limit, 500)))
    with get_connection() as conn:
        rows = conn.execute(sql, tuple(params)).fetchall()
    return [_event_to_view(r) for r in rows]


def _event_to_view(r: Any) -> SalesEventView:
    d = dict(r)
    return SalesEventView(
        event_id=int(d["id"]),
        client_id=int(d["client_id"]),
        relationship_id=int(d["relationship_id"]) if d.get("relationship_id") else None,
        company_id=int(d["company_id"]) if d.get("company_id") else None,
        contact_id=int(d["contact_id"]) if d.get("contact_id") else None,
        thread_key=_blank(d.get("thread_key")),
        parent_event_id=int(d["parent_event_id"]) if d.get("parent_event_id") else None,
        event_type=_blank(d.get("event_type")),
        event_family=_blank(d.get("event_family")),
        source_date_time_text=_blank(d.get("source_date_time_text")),
        event_date=_blank(d.get("event_date")),
        event_time=_blank(d.get("event_time")),
        timezone=_blank(d.get("timezone")),
        meeting_type=_blank(d.get("meeting_type")),
        company_name=_blank(d.get("company_name")),
        contact_name=_blank(d.get("contact_name")),
        contact_title=_blank(d.get("contact_title")),
        phone=_blank(d.get("phone")),
        email=_blank(d.get("email")),
        address=_blank(d.get("address")),
        city_state_zip=_blank(d.get("city_state_zip")),
        caller_notes=_blank(d.get("caller_notes")),
        sales_notes=_blank(d.get("sales_notes")),
        source_appointment_grade=_blank(d.get("source_appointment_grade")),
        source_rev_spec_text=_blank(d.get("source_rev_spec_text")),
        rev_spec_user_id=int(d["rev_spec_user_id"]) if d.get("rev_spec_user_id") else None,
        quoted_amount=float(d["quoted_amount"]) if d.get("quoted_amount") is not None else None,
        source_quoted_value=_blank(d.get("source_quoted_value")),
        potential_quote_signal=bool(d.get("potential_quote_signal")),
        outcome_normalized=_blank(d.get("outcome_normalized")),
        source_outcome=_blank(d.get("source_outcome")),
        source_document_id=int(d["source_document_id"]) if d.get("source_document_id") else None,
        source_batch_id=int(d["source_batch_id"]) if d.get("source_batch_id") else None,
        source_sheet=_blank(d.get("source_sheet")),
        source_row=int(d.get("source_row") or 0),
        source_file_name=_blank(d.get("source_file_name")),
        source_record_number=_blank(d.get("source_record_number")),
        source_row_fingerprint=_blank(d.get("source_row_fingerprint")),
        imported_at=_blank(d.get("imported_at")),
        imported_by=_blank(d.get("imported_by_name")),
    )
